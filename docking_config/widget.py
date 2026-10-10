"""PyQt workbench for batch docking configuration updates."""

from __future__ import annotations

import csv
import platform
import queue
import re
import threading
from collections.abc import Callable
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

from PyQt6.QtCore import QAbstractTableModel, QModelIndex, QTimer, Qt, pyqtSignal
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFrame,
    QHeaderView,
    QHBoxLayout,
    QLineEdit,
    QPlainTextEdit,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QLabel,
    QMessageBox,
    QSizePolicy,
    QTableView,
    QWidget,
)

try:
    from qfluentwidgets import (
        BodyLabel,
        LineEdit,
        PlainTextEdit,
        PrimaryPushButton,
        PushButton,
        TableWidget,
        TextEdit,
    )
except Exception:  # pragma: no cover - native fallback
    BodyLabel = QLabel
    LineEdit = QLineEdit
    PlainTextEdit = QPlainTextEdit
    PrimaryPushButton = QPushButton
    PushButton = QPushButton
    TableWidget = QTableWidget
    TextEdit = QPlainTextEdit

from .api import ALL_SITE_QUERY_TARGET, DockingConfigApi, MAX_PARALLEL_REQUESTS
from .template import DockingTemplateChange, export_docking_template, import_docking_template
from .workflow import DockingBatchResult, DockingConfigWorkflow

TransportFactory = Callable[[], Any]
ConfirmBatch = Callable[[QWidget, int], bool]


DEFAULT_REQUEST_INTERVAL_SECONDS = 0.5
DEFAULT_PARALLEL_REQUESTS = 5
MIN_REQUEST_INTERVAL_SECONDS = 0.0
MAX_REQUEST_INTERVAL_SECONDS = 10.0

STATUS_UNLOADED = "未提交"
STATUS_SUCCESS = "成功"
STATUS_FAILED = "失败"
STATUS_PENDING = "待提交"
STATUS_READY = "已预载"
STATUS_NOCHANGE = "无差异"

FIELD_LABELS = {
    "vendor_resource_id": "供应商资源ID",
    "ticket_code": "票种代码",
    "preorder_code": "支付前对接",
    "allow_sync_sale_price": "同步卖价",
    "allow_sync_cost_price": "同步底价",
    "allow_sync_sale_status": "同步开关班",
}


def excel_file_dialog_options() -> QFileDialog.Option:
    options = QFileDialog.Option(0)
    if platform.system() == "Darwin":
        options |= QFileDialog.Option.DontUseNativeDialog
    return options


@dataclass
class DockingPreviewRow:
    row_kind: str
    source_row: int
    product_id: str
    product_name: str
    resource_id: str
    resource_name: str
    field_name: str
    old_value: str
    new_value: str
    status: str
    message: str = ""
    docking_code: str = ""

    def as_report_row(self) -> tuple[str, ...]:
        return (
            self.status,
            str(self.source_row),
            self.product_id,
            self.product_name,
            self.resource_id,
            self.resource_name,
            self.docking_code,
            self.field_name,
            self.old_value,
            self.new_value,
            self.message,
        )


class DockingResultTableModel(QAbstractTableModel):
    COLUMNS = (
        "状态",
        "来源行",
        "产品ID",
        "产品名称",
        "资源ID",
        "资源名称",
        "当前对接代码",
        "字段",
        "原值",
        "目标值",
        "说明",
    )

    def __init__(self, rows: list[DockingPreviewRow] | None = None, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.rows: list[DockingPreviewRow] = rows or []

    def rowCount(self, _parent: QModelIndex = QModelIndex()) -> int:  # pragma: no cover - Qt model contract
        return 0 if _parent.isValid() else len(self.rows)

    def columnCount(self, _parent: QModelIndex = QModelIndex()) -> int:  # pragma: no cover - Qt model contract
        return 0 if _parent.isValid() else len(self.COLUMNS)

    def data(self, index: QModelIndex, role: int = Qt.ItemDataRole.DisplayRole) -> str | None:  # pragma: no cover - UI
        if not index.isValid():
            return None
        if role not in {Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.ToolTipRole}:
            return None
        row = self.rows[index.row()]
        values = (
            row.status,
            str(row.source_row),
            row.product_id,
            row.product_name,
            row.resource_id,
            row.resource_name,
            row.docking_code,
            row.field_name,
            row.old_value,
            row.new_value,
            row.message,
        )
        return values[index.column()]

    def headerData(
        self,
        section: int,
        orientation: Qt.Orientation,
        role: int = Qt.ItemDataRole.DisplayRole,
    ) -> str | None:  # pragma: no cover - Qt model contract
        if role != Qt.ItemDataRole.DisplayRole or orientation != Qt.Orientation.Horizontal:
            return None
        return self.COLUMNS[section]

    def set_rows(self, rows: list[DockingPreviewRow]) -> None:
        self.beginResetModel()
        self.rows = list(rows)
        self.endResetModel()

    def row_at(self, row: int) -> DockingPreviewRow | None:
        if row < 0 or row >= len(self.rows):
            return None
        return self.rows[row]


class ProductIdsDialog(QDialog):
    def __init__(self, parent: QWidget | None = None):
        super().__init__(parent)
        self.setWindowTitle("输入产品/资源ID")
        self.resize(520, 320)
        self.setMinimumSize(420, 260)
        layout = QVBoxLayout(self)
        layout.addWidget(BodyLabel("请输入一个或多个产品/资源ID，支持换行、逗号、空格分隔：", self))
        sitewide_hint = BodyLabel(
            "留空点击确定表示查询当前供应商和对接系统下的全站数据，"
            "程序会先读取全站产品清单，再逐产品完整读取对接配置。",
            self,
        )
        sitewide_hint.setWordWrap(True)
        layout.addWidget(sitewide_hint)
        self.text_box = TextEdit(self)
        layout.addWidget(self.text_box, 1)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel,
            self,
        )
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def value(self) -> str:
        return self.text_box.toPlainText().strip()


class DockingPreviewDialog(QDialog):
    def __init__(self, rows: list[DockingPreviewRow], parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("对接配置结果预览")
        self.resize(1280, 760)
        layout = QVBoxLayout(self)
        table = TableWidget(self)
        table.setColumnCount(len(DockingResultTableModel.COLUMNS))
        table.setHorizontalHeaderLabels(list(DockingResultTableModel.COLUMNS))
        table.setRowCount(len(rows))
        for row_index, row in enumerate(rows):
            values = (
                row.status,
                str(row.source_row),
                row.product_id,
                row.product_name,
                row.resource_id,
                row.resource_name,
                row.docking_code,
                row.field_name,
                row.old_value,
                row.new_value,
                row.message,
            )
            for column_index, value in enumerate(values):
                table.setItem(row_index, column_index, QTableWidgetItem(str(value)))
        table.resizeColumnsToContents()
        table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(table, 1)
        close_btn = PushButton("关闭", self)
        close_btn.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        close_btn.clicked.connect(self.accept)
        actions = QHBoxLayout()
        actions.addStretch(1)
        actions.addWidget(close_btn)
        layout.addLayout(actions)


def _default_confirm_batch(parent: QWidget, count: int) -> bool:
    return QMessageBox.question(
        parent,
        "确认批量写入",
        f"确认按模板写入 {count} 个资源？\n\n" +
        f"写入按每批 {10000} 条提交，出现失败后立即停止。",
        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        QMessageBox.StandardButton.No,
    ) == QMessageBox.StandardButton.Yes


class DockingConfigWidget(QWidget):
    write_started = pyqtSignal(str)
    write_finished = pyqtSignal(str)
    auth_refresh_requested = pyqtSignal()
    settings_changed = pyqtSignal()
    request_failed = pyqtSignal(str, int)

    def __init__(
        self,
        transport_factory: TransportFactory,
        parent: QWidget | None = None,
        *,
        confirm_batch: ConfirmBatch | None = None,
    ) -> None:
        super().__init__(parent)
        self._transport_factory = transport_factory
        self._confirm_batch = confirm_batch or _default_confirm_batch
        self._resources: list[Any] = []
        self._template_changes: tuple[DockingTemplateChange, ...] = ()
        self._preview_rows: list[DockingPreviewRow] = []
        self._running = False
        self._writing = False
        self._shutdown = False
        self._status_base = "就绪"
        self._request_stats_text = ""
        self._events: queue.Queue[tuple[Any, ...]] = queue.Queue()
        self._stop_event = threading.Event()
        self._worker_thread: threading.Thread | None = None

        self._build_ui()
        self._timer = QTimer(self)
        self._timer.setInterval(25)
        self._timer.timeout.connect(self._drain_events)
        self._timer.start()

        self.request_interval_edit.textChanged.connect(self._on_setting_changed)
        self.parallel_requests_edit.textChanged.connect(self._on_setting_changed)
        self.vendor_edit.textChanged.connect(self._on_setting_changed)
        self.connect_no_id_edit.textChanged.connect(self._on_setting_changed)

        # compatibility with legacy naming used by some external UI states
        self.query_button = self.preload_btn
        self.import_template_button = self.load_template_button
        self.apply_template_button = self.apply_btn
        self.stop_button = self.stop_btn

    @staticmethod
    def _yes_no_text(value: Any) -> str:
        text = str(value or "").strip().lower()
        if text in {"1", "true", "是", "yes"}:
            return "是"
        if text in {"0", "false", "否", "no"}:
            return "否"
        return text or "—"

    @staticmethod
    def _resource_summary(resource: Any) -> str:
        return "\n".join(
            [
                f"产品名称：{resource.product_name}",
                f"资源名称：{resource.resource_name}",
                f"{FIELD_LABELS['vendor_resource_id']}：{resource.vendor_resource_id}",
                f"{FIELD_LABELS['ticket_code']}：{resource.ticket_code}",
                f"{FIELD_LABELS['preorder_code']}：{DockingConfigWidget._yes_no_text(resource.preorder_code)}",
                f"{FIELD_LABELS['allow_sync_sale_price']}：{DockingConfigWidget._yes_no_text(resource.allow_sync_sale_price)}",
                f"{FIELD_LABELS['allow_sync_cost_price']}：{DockingConfigWidget._yes_no_text(resource.allow_sync_cost_price)}",
                f"{FIELD_LABELS['allow_sync_sale_status']}：{DockingConfigWidget._yes_no_text(resource.allow_sync_sale_status)}",
            ]
        )

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(10)

        workflow_card = QFrame(self)
        workflow_card.setObjectName("dockingWorkflowHeader")
        workflow_card.setFrameShape(QFrame.Shape.StyledPanel)
        workflow_layout = QHBoxLayout(workflow_card)
        workflow_layout.setContentsMargins(18, 13, 18, 13)
        workflow_layout.setSpacing(20)

        workflow_copy = QVBoxLayout()
        workflow_copy.setSpacing(3)
        workflow_title = BodyLabel("对接配置批量修改", workflow_card)
        workflow_title.setObjectName("dockingWorkflowTitle")
        workflow_title.setStyleSheet("font-size: 20px; font-weight: 700;")
        workflow_subtitle = BodyLabel("用一份统一模板管理供应商对接配置，预载后可回填并预览差异", workflow_card)
        workflow_subtitle.setObjectName("dockingWorkflowSubtitle")
        workflow_subtitle.setStyleSheet("font-size: 12px;")
        workflow_steps = BodyLabel("预载数据  →  导出模板  →  回填预览  →  传出修改", workflow_card)
        workflow_steps.setObjectName("dockingWorkflowSteps")
        workflow_steps.setStyleSheet("font-size: 12px; font-weight: 600;")
        workflow_copy.addWidget(workflow_title)
        workflow_copy.addWidget(workflow_subtitle)
        workflow_copy.addWidget(workflow_steps)
        workflow_layout.addLayout(workflow_copy, 1)

        auth_copy = QVBoxLayout()
        auth_copy.setSpacing(3)
        auth_caption = BodyLabel("第一步：登录态", workflow_card)
        auth_caption.setStyleSheet("font-size: 11px;")
        auth_row = QHBoxLayout()
        auth_row.setSpacing(8)
        self.auth_button = PrimaryPushButton("刷新登录态", workflow_card)
        self.auth_button.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        self.auth_button.clicked.connect(self.auth_refresh_requested.emit)
        self.product_auth_status_label = BodyLabel("登录态：检查中…", workflow_card)
        self.product_auth_status_label.setObjectName("dockingAuthStatus")
        auth_row.addWidget(self.auth_button)
        auth_row.addWidget(self.product_auth_status_label)
        auth_copy.addWidget(auth_caption)
        auth_copy.addLayout(auth_row)
        workflow_layout.addLayout(auth_copy)

        status_copy = QVBoxLayout()
        status_copy.setSpacing(2)
        status_caption = BodyLabel("当前运行状态", workflow_card)
        status_caption.setObjectName("dockingStatusCaption")
        status_caption.setStyleSheet("font-size: 11px;")
        self.status_label = BodyLabel(self._status_base, workflow_card)
        self.status_label.setObjectName("dockingStatusLabel")
        self.status_label.setStyleSheet("font-size: 12px; color: #0d6efd;")
        status_copy.addWidget(status_caption)
        status_copy.addWidget(self.status_label)
        workflow_layout.addLayout(status_copy)

        runtime_settings = QWidget(workflow_card)
        runtime_settings.setObjectName("dockingRuntimeSettings")
        runtime_settings_layout = QVBoxLayout(runtime_settings)
        runtime_settings_layout.setContentsMargins(0, 0, 0, 0)
        runtime_settings_layout.setSpacing(3)
        runtime_caption = BodyLabel("请求设置", runtime_settings)
        runtime_caption.setObjectName("dockingRuntimeSettingsCaption")
        runtime_caption.setStyleSheet("font-size: 11px;")
        runtime_row = QHBoxLayout()
        runtime_row.setSpacing(6)
        runtime_row.addWidget(BodyLabel("间隔", runtime_settings))
        self.request_interval_edit = LineEdit(runtime_settings)
        self.request_interval_edit.setText(str(DEFAULT_REQUEST_INTERVAL_SECONDS))
        self.request_interval_edit.setObjectName("dockingRequestInterval")
        self.request_interval_edit.setFixedWidth(60)
        runtime_row.addWidget(self.request_interval_edit)
        runtime_row.addWidget(BodyLabel("秒", runtime_settings))
        runtime_row.addSpacing(4)
        runtime_row.addWidget(BodyLabel("并行数", runtime_settings))
        self.parallel_requests_edit = LineEdit(runtime_settings)
        self.parallel_requests_edit.setText(str(DEFAULT_PARALLEL_REQUESTS))
        self.parallel_requests_edit.setObjectName("dockingParallelRequests")
        self.parallel_requests_edit.setFixedWidth(52)
        runtime_row.addWidget(self.parallel_requests_edit)
        runtime_settings_layout.addWidget(runtime_caption)
        runtime_settings_layout.addLayout(runtime_row)
        workflow_layout.addWidget(runtime_settings)
        layout.addWidget(workflow_card, 0)

        template_card = QFrame(self)
        template_card.setObjectName("dockingTemplateCard")
        template_card.setFrameShape(QFrame.Shape.StyledPanel)
        template_layout = QVBoxLayout(template_card)
        template_layout.setContentsMargins(14, 10, 14, 12)
        template_layout.setSpacing(8)

        template_heading = QHBoxLayout()
        template_index = BodyLabel("01", template_card)
        template_index.setObjectName("dockingSectionIndex")
        template_index.setStyleSheet("font-size: 11px; font-weight: 700;")
        template_title = BodyLabel("准备数据与模板", template_card)
        template_title.setObjectName("dockingSectionTitle")
        template_title.setStyleSheet("font-size: 14px; font-weight: 700;")
        template_hint = BodyLabel("先获取现有资源，再导出、填写并加载统一模板", template_card)
        template_hint.setObjectName("dockingSectionHint")
        template_hint.setStyleSheet("font-size: 11px;")
        template_heading.addWidget(template_index)
        template_heading.addWidget(template_title)
        template_heading.addSpacing(6)
        template_heading.addWidget(template_hint)
        template_heading.addStretch(1)
        template_layout.addLayout(template_heading)

        actions = QHBoxLayout()
        actions.setSpacing(8)
        self.preload_btn = PrimaryPushButton("输入产品/资源ID并预载", self)
        self.preload_btn.clicked.connect(self.start_preload)
        self.preload_btn.setObjectName("dockingStartPreloadButton")
        self.export_template_btn = PushButton("导出修改模板", self)
        self.export_template_btn.setObjectName("dockingExportTemplateButton")
        self.export_template_btn.clicked.connect(
            lambda _checked=False: self.export_template()
        )
        actions.addWidget(self.preload_btn)
        actions.addWidget(self.export_template_btn)

        actions.addSpacing(12)
        actions.addWidget(BodyLabel("回填文件", self))
        self.template_path = LineEdit(self)
        self.template_path.setObjectName("dockingTemplatePath")
        self.template_path.setPlaceholderText("选择已回填的 .xlsx 修改模板")
        actions.addWidget(self.template_path, 1)
        self.choose_template_button = PushButton("选择", self)
        self.choose_template_button.setObjectName("dockingChooseTemplateButton")
        self.choose_template_button.clicked.connect(self.choose_template)
        actions.addWidget(self.choose_template_button)
        self.load_template_button = PushButton("加载回填并预览Diff", self)
        self.load_template_button.setObjectName("dockingLoadTemplateButton")
        self.load_template_button.clicked.connect(self.load_template)
        actions.addWidget(self.load_template_button)
        template_layout.addLayout(actions)

        condition_row = QHBoxLayout()
        condition_row.setSpacing(16)
        condition_row.addWidget(BodyLabel("供应商 ID", self))
        self.vendor_edit = LineEdit(self)
        self.vendor_edit.setText("")
        self.vendor_edit.setObjectName("dockingVendorEdit")
        self.vendor_edit.setFixedWidth(120)
        condition_row.addWidget(self.vendor_edit)

        condition_row.addWidget(BodyLabel("对接系统 ID", self))
        self.connect_no_id_edit = LineEdit(self)
        self.connect_no_id_edit.setText("")
        self.connect_no_id_edit.setObjectName("dockingConnectNoIdEdit")
        self.connect_no_id_edit.setFixedWidth(120)
        condition_row.addWidget(self.connect_no_id_edit)
        condition_row.addStretch(1)
        template_layout.addLayout(condition_row)
        layout.addWidget(template_card)

        execution_card = QFrame(self)
        execution_card.setObjectName("dockingExecutionCard")
        execution_card.setFrameShape(QFrame.Shape.StyledPanel)
        execution_layout = QHBoxLayout(execution_card)
        execution_layout.setContentsMargins(14, 10, 14, 10)
        execution_layout.setSpacing(8)

        execution_index = BodyLabel("02", execution_card)
        execution_index.setObjectName("dockingSectionIndex")
        execution_index.setStyleSheet("font-size: 11px; font-weight: 700;")
        execution_title = BodyLabel("校验与执行", execution_card)
        execution_title.setObjectName("dockingSectionTitle")
        execution_title.setStyleSheet("font-size: 14px; font-weight: 700;")
        execution_layout.addWidget(execution_index)
        execution_layout.addWidget(execution_title)
        execution_layout.addSpacing(8)

        self.apply_btn = PrimaryPushButton("传出修改", self)
        self.apply_btn.setObjectName("dockingApplyButton")
        self.apply_btn.clicked.connect(self.start_apply)
        self.export_report_btn = PushButton("导出结果报告", self)
        self.export_report_btn.setObjectName("dockingExportReportButton")
        self.export_report_btn.clicked.connect(self.export_report)
        self.stop_btn = PushButton("停止", self)
        self.stop_btn.setObjectName("dockingStopButton")
        self.stop_btn.clicked.connect(self.stop)
        for button in (self.apply_btn, self.export_report_btn, self.stop_btn):
            execution_layout.addWidget(button)
        execution_layout.addStretch(1)
        layout.addWidget(execution_card)

        results_card = QFrame(self)
        results_card.setObjectName("dockingResultsCard")
        results_card.setFrameShape(QFrame.Shape.StyledPanel)
        results_layout = QVBoxLayout(results_card)
        results_layout.setContentsMargins(8, 8, 8, 8)
        results_layout.setSpacing(7)

        results_heading = QHBoxLayout()
        results_index = BodyLabel("03", results_card)
        results_index.setObjectName("dockingSectionIndex")
        results_index.setStyleSheet("font-size: 11px; font-weight: 700;")
        results_title = BodyLabel("修改预览", results_card)
        results_title.setObjectName("dockingSectionTitle")
        results_title.setStyleSheet("font-size: 14px; font-weight: 700;")
        results_hint = BodyLabel("单击行查看字段差异，双击打开完整预览", results_card)
        results_hint.setObjectName("dockingSectionHint")
        results_hint.setStyleSheet("font-size: 11px;")
        results_heading.addWidget(results_index)
        results_heading.addWidget(results_title)
        results_heading.addSpacing(6)
        results_heading.addWidget(results_hint)
        results_heading.addStretch(1)
        results_layout.addLayout(results_heading)

        splitter = QSplitter(Qt.Orientation.Vertical, self)
        results_layout.addWidget(splitter, 1)
        layout.addWidget(results_card, 1)

        self.table = QTableView(self)
        self.table.setObjectName("dockingResultTable")
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table_model = DockingResultTableModel(parent=self)
        self.table.setModel(self.table_model)
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.horizontalHeader().setStretchLastSection(False)
        self.table.verticalHeader().setVisible(False)
        self.table.selectionModel().selectionChanged.connect(self._on_row_selection_changed)
        self.table.doubleClicked.connect(self.open_preview)
        widths = [80, 64, 90, 190, 90, 190, 220, 130, 220, 220, 240]
        for index, width in enumerate(widths):
            self.table.setColumnWidth(index, width)
        splitter.addWidget(self.table)

        bottom = QSplitter(Qt.Orientation.Horizontal, self)
        splitter.addWidget(bottom)
        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 2)

        log_panel = QVBoxLayout()
        log_widget = QFrame(self)
        log_widget.setFrameShape(QFrame.Shape.StyledPanel)
        log_widget.setLayout(log_panel)
        log_panel.addWidget(BodyLabel("运行日志", self))
        self.log_text = PlainTextEdit(self)
        self.log_text.setObjectName("dockingRuntimeLog")
        self.log_text.setReadOnly(True)
        log_panel.addWidget(self.log_text)
        bottom.addWidget(log_widget)

        detail_panel = QVBoxLayout()
        detail_widget = QFrame(self)
        detail_widget.setFrameShape(QFrame.Shape.StyledPanel)
        detail_widget.setLayout(detail_panel)
        detail_panel.addWidget(BodyLabel("选中项详情 / Diff", self))
        self.detail_text = PlainTextEdit(self)
        self.detail_text.setObjectName("dockingDetailText")
        self.detail_text.setReadOnly(True)
        detail_panel.addWidget(self.detail_text)
        bottom.addWidget(detail_widget)
        bottom.setStretchFactor(0, 1)
        bottom.setStretchFactor(1, 1)

        self._set_result_rows([])
        self._refresh_controls()

    def set_request_stats_text(self, text: str) -> None:
        self._request_stats_text = str(text or "").strip()
        self._render_status()

    def _set_status(self, text: str) -> None:
        self._status_base = str(text or "就绪")
        self._render_status()

    def _render_status(self) -> None:
        suffix = f" | {self._request_stats_text}" if self._request_stats_text else ""
        self.status_label.setText(f"{self._status_base}{suffix}")

    @staticmethod
    def _parse_ids(raw: str, *, allow_empty: bool = False) -> tuple[str, ...]:
        tokens = re.split(r"[\s,，;]+", raw.strip())
        seen: set[str] = set()
        result: list[str] = []
        for token in tokens:
            value = token.strip()
            if not value:
                continue
            if not value.isdigit():
                raise ValueError("请输入数字格式的产品/资源 ID")
            if value in seen:
                continue
            seen.add(value)
            result.append(value)
        if not result:
            if allow_empty:
                return ()
            raise ValueError("请输入至少一个产品/资源 ID")
        return tuple(result)

    def _normalize_runtime_settings(self) -> tuple[float, int]:
        try:
            interval = float(str(self.request_interval_edit.text()).strip())
        except Exception as error:
            raise ValueError("请求间隔必须是数字") from error
        if interval < MIN_REQUEST_INTERVAL_SECONDS or interval > MAX_REQUEST_INTERVAL_SECONDS:
            raise ValueError(f"请求间隔必须在 {MIN_REQUEST_INTERVAL_SECONDS}-{MAX_REQUEST_INTERVAL_SECONDS} 秒之间")

        try:
            parallel = int(str(self.parallel_requests_edit.text()).strip())
        except Exception as error:
            raise ValueError("并行数必须是整数") from error
        if parallel < 1 or parallel > MAX_PARALLEL_REQUESTS:
            raise ValueError(f"并行数需要在 1-{MAX_PARALLEL_REQUESTS} 之间")
        return interval, parallel

    def _parse_query_inputs(self) -> tuple[tuple[str, ...], int, int, float, int]:
        dialog = ProductIdsDialog(self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            raise ValueError("已取消预载")
        target_ids = self._parse_ids(dialog.value(), allow_empty=True)
        if not target_ids:
            target_ids = (ALL_SITE_QUERY_TARGET,)
        try:
            vendor_id = int(self.vendor_edit.text().strip())
            connect_no_id = int(self.connect_no_id_edit.text().strip())
        except ValueError as error:
            raise ValueError("供应商 ID 和对接系统 ID 必须是数字") from error
        request_interval_seconds, parallel_requests = self._normalize_runtime_settings()
        return target_ids, vendor_id, connect_no_id, request_interval_seconds, parallel_requests

    def _parse_template_inputs(self) -> tuple[str, float, int]:
        path = self.template_path.text().strip()
        if not path:
            raise ValueError("请先选择回填模板")
        request_interval_seconds, parallel_requests = self._normalize_runtime_settings()
        return path, request_interval_seconds, parallel_requests

    def _new_workflow(self) -> DockingConfigWorkflow:
        return DockingConfigWorkflow(DockingConfigApi(self._transport_factory()))

    def start_preload(self) -> None:
        if self._running or self._shutdown:
            return
        try:
            target_ids, vendor_id, connect_no_id, request_interval_seconds, parallel_requests = self._parse_query_inputs()
        except ValueError as error:
            if str(error) != "已取消预载":
                self._set_status(str(error))
            return
        self._start_worker(
            "preload",
            target_ids,
            vendor_id,
            connect_no_id,
            reset_stop=True,
            request_interval_seconds=request_interval_seconds,
            parallel_requests=parallel_requests,
        )

    def start_query(self) -> None:
        self.start_preload()

    def start_apply(self) -> None:
        if self._running or self._shutdown:
            return
        if not self._template_changes:
            self._set_status("请先导入包含有效修改的模板")
            return
        try:
            request_interval_seconds, _parallel = self._normalize_runtime_settings()
        except ValueError as error:
            self._set_status(str(error))
            return
        if not self._confirm_batch(self, len(self._template_changes)):
            self._set_status("已取消批量修改")
            return
        if not self._template_changes:
            self._set_status("请先导入包含有效修改的模板")
            return
        self._set_row_status(STATUS_PENDING)
        self._writing = True
        self._stop_event.clear()
        self.write_started.emit("docking_config")
        if self._stop_event.is_set():
            self._writing = False
            self._set_status("另一个业务组件正在写入，本次批量任务未启动")
            self._refresh_controls()
            return
        self._start_worker(
            "apply",
            self._template_changes,
            reset_stop=False,
            request_interval_seconds=request_interval_seconds,
            parallel_requests=1,
        )

    def start_apply_template(self) -> None:
        self.start_apply()

    def export_template(self, path: str | None = None) -> str | None:
        # QPushButton.clicked emits a checked bool. Keep this public helper safe
        # even if it is connected directly by another host or an older adapter.
        if isinstance(path, bool):
            path = None
        if self._running:
            self._set_status("任务运行中，暂不可导出模板")
            return None
        if not self._resources:
            self._set_status("请先预载资源")
            return None

        if path is None:
            default_path = Path.home() / "Downloads" / f"产品对接配置_{self._resources[0].product_id}.xlsx"
            path, _selected = QFileDialog.getSaveFileName(
                self,
                "导出产品对接配置模板",
                str(default_path),
                "Excel 工作簿 (*.xlsx)",
                options=excel_file_dialog_options(),
            )
        if not path:
            return None

        if not str(path).lower().endswith(".xlsx"):
            path = f"{path}.xlsx"
        try:
            exported_path = export_docking_template(path, self._resources)
        except Exception as error:
            self._append_log(f"模板导出失败：{error}")
            self._set_status(f"模板导出失败：{error}")
            return None
        self._append_log(f"已导出产品对接配置模板：{exported_path}")
        self._set_status(f"模板导出成功：{exported_path}")
        return str(exported_path)

    def choose_template(self) -> None:
        path, _selected = QFileDialog.getOpenFileName(
            self,
            "导入产品对接配置模板",
            str(Path.home() / "Downloads"),
            "Excel 工作簿 (*.xlsx)",
            options=excel_file_dialog_options(),
        )
        if not path:
            self._set_status("已取消选择模板")
            return
        self.template_path.setText(path)
        self.load_template()

    def load_template(self) -> tuple[DockingTemplateChange, ...]:
        if self._running:
            return ()
        try:
            path = self._parse_template_inputs()[0]
        except ValueError as error:
            self._set_status(str(error))
            return ()
        return self.import_template(path)

    def import_template(self, path: str | None = None) -> tuple[DockingTemplateChange, ...]:
        if self._running:
            return ()
        if path is None:
            if not self.template_path.text().strip():
                self._set_status("请先选择模板文件")
                return ()
            path = self.template_path.text().strip()
        try:
            changes = import_docking_template(path)
        except Exception as error:
            self._set_status(f"模板导入失败：{error}")
            self._append_log(f"模板导入失败：{error}")
            return ()
        self.template_path.setText(path)
        self._template_changes = changes
        self._preview_rows = self._rows_from_changes(changes)
        self._set_result_rows(self._preview_rows)
        self._append_log(f"已导入模板：{path}；有效修改 {len(changes)} 条")
        self._set_status(
            "模板中没有与当前值不同的目标修改" if not changes else f"模板导入完成：{len(changes)} 条有效修改"
        )
        self._refresh_controls()
        return changes

    def _start_worker(
        self,
        operation: str,
        *args: Any,
        reset_stop: bool = True,
        request_interval_seconds: float = 0.0,
        parallel_requests: int = 1,
    ) -> None:
        if reset_stop:
            self._stop_event.clear()
        if self._shutdown:
            return

        self._running = True
        is_all_site_preload = (
            operation == "preload"
            and bool(args)
            and tuple(args[0]) == (ALL_SITE_QUERY_TARGET,)
        )
        if is_all_site_preload:
            self._set_status("正在读取全站产品清单并逐产品预载...")
        else:
            self._set_status("正在预载..." if operation == "preload" else "正在传出修改...")
        self._refresh_controls()
        self._append_log("开始执行全站分页预载 ..." if is_all_site_preload else f"开始执行 {operation} ...")

        def run() -> None:
            try:
                workflow = self._new_workflow()
                if operation == "preload":
                    target_ids, vendor_id, connect_no_id = args
                    result = workflow.query_many(
                        target_ids,
                        vendor_id,
                        connect_no_id,
                        self._stop_event,
                        parallel_requests=parallel_requests,
                        request_interval_seconds=request_interval_seconds,
                        progress=self.enqueue_log,
                    )
                    self._events.put(("success", operation, result))
                else:
                    result = workflow.apply_template_changes(
                        args[0],
                        self._stop_event,
                        request_interval_seconds=request_interval_seconds,
                    )
                    self._events.put(("success", operation, result))
            except Exception as error:
                self._events.put(("failed", operation, str(error)))

        self._worker_thread = threading.Thread(
            target=run,
            daemon=True,
            name=f"docking-config-{operation}",
        )
        self._worker_thread.start()

    def _drain_events(self) -> None:
        while True:
            try:
                event = self._events.get_nowait()
            except queue.Empty:
                return

            kind, operation, payload = event
            if kind == "log":
                self._append_log(str(payload))
                continue

            self._running = False
            if kind == "failed":
                self._set_status(f"操作失败：{payload}")
                self._append_log(f"{operation} 失败：{payload}")
                self.request_failed.emit("docking", 1)
            elif operation == "preload":
                self._show_preload_result(payload)
            else:
                self._show_batch_result(payload)

            if self._writing:
                self._writing = False
                self.write_finished.emit("docking_config")
            self._refresh_controls()

    def _show_preload_result(self, result: Any) -> None:
        resources = list(getattr(result, "resources", ()))
        self._resources = resources
        self._template_changes = ()
        self._preview_rows = self._rows_from_resources(resources)
        self._set_result_rows(self._preview_rows)
        self._append_log(f"已预载 {len(resources)} 个对接资源")
        if not resources:
            self._set_status("预载完成：当前无结果")
            return
        self._set_status(f"预载完成：返回 {len(resources)} 条资源，当前页大小 {result.page_size}")

    def _show_batch_result(self, result: DockingBatchResult) -> None:
        self._preview_rows = self._update_rows_for_batch(self._preview_rows, result)
        self._set_result_rows(self._preview_rows)
        self._set_status(
            f"模板写入完成：成功 {result.succeeded}，失败 {result.failed}，未提交 {result.unsubmitted}"
        )
        self._append_log(
            f"模板结果：成功 {result.succeeded}，失败 {result.failed}，未提交 {result.unsubmitted}"
        )
        if result.failure is not None:
            failure = result.failure
            self._append_log(
                f"第 {failure.batch_number} 批失败：提交序号 "
                f"{failure.item_start}-{failure.item_end}，模板行 "
                f"{failure.source_row_start}-{failure.source_row_end}；{failure.message}"
            )
            self.request_failed.emit("docking", result.failed + result.unsubmitted)
        elif result.failed or result.unsubmitted:
            self.request_failed.emit("docking", result.failed + result.unsubmitted)
        if result.failed == 0 and result.unsubmitted == 0:
            self._template_changes = ()

    def _rows_from_resources(self, resources: list[Any]) -> list[DockingPreviewRow]:
        rows: list[DockingPreviewRow] = []
        for resource in resources:
            rows.append(
                DockingPreviewRow(
                    row_kind="resource",
                    source_row=0,
                    product_id=str(resource.product_id),
                    product_name=str(resource.product_name),
                    resource_id=str(resource.resource_id),
                    resource_name=str(resource.resource_name),
                    field_name="资源信息",
                    old_value=self._resource_summary(resource),
                    new_value="",
                    status=STATUS_READY,
                    message="预载到的现有资源",
                    docking_code=str(resource.current_docking_code),
                )
            )
        if not rows:
            rows.append(
                DockingPreviewRow(
                    row_kind="resource",
                    source_row=0,
                    product_id="",
                    product_name="",
                    resource_id="",
                    resource_name="",
                    field_name="整体",
                    old_value="",
                    new_value="",
                    status=STATUS_NOCHANGE,
                    message="无资源结果",
                )
            )
        return rows

    def _rows_from_changes(self, changes: tuple[DockingTemplateChange, ...]) -> list[DockingPreviewRow]:
        rows: list[DockingPreviewRow] = []
        for change in changes:
            old_values = {
                "vendor_resource_id": str(change.baseline.vendor_resource_id),
                "ticket_code": str(change.baseline.ticket_code),
                "preorder_code": self._yes_no_text(change.baseline.preorder_code),
                "allow_sync_sale_price": self._yes_no_text(change.baseline.allow_sync_sale_price),
                "allow_sync_cost_price": self._yes_no_text(change.baseline.allow_sync_cost_price),
                "allow_sync_sale_status": self._yes_no_text(change.baseline.allow_sync_sale_status),
            }
            new_values = {
                "vendor_resource_id": str(change.changes.get("vendor_resource_id")),
                "ticket_code": str(change.changes.get("ticket_code")),
                "preorder_code": self._yes_no_text(change.changes.get("preorder_code")),
                "allow_sync_sale_price": self._yes_no_text(change.changes.get("allow_sync_sale_price")),
                "allow_sync_cost_price": self._yes_no_text(change.changes.get("allow_sync_cost_price")),
                "allow_sync_sale_status": self._yes_no_text(change.changes.get("allow_sync_sale_status")),
            }
            changed_rows = []
            for key, old_label in FIELD_LABELS.items():
                old_value = old_values[key]
                new_value = new_values[key]
                if str(old_value) == str(new_value):
                    continue
                changed_rows.append(
                    DockingPreviewRow(
                        row_kind="field",
                        source_row=change.source_row,
                        product_id=str(change.baseline.product_id),
                        product_name=str(change.baseline.product_name),
                        resource_id=str(change.baseline.resource_id),
                        resource_name=str(change.baseline.resource_name),
                        field_name=old_label,
                        old_value=old_value,
                        new_value=new_value,
                        status=STATUS_PENDING,
                        message="待提交",
                        docking_code=str(change.baseline.current_docking_code),
                    )
                )
            rows.extend(changed_rows)
        if not rows:
            rows.append(
                DockingPreviewRow(
                    row_kind="field",
                    source_row=0,
                    product_id="",
                    product_name="",
                    resource_id="",
                    resource_name="",
                    field_name="整体",
                    old_value="",
                    new_value="",
                    status=STATUS_NOCHANGE,
                    message="模板中无差异字段",
                )
            )
        return rows

    def _set_row_status(self, status: str) -> None:
        self._preview_rows = [
            replace(row, status=status) if row.row_kind == "field" and row.status == STATUS_PENDING else row
            for row in self._preview_rows
        ]
        self._set_result_rows(self._preview_rows)

    def _set_result_rows(self, rows: list[DockingPreviewRow]) -> None:
        self._preview_rows = list(rows)
        self.table_model.set_rows(self._preview_rows)

    def _update_rows_for_batch(self, rows: list[DockingPreviewRow], result: DockingBatchResult) -> list[DockingPreviewRow]:
        if not rows:
            return []
        position_by_source_row = {
            change.source_row: index
            for index, change in enumerate(self._template_changes, start=1)
        }
        updated: list[DockingPreviewRow] = []
        for row in rows:
            if row.row_kind != "field":
                updated.append(row)
                continue
            if row.status == STATUS_NOCHANGE:
                updated.append(row)
                continue
            position = position_by_source_row.get(row.source_row)
            if position is None:
                updated.append(row)
                continue
            if position <= result.succeeded:
                updated.append(replace(row, status=STATUS_SUCCESS, message="写入成功"))
            elif position <= result.succeeded + result.failed:
                failure = result.failure
                updated.append(
                    replace(
                        row,
                        status=STATUS_FAILED,
                        message=(
                            f"第 {failure.batch_number} 批失败：{failure.message}"
                            if failure is not None
                            else "写入失败"
                        ),
                    )
                )
            else:
                updated.append(replace(row, status=STATUS_UNLOADED, message="尚未提交"))
        return updated

    def _on_row_selection_changed(self, _selected, _deselected) -> None:  # pragma: no cover - UI callback
        model = self.table.selectionModel()
        indexes = model.selectedRows() if model is not None else []
        if not indexes:
            self.detail_text.clear()
            return
        row = self.table_model.row_at(indexes[0].row())
        if row is None:
            return
        self._render_row_detail(row)

    def _render_row_detail(self, row: DockingPreviewRow) -> None:
        lines = [
            f"状态: {row.status}",
            f"来源行: {row.source_row}",
            f"产品: {row.product_id} - {row.product_name}",
            f"资源: {row.resource_id} - {row.resource_name}",
            f"当前对接代码: {row.docking_code or '（无）'}",
            f"字段: {row.field_name}",
            "",
            "原值:",
            row.old_value or "（无）",
            "",
            "目标值:",
            row.new_value or "（无）",
            "",
            "说明:",
            row.message or "（无）",
        ]
        self.detail_text.setPlainText("\n".join(lines))

    def open_preview(self, _index: Any = None) -> None:
        selected = self.table.selectionModel().selectedRows() if self.table.selectionModel() is not None else []
        rows: list[DockingPreviewRow]
        if selected:
            rows = [self.table_model.row_at(selected[0].row())] if self.table_model.row_at(selected[0].row()) else []
        else:
            rows = self._preview_rows
        dialog = DockingPreviewDialog([row for row in rows if row is not None], self)
        dialog.exec()

    def export_report(self) -> None:
        if self._running:
            self._set_status("任务运行中，暂不可导出报告")
            return
        if not self._preview_rows:
            self._set_status("无可导出的报告内容")
            return

        default_path = Path.home() / "Downloads" / "对接配置批量修改报告"
        path, selected_filter = QFileDialog.getSaveFileName(
            self,
            "导出结果报告",
            str(default_path),
            "Excel 工作簿 (*.xlsx);;CSV 文件 (*.csv)",
            options=excel_file_dialog_options(),
        )
        if not path:
            return

        try:
            is_csv = str(path).lower().endswith(".csv") or str(selected_filter).startswith("CSV")
            if is_csv:
                if not str(path).lower().endswith(".csv"):
                    path = f"{path}.csv"
                with open(path, "w", encoding="utf-8", newline="") as csv_file:
                    writer = csv.writer(csv_file)
                    writer.writerow(DockingResultTableModel.COLUMNS)
                    for row in self._preview_rows:
                        writer.writerow(row.as_report_row())
            else:
                if not str(path).lower().endswith(".xlsx"):
                    path = f"{path}.xlsx"
                from openpyxl import Workbook
                from openpyxl.utils import get_column_letter

                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "结果报告"
                sheet.append(list(DockingResultTableModel.COLUMNS))
                for row in self._preview_rows:
                    sheet.append(list(row.as_report_row()))
                for column in range(1, len(DockingResultTableModel.COLUMNS) + 1):
                    sheet.column_dimensions[get_column_letter(column)].width = 24 if column != 10 else 52
                workbook.save(path)
        except Exception as error:
            self._set_status(f"导出结果报告失败：{error}")
            self._append_log(f"导出结果报告失败：{error}")
            return
        self._set_status(f"结果报告已导出：{path}")
        self._append_log(f"已导出结果报告：{path}")

    def _append_log(self, message: str) -> None:
        self.log_text.appendPlainText(message)

    def enqueue_log(self, message: str) -> None:
        self._events.put(("log", "", message))

    def _refresh_controls(self) -> None:
        idle = not self._running and not self._shutdown
        self.preload_btn.setEnabled(idle)
        self.export_template_btn.setEnabled(idle and bool(self._resources))
        self.apply_btn.setEnabled(idle and bool(self._template_changes))
        self.export_report_btn.setEnabled(bool(self._preview_rows))
        self.stop_btn.setEnabled(self._running)
        self.choose_template_button.setEnabled(idle)
        self.load_template_button.setEnabled(idle)
        self.template_path.setEnabled(idle)
        self.vendor_edit.setEnabled(idle)
        self.connect_no_id_edit.setEnabled(idle)
        self.request_interval_edit.setEnabled(idle)
        self.parallel_requests_edit.setEnabled(idle)

    def _on_setting_changed(self) -> None:
        if self._running:
            return
        self.settings_changed.emit()

    def stop(self) -> None:
        self._stop_event.set()
        if self._running:
            self._set_status("正在停止…")

    def is_running(self) -> bool:
        return self._running

    def reset(self) -> None:
        if self._running:
            return
        self.template_path.clear()
        self.vendor_edit.setText("")
        self.connect_no_id_edit.setText("")
        self.request_interval_edit.setText(str(DEFAULT_REQUEST_INTERVAL_SECONDS))
        self.parallel_requests_edit.setText(str(DEFAULT_PARALLEL_REQUESTS))
        self._resources = []
        self._template_changes = ()
        self._preview_rows = []
        self._set_result_rows([])
        self._set_status("就绪")
        self.log_text.clear()
        self.detail_text.clear()
        self._refresh_controls()

    def export_settings(self) -> dict[str, str]:
        settings: dict[str, str] = {}
        if self.template_path.text().strip():
            settings["template_path"] = self.template_path.text().strip()
        if self.vendor_edit.text().strip():
            settings["vendor_id"] = self.vendor_edit.text().strip()
        if self.connect_no_id_edit.text().strip():
            settings["connect_no_id"] = self.connect_no_id_edit.text().strip()
        settings["request_interval_seconds"] = str(float(self.request_interval_edit.text() or DEFAULT_REQUEST_INTERVAL_SECONDS))
        settings["parallel_requests"] = str(int(self.parallel_requests_edit.text() or DEFAULT_PARALLEL_REQUESTS))
        return settings

    def current_request_interval_seconds(self) -> float:
        try:
            value = float(self.request_interval_edit.text().strip())
        except (TypeError, ValueError):
            return DEFAULT_REQUEST_INTERVAL_SECONDS
        return min(max(value, MIN_REQUEST_INTERVAL_SECONDS), MAX_REQUEST_INTERVAL_SECONDS)

    def request_close(self, parent: QWidget) -> bool:
        if not self._running:
            return True
        QMessageBox.information(parent, "任务运行中", "请先停止产品对接配置任务。")
        return False

    def shutdown(self) -> None:
        self._shutdown = True
        self.stop()
        self._timer.stop()
        self._refresh_controls()

    def _refresh_result_panel(self) -> None:
        self._set_result_rows(self._preview_rows)

    def clear(self) -> None:
        self.reset()
