"""Excel template round-trip for docking configuration changes."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .models import DockingResource


SHEET_NAME = "对接配置"
GUIDE_SHEET_NAME = "填写说明"
TEMPLATE_VERSION = "1"
HEADERS = (
    "产品ID",
    "产品名称",
    "资源ID",
    "资源名称",
    "供应商ID",
    "对接系统ID",
    "当前供应商资源ID",
    "目标供应商资源ID",
    "当前票种代码",
    "目标票种代码",
    "当前支付前对接",
    "目标支付前对接",
    "当前同步卖价",
    "目标同步卖价",
    "当前同步底价",
    "目标同步底价",
    "当前同步开关班",
    "目标同步开关班",
)
TARGET_HEADERS = {
    "目标供应商资源ID",
    "目标票种代码",
    "目标支付前对接",
    "目标同步卖价",
    "目标同步底价",
    "目标同步开关班",
}


@dataclass(frozen=True)
class DockingTemplateChange:
    baseline: DockingResource
    changes: dict[str, Any]
    source_row: int


def _yes_no(value: int) -> str:
    return "是" if int(value) == 1 else "否"


def _normalize_yes_no(value: Any, field_name: str, row_number: int) -> str:
    text = str(value or "").strip()
    mapping = {"是": "1", "否": "0", "1": "1", "0": "0"}
    if text not in mapping:
        raise ValueError(f"第 {row_number} 行“{field_name}”只能填写是或否")
    return mapping[text]


def export_docking_template(path: str | Path, resources: Iterable[DockingResource]) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    rows = tuple(resources)
    if not rows:
        raise ValueError("没有可导出的产品对接资源")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{len(rows) + 1}"
    sheet.append(list(HEADERS))
    for resource in rows:
        sheet.append([
            str(resource.product_id),
            resource.product_name,
            str(resource.resource_id),
            resource.resource_name,
            str(resource.vendor_id),
            str(resource.connect_no_id),
            resource.vendor_resource_id,
            "",
            resource.ticket_code,
            "",
            resource.preorder_label,
            "",
            _yes_no(resource.allow_sync_sale_price),
            "",
            _yes_no(resource.allow_sync_cost_price),
            "",
            _yes_no(resource.allow_sync_sale_status),
            "",
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    target_fill = PatternFill("solid", fgColor="FFF2CC")
    current_fill = PatternFill("solid", fgColor="EAF2F8")
    change_fill = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="D9E2F3")
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, column)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color="17365D"))
        fill = target_fill if header in TARGET_HEADERS else current_fill
        for row in range(2, len(rows) + 2):
            data_cell = sheet.cell(row, column)
            data_cell.fill = fill
            data_cell.border = Border(bottom=thin)
            data_cell.alignment = Alignment(vertical="center")
            if header.endswith("ID") or header in {"产品ID", "资源ID", "供应商ID", "对接系统ID"}:
                data_cell.number_format = "@"

    for column in (8, 10, 12, 14, 16, 18):
        letter = sheet.cell(1, column).column_letter
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}{len(rows) + 1}",
            FormulaRule(formula=[f'LEN({letter}2)>0'], fill=change_fill),
        )

    yes_no_validation = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    yes_no_validation.error = "请选择“是”或“否”"
    yes_no_validation.errorTitle = "无效选项"
    sheet.add_data_validation(yes_no_validation)
    for column in (12, 14, 16, 18):
        letter = sheet.cell(1, column).column_letter
        yes_no_validation.add(f"{letter}2:{letter}{len(rows) + 1}")

    widths = {
        "A": 14, "B": 34, "C": 14, "D": 34, "E": 13, "F": 14,
        "G": 24, "H": 24, "I": 24, "J": 24, "K": 16, "L": 16,
        "M": 14, "N": 14, "O": 14, "P": 14, "Q": 16, "R": 16,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 28

    guide = workbook.create_sheet(GUIDE_SHEET_NAME)
    guide.sheet_view.showGridLines = False
    guide["A1"] = "产品对接单独配置 · 模板填写说明"
    guide["A1"].font = Font(size=16, bold=True, color="1F4E78")
    guide["A3"] = "模板版本"
    guide["B3"] = TEMPLATE_VERSION
    guide["A5"] = "1. 只填写黄色的“目标”列；蓝色当前值用于判断变化并保留未修改字段。"
    guide["A6"] = "2. 目标列留空表示该字段保持不变；目标值与当前值相同的行不会产生写入。"
    guide["A7"] = "3. 支付前对接和同步开关列只能选择“是”或“否”。"
    guide["A8"] = "4. 导入后程序只提取发生变化的行，并按每批最多 10000 条提交。"
    guide["A9"] = "5. 不要修改产品 ID、资源 ID、供应商 ID、对接系统 ID 或当前值列。"
    guide.column_dimensions["A"].width = 104
    guide.column_dimensions["B"].width = 18
    for row in range(5, 10):
        guide.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row].height = 30

    workbook.save(target)
    return target


def import_docking_template(path: str | Path) -> tuple[DockingTemplateChange, ...]:
    source = Path(path)
    workbook = load_workbook(source, data_only=True, read_only=False)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"模板缺少工作表“{SHEET_NAME}”")
    sheet = workbook[SHEET_NAME]
    actual_headers = tuple(str(sheet.cell(1, column).value or "").strip() for column in range(1, len(HEADERS) + 1))
    if actual_headers != HEADERS:
        raise ValueError("模板列结构不匹配，请重新从程序导出模板")

    seen_keys: set[tuple[int, int, int, int]] = set()
    changes: list[DockingTemplateChange] = []
    for row_number in range(2, sheet.max_row + 1):
        values = {
            header: sheet.cell(row_number, column).value
            for column, header in enumerate(HEADERS, start=1)
        }
        if not any(value not in (None, "") for value in values.values()):
            continue
        try:
            product_id = int(str(values["产品ID"] or "").strip())
            resource_id = int(str(values["资源ID"] or "").strip())
            vendor_id = int(str(values["供应商ID"] or "").strip())
            connect_no_id = int(str(values["对接系统ID"] or "").strip())
        except ValueError as error:
            raise ValueError(f"第 {row_number} 行产品、资源、供应商和对接系统 ID 必须是数字") from error
        resource_key = (product_id, resource_id, vendor_id, connect_no_id)
        if resource_key in seen_keys:
            raise ValueError(
                f"第 {row_number} 行重复的唯一键 "
                f"(产品ID={product_id}、资源ID={resource_id}、供应商ID={vendor_id}、对接系统ID={connect_no_id})"
            )
        seen_keys.add(resource_key)

        current_preorder = _normalize_yes_no(values["当前支付前对接"], "当前支付前对接", row_number)
        current_sale = _normalize_yes_no(values["当前同步卖价"], "当前同步卖价", row_number)
        current_cost = _normalize_yes_no(values["当前同步底价"], "当前同步底价", row_number)
        current_status = _normalize_yes_no(values["当前同步开关班"], "当前同步开关班", row_number)
        baseline = DockingResource(
            product_id=product_id,
            product_name=str(values["产品名称"] or ""),
            resource_id=resource_id,
            resource_name=str(values["资源名称"] or ""),
            vendor_id=vendor_id,
            connect_no_id=connect_no_id,
            vendor_resource_id=str(values["当前供应商资源ID"] or ""),
            allow_sync_sale_price=1 if current_sale == "1" else 2,
            allow_sync_cost_price=1 if current_cost == "1" else 2,
            allow_sync_stock=0,
            allow_sync_sale_status=1 if current_status == "1" else 2,
            connection_properties=(
                {"key": "4", "value": str(values["当前票种代码"] or "")},
                {"key": "203", "value": current_preorder},
            ),
        )

        raw_targets = (
            values["目标供应商资源ID"], values["目标票种代码"], values["目标支付前对接"],
            values["目标同步卖价"], values["目标同步底价"], values["目标同步开关班"],
        )
        if not any(str(value or "").strip() for value in raw_targets):
            continue
        target_vendor = str(values["目标供应商资源ID"] or "").strip() or baseline.vendor_resource_id
        target_ticket = str(values["目标票种代码"] or "").strip() or baseline.ticket_code
        target_preorder = (
            _normalize_yes_no(values["目标支付前对接"], "目标支付前对接", row_number)
            if str(values["目标支付前对接"] or "").strip()
            else baseline.preorder_code
        )
        target_sale = (
            _normalize_yes_no(values["目标同步卖价"], "目标同步卖价", row_number)
            if str(values["目标同步卖价"] or "").strip()
            else "1" if baseline.allow_sync_sale_price == 1 else "0"
        )
        target_cost = (
            _normalize_yes_no(values["目标同步底价"], "目标同步底价", row_number)
            if str(values["目标同步底价"] or "").strip()
            else "1" if baseline.allow_sync_cost_price == 1 else "0"
        )
        target_status = (
            _normalize_yes_no(values["目标同步开关班"], "目标同步开关班", row_number)
            if str(values["目标同步开关班"] or "").strip()
            else "1" if baseline.allow_sync_sale_status == 1 else "0"
        )
        target = {
            "vendor_resource_id": target_vendor,
            "ticket_code": target_ticket,
            "preorder_code": target_preorder,
            "allow_sync_sale_price": 1 if target_sale == "1" else 2,
            "allow_sync_cost_price": 1 if target_cost == "1" else 2,
            "allow_sync_sale_status": 1 if target_status == "1" else 2,
        }
        unchanged = (
            target_vendor == baseline.vendor_resource_id
            and target_ticket == baseline.ticket_code
            and target_preorder == baseline.preorder_code
            and target["allow_sync_sale_price"] == baseline.allow_sync_sale_price
            and target["allow_sync_cost_price"] == baseline.allow_sync_cost_price
            and target["allow_sync_sale_status"] == baseline.allow_sync_sale_status
        )
        if not unchanged:
            changes.append(DockingTemplateChange(baseline, target, row_number))
    return tuple(changes)
