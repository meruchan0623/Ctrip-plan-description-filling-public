# -*- coding: utf-8 -*-
"""Trip_detail_filling translation rules for interface-based writes."""

from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

ALL_TEMPLATE_LANGUAGES = sorted([
    "ar", "de", "en", "es", "fr", "id", "it", "ja", "ko", "ms", "nl", "ru", "th",
    "vi", "zh-hk", "zh-tw", "zh-sgmy", "fi", "uk", "da", "sv", "el", "pl", "tr", "pt",
])

DEFAULT_SELECTED_LANGUAGES = [
    "ar", "de", "en", "es", "fr", "id", "it", "ja", "ko", "ms", "nl", "ru", "th",
    "vi", "zh-hk", "zh-tw", "zh-sgmy", "fi", "uk", "da", "sv", "el", "pl", "tr", "pt",
]

DEFAULT_SELECTED_MODULES = [
    "product_name", "mm_detail", "fee_include", "fee_exclude",
    "usage", "description", "notice", "post_booking_reminder", "operation_attention", "sales_property",
]

INTERFACE_LOCALES_BY_LANG_GROUP = {
    "ar": ["ar-XX"],
    "zh-sgmy": ["zh-SG", "zh-MY"],
}

RTL_LANG_GROUPS = {"ar"}

LANGUAGE_LABELS = {
    "ar": "阿拉伯语",
    "de": "德语", "en": "英语", "es": "西班牙语", "fr": "法语", "id": "印尼语",
    "it": "意大利语", "ja": "日语", "ko": "韩语", "ms": "马来语", "nl": "荷兰语",
    "ru": "俄语", "th": "泰语", "vi": "越南语", "zh-hk": "繁中(香港)",
    "zh-tw": "繁中(台湾)", "zh-sgmy": "中文(新马)", "fi": "芬兰语",
    "uk": "乌克兰语", "da": "丹麦语", "sv": "瑞典语", "el": "希腊语",
    "pl": "波兰语", "tr": "土耳其语", "pt": "葡萄牙语",
}

MODULE_LABELS = {
    "product_name": "产品名称",
    "mm_detail": "商品图文详情",
    "fee_include": "费用包含",
    "fee_exclude": "费用不包含",
    "usage": "使用方法",
    "description": "产品说明",
    "notice": "注意事项",
    "post_booking_reminder": "订后出行提醒",
    "operation_attention": "携程处理注意事项",
    "sales_property": "销售属性",
    "package_detail": "套餐规格详情",
    "package_quantity": "套餐规格量词",
}

SPECIAL_RULE_LABELS = {
    "force_365_total_stop": "365天总量强制后停用模板",
    "daily_100gb_template": "每日100GB使用专用模板",
    "daily_100gb_unlimited_template": "每日100GB不限速使用专用模板",
    "unlimited_daily_20gb_template": "吃到饱每日20GB使用专用模板",
    "unlimited_capacity_whitelist_skip": "吃到饱仅填每日100GB/20GB",
    "mixed_daily_20gb_template": "总量停用每日降速中20GB使用吃到饱模板",
}

SPECIAL_RULE_DESCRIPTIONS = {
    "force_365_total_stop": "365天 + 总量/總量 时强制使用后停用 total 模板",
    "daily_100gb_template": "每日套餐容量包含100GB时优先使用 daily_100gb",
    "daily_100gb_unlimited_template": "每日100GB-不限速时优先使用 daily_100gb_unlimited",
    "unlimited_daily_20gb_template": "吃到饱_提前安装中每日20GB使用 daily_20gb",
    "unlimited_capacity_whitelist_skip": "吃到饱_提前安装中非每日100GB/20GB静默跳过",
    "mixed_daily_20gb_template": "总量停用每日降速中每日20GB使用 daily_20gb",
}

DETAIL_SHEET_NAME = "详情模板"
DETAIL_REQUIRED_COLUMNS = {"语言", "使用方法", "产品说明", "注意事项", "订后出行提醒"}
DETAIL_COLUMNS = [
    ("使用方法", "usage"),
    ("产品说明", "description"),
    ("注意事项", "notice"),
    ("订后出行提醒", "post_booking_reminder"),
    ("携程处理注意事项", "operation_attention"),
]
DETAIL_RICH_KEYS = {"description", "notice", "post_booking_reminder"}

PACKAGE_COMBINATION_SHEET = "组合模板"
PACKAGE_CAPACITY_SHEET = "容量模板"
PACKAGE_INSTALL_SHEET = "安装说明"
PACKAGE_DAY_CALC_SHEET = "天数计算"
PACKAGE_RESET_CALC_SHEET = "重置计时"
PACKAGE_RESET_OPTIONS_SHEET = "重置选项"

QUANTITY_WORDS = {
    "ar": "الكمية", "de": "Anzahl der Exemplare", "en": "Quantity", "es": "Cantidad",
    "fr": "Nombre d'exemplaires", "id": "Jumlah salinan", "it": "Numero di copie",
    "ja": "数量", "ko": "수량", "ms": "Bilangan salinan", "nl": "Aantal exemplaren",
    "th": "จำนวน", "vi": "Số lượng bản sao", "zh-hk": "份數", "zh-tw": "份數",
    "zh-sgmy": "份数", "ru": "Количество", "fi": "Määrä", "uk": "Кількість",
    "da": "Antal", "sv": "Antal", "el": "Ποσότητα", "pl": "Ilość",
    "tr": "Miktar", "pt": "Quantidade",
}

FIXED_CONTENT = {
    "ar-XX": {"fee_include": "eSIM + بيانات\nمشاركة نقطة الاتصال", "fee_exclude": "مكالمات صوتية\nرسائل SMS", "sales_property": ["طريقة الاستلام", "الأيام", "تفاصيل الباقة"]},
    "zh-HK": {"fee_include": "eSIM\n流動數據\n支援熱點共享", "fee_exclude": "語音通話\n短訊", "sales_property": ["取卡方式", "套餐天數", "套餐內容"]},
    "zh-TW": {"fee_include": "eSIM\n行動數據\n支援熱點分享", "fee_exclude": "通話\n簡訊", "sales_property": ["取卡方式", "套餐天數", "套餐內容"]},
    "zh-SGMY": {"fee_include": "eSIM\n移动数据\n支持热点共享", "fee_exclude": "语音通话\n短信", "sales_property": ["取卡方式", "套餐天数", "套餐内容"]},
    "en-US": {"fee_include": "eSIM + Data\nHotspot Sharing", "fee_exclude": "Voice\nSMS", "sales_property": ["Collection Method", "Days", "Plan Details"]},
    "ja-JP": {"fee_include": "eSIM + データ通信\nテザリング対応", "fee_exclude": "音声通話\nSMS", "sales_property": ["受取方法", "日数", "プラン内容"]},
    "ko-KR": {"fee_include": "eSIM + 데이터 \n핫스팟 공유 지원", "fee_exclude": "음성 통화\nSMS", "sales_property": ["수령 방법", "일수", "요금제 정보"]},
    "th-TH": {"fee_include": "eSIM + ดาต้า\nรองรับการแชร์ฮอตสปอต", "fee_exclude": "โทร\nSMS", "sales_property": ["วิธีรับ", "จำนวนวัน", "รายละเอียดแพ็กเกจ"]},
    "de-DE": {"fee_include": "eSIM + Daten\nHotspot-Freigabe", "fee_exclude": "Anrufe\nSMS", "sales_property": ["Empfangsmethode", "Tage", "Tarifdetails"]},
    "es-ES": {"fee_include": "eSIM + Datos\nCompartir hotspot", "fee_exclude": "Llamadas\nSMS", "sales_property": ["Método de recepción", "Días", "Detalles del plan"]},
    "fr-FR": {"fee_include": "eSIM + Données\nPartage de connexion", "fee_exclude": "Appels\nSMS", "sales_property": ["Mode de réception", "Jours", "Détails du forfait"]},
    "id-ID": {"fee_include": "eSIM + Data\nBerbagi Hotspot", "fee_exclude": "Panggilan\nSMS", "sales_property": ["Metode Penerimaan", "Hari", "Detail Paket"]},
    "ms-MY": {"fee_include": "eSIM + Data\nHotspot Sharing", "fee_exclude": "Panggilan\nSMS", "sales_property": ["Kaedah Penerimaan", "Hari", "Maklumat Pelan"]},
    "vi-VN": {"fee_include": "eSIM + Dữ liệu\nChia sẻ điểm phát sóng", "fee_exclude": "Gọi thoại\nSMS", "sales_property": ["Phương thức nhận", "Ngày", "Thông tin gói"]},
    "it-IT": {"fee_include": "eSIM + Dati\nHotspot personale", "fee_exclude": "Chiamate\nSMS", "sales_property": ["Metodo di ricezione", "Giorni", "Dettagli del piano"]},
    "nl-NL": {"fee_include": "eSIM + Gegevens\nHotspot delen", "fee_exclude": "Bellen\nSMS", "sales_property": ["Ontvangstmethode", "Dagen", "Pakketdetails"]},
    "ru-RU": {"fee_include": "eSIM + Данные\nРежим модема", "fee_exclude": "Звонки\nSMS", "sales_property": ["Способ получения", "Дни", "Информация о тарифе"]},
    "fi-FI": {"fee_include": "eSIM + Data\nHotspot Sharing", "fee_exclude": "Puhelut\nSMS", "sales_property": ["Toimitustapa", "Päivää", "Paketin tiedot"]},
    "uk-UA": {"fee_include": "eSIM + Дані\nРежим модема", "fee_exclude": "Дзвінки\nSMS", "sales_property": ["Спосіб отримання", "Дні", "Інформація про тариф"]},
    "da-DK": {"fee_include": "eSIM + Data\nHotspotdeling", "fee_exclude": "Opkald\nSMS", "sales_property": ["Modtagelsesmetode", "Dage", "Plandetaljer"]},
    "sv-SE": {"fee_include": "eSIM + Data\nHotspotdelning", "fee_exclude": "Samtal\nSMS", "sales_property": ["Mottagningssätt", "Dagar", "Plandetaljer"]},
    "el-GR": {"fee_include": "eSIM + Δεδομένα\nΚοινή χρήση hotspot", "fee_exclude": "Κλήσεις\nSMS", "sales_property": ["Τρόπος παραλαβής", "Ημέρες", "Λεπτομέρειες πακέτου"]},
    "pl-PL": {"fee_include": "eSIM + Dane\nUdostępnianie hotspotu", "fee_exclude": "Połączenia\nSMS", "sales_property": ["Metoda odbioru", "Dni", "Szczegóły pakietu"]},
    "tr-TR": {"fee_include": "eSIM + Veri\nHotspot Paylaşımı", "fee_exclude": "Arama\nSMS", "sales_property": ["Teslim Alma Yöntemi", "Günler", "Plan Detayları"]},
    "pt-PT": {"fee_include": "eSIM + Dados\nPartilha de hotspot", "fee_exclude": "Chamadas\nSMS", "sales_property": ["Método de recebimento", "Dias", "Detalhes do plano"]},
    "pt-BR": {"fee_include": "eSIM + Dados\nPartilha de hotspot", "fee_exclude": "Chamadas\nSMS", "sales_property": ["Método de recebimento", "Dias", "Detalhes do plano"]},
}

def _spec(decimal, unit, unit_space, daily, total, speed_limited, unlimited, speed="{v}Mbps"):
    return {
        "decimal": decimal, "unit": unit, "unit_space": unit_space,
        "daily": daily, "total": total, "speed_limited": speed_limited,
        "unlimited": unlimited, "speed": speed,
    }


# 销售属性值（规格名称）格式规格表。key 为语言组（见 get_lang_group）。
# {n} = 已按 decimal / unit / unit_space 格式化的“数字+单位”，{speed} = 已格式化的限速值。
SALES_PROPERTY_FORMATS = {
    "ar": _spec(".", "GB", False, "{n} يوميًا", "{n} إجمالي", "{n} يوميًا ({speed})", "{n} يوميًا (بدون حد للسرعة)"),
    "zh-hk": _spec(".", "GB", False, "每日{n}", "總量{n}", "每日{n}（限速{speed}）", "每日{n}（不限速）"),
    "zh-tw": _spec(".", "GB", False, "每日{n}", "總量{n}", "每日{n}（限速{speed}）", "每日{n}（不限速）"),
    "zh-sgmy": _spec(".", "GB", False, "每日{n}", "总量{n}", "每日{n}（限速{speed}）", "每日{n}（不限速）"),
    "en": _spec(".", "GB", False, "Daily {n}", "Total {n}", "Daily {n} ({speed})", "Daily {n} (No Speed Limit)"),
    "ja": _spec(".", "GB", False, "毎日{n}", "合計{n}", "毎日{n}（{speed}）", "毎日{n}（速度制限なし）"),
    "ko": _spec(".", "GB", False, "매일 {n}", "총 {n}", "매일 {n} ({speed})", "매일 {n} (속도 제한 없음)"),
    "th": _spec(".", "GB", False, "{n}/วัน", "รวม {n}", "{n}/วัน ({speed})", "{n}/วัน (ไม่จำกัดความเร็ว)"),
    "de": _spec(",", "GB", True, "{n}/Tag", "Gesamt {n}", "{n}/Tag ({speed})", "{n}/Tag (volle Geschwindigkeit)", "{v} Mbps"),
    "es": _spec(",", "GB", True, "{n}/día", "Total {n}", "{n}/día ({speed})", "{n}/día (sin límite de velocidad)"),
    "fr": _spec(",", "Go", True, "{n}/jour", "Total {n}", "{n}/jour ({speed})", "{n}/jour (débit non limité)"),
    "id": _spec(".", "GB", False, "{n}/hari", "Total {n}", "{n}/hari ({speed})", "{n}/hari (tanpa batas kecepatan)"),
    "ms": _spec(".", "GB", False, "{n}/hari", "Jumlah {n}", "{n}/hari ({speed})", "{n}/hari (kelajuan penuh)", "{v} Mbps"),
    "vi": _spec(",", "GB", False, "{n}/ngày", "Tổng {n}", "{n}/ngày ({speed})", "{n}/ngày (không giới hạn tốc độ)"),
    "it": _spec(",", "GB", True, "{n}/giorno", "Totale {n}", "{n}/giorno ({speed})", "{n}/giorno (senza limiti di velocità)"),
    "nl": _spec(",", "GB", True, "{n}/dag", "Totaal {n}", "{n}/dag ({speed})", "{n}/dag (geen snelheidslimiet)"),
    "ru": _spec(",", "ГБ", True, "{n}/день", "Всего {n}", "{n}/день ({speed})", "{n}/день (без ограничения скорости)", "{v} Мбит/с"),
    "fi": _spec(",", "GB", True, "{n}/päivä", "Yhteensä {n}", "{n}/päivä ({speed})", "{n}/päivä (ei nopeusrajoitusta)"),
    "uk": _spec(",", "ГБ", True, "{n}/день", "Всього {n}", "{n}/день ({speed})", "{n}/день (без обмеження швидкості)", "{v} Мбіт/с"),
    "da": _spec(",", "GB", True, "{n}/dag", "I alt {n}", "{n}/dag ({speed})", "{n}/dag (ingen hastighedsgrænse)"),
    "sv": _spec(",", "GB", True, "{n}/dag", "Totalt {n}", "{n}/dag ({speed})", "{n}/dag (ingen hastighetsgräns)"),
    "el": _spec(",", "GB", True, "{n}/ημέρα", "Σύνολο {n}", "{n}/ημέρα ({speed})", "{n}/ημέρα (χωρίς όριο ταχύτητας)"),
    "pl": _spec(",", "GB", True, "{n}/dzień", "Łącznie {n}", "{n}/dzień ({speed})", "{n}/dzień (bez limitu prędkości)"),
    "tr": _spec(",", "GB", True, "Günlük {n}", "Toplam {n}", "Günlük {n} ({speed})", "Günlük {n} (hız sınırı yok)"),
    "pt": _spec(",", "GB", True, "{n}/dia", "Total {n}", "{n}/dia ({speed})", "{n}/dia (sem limite de velocidade)"),
}

DEFAULT_SPEED_LIMIT_MBPS = "10"

PACKAGE_TRANSLATION_KEYS_BY_LANG_GROUP = {
    "ar": "ar-XX", "zh-hk": "zh-HK", "zh-tw": "zh-TW", "zh-sgmy": "zh-SGMY", "en": "en-US",
    "ja": "ja-JP", "ko": "ko-KR", "th": "th-TH", "de": "de-DE", "es": "es-ES",
    "fr": "fr-FR", "id": "id-ID", "ms": "ms-MY", "vi": "vi-VN", "it": "it-IT",
    "nl": "nl-NL", "ru": "ru-RU", "fi": "fi-FI", "uk": "uk-UA", "da": "da-DK",
    "sv": "sv-SE", "el": "el-GR", "pl": "pl-PL", "tr": "tr-TR", "pt": "pt-BR",
}


@dataclass
class TranslationRuleTask:
    row_number: int
    product_id: str
    source: str
    row_data: dict[str, str]


@dataclass
class TranslationExecutionConfig:
    selected_languages: list[str] = field(default_factory=lambda: DEFAULT_SELECTED_LANGUAGES.copy())
    selected_modules: list[str] = field(default_factory=lambda: DEFAULT_SELECTED_MODULES.copy())
    sales_attr_mode: str = "strict"
    long_day_package_template_mode: str = "none"
    package_special_rules: dict[str, bool] = field(default_factory=dict)
    online: bool = True
    include_unopened_locales: bool = False


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("_x000D_", "")


def normalize_field_key(key: str) -> str:
    text = str(key or "").strip()
    if text.startswith("{") and text.endswith("}"):
        return "{" + text[1:-1].lower() + "}"
    return text


def parse_product_id_cell(value: str) -> str:
    text = normalize_cell(value)
    return text if text.isdigit() else ""


def load_translation_rule_tasks(path: str) -> list[TranslationRuleTask]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [normalize_field_key(cell.value) for cell in ws[1]]
    if "productId" not in headers and "productid" not in {h.lower() for h in headers}:
        raise ValueError("任务模板缺少 productId 列")
    lower_headers = {h.lower(): h for h in headers}
    product_col = headers.index(lower_headers["productid"])

    tasks: list[TranslationRuleTask] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cleaned = [normalize_cell(value) for value in values]
        row_data = dict(zip(headers, cleaned))
        source = cleaned[product_col] if product_col < len(cleaned) else ""
        product_id = parse_product_id_cell(source)
        if product_id:
            tasks.append(TranslationRuleTask(row_number=row_number, product_id=product_id, source=source, row_data=row_data))
    if not tasks:
        raise ValueError("未找到有效 productId 任务行")
    return tasks


def get_lang_group(locale: str) -> str:
    lower = str(locale or "").strip().replace("_", "-").lower()
    if lower.startswith("ar-") or lower == "ar" or "arabic" in lower:
        return "ar"
    if "zh-hk" in lower or "hong" in lower:
        return "zh-hk"
    if "zh-tw" in lower or "taiwan" in lower:
        return "zh-tw"
    if "zh-sg" in lower or "zh-my" in lower or "singapore" in lower or "malaysia" in lower:
        return "zh-sgmy"
    return lower.split("-")[0]


def interface_locales_for_lang_group(lang_group: str) -> list[str]:
    return list(INTERFACE_LOCALES_BY_LANG_GROUP.get(str(lang_group or "").lower(), []))


def is_rtl_lang_group(lang_group: str) -> bool:
    return str(lang_group or "").lower() in RTL_LANG_GROUPS


def is_rtl_locale(locale: str) -> bool:
    return is_rtl_lang_group(get_lang_group(locale))


def html_paragraph_attrs(locale: str) -> str:
    return ' dir="rtl" style="text-align:right;"' if is_rtl_locale(locale) else ""


def fixed_content_key(locale: str) -> str:
    group = get_lang_group(locale)
    if group in PACKAGE_TRANSLATION_KEYS_BY_LANG_GROUP:
        return PACKAGE_TRANSLATION_KEYS_BY_LANG_GROUP[group]
    return str(locale or "")


def normalize_special_rule_config(config: dict[str, Any] | None) -> dict[str, bool]:
    result = {key: True for key in SPECIAL_RULE_LABELS}
    if isinstance(config, dict):
        for key in result:
            if key in config:
                result[key] = bool(config[key])
    return result


def should_force_total_stop(days: str, data_type: str, config: dict[str, Any]) -> bool:
    return normalize_special_rule_config(config).get("force_365_total_stop", True) and str(days).strip() == "365" and str(data_type).strip() in {"总量", "總量"}


def get_long_day_forced_script_key(days: str, mode: str) -> str | None:
    if str(days).strip() not in {"90", "180"}:
        return None
    if mode == "stop":
        return "后停用"
    if mode == "throttle":
        return "后降速"
    return None


def select_usedup_template_key(script_key: str, is_daily: bool, data_capacity: str, available_templates: dict[str, Any], config: dict[str, Any]) -> str | None:
    rules = normalize_special_rule_config(config)
    normalized = str(data_capacity or "").replace(" ", "").replace("－", "-").upper()
    has_100gb = "100GB" in normalized
    has_100gb_unlimited = has_100gb and "不限速" in normalized
    has_20gb = "20GB" in normalized
    templates = available_templates if isinstance(available_templates, dict) else {}
    if script_key == "吃到饱_提前安装":
        if is_daily and has_100gb_unlimited and rules["daily_100gb_unlimited_template"] and "daily_100gb_unlimited" in templates:
            return "daily_100gb_unlimited"
        if is_daily and has_100gb and rules["daily_100gb_template"]:
            return "daily_100gb"
        if is_daily and has_20gb and rules["unlimited_daily_20gb_template"]:
            return "daily_20gb"
        if rules["unlimited_capacity_whitelist_skip"] and not (is_daily and (has_100gb or has_20gb)):
            return None
        return "daily_generic" if is_daily else "total"
    if script_key == "总量停用每日降速":
        if is_daily and has_20gb and rules["mixed_daily_20gb_template"] and "daily_20gb" in templates:
            return "daily_20gb"
        return "daily_generic" if is_daily else "total"
    if is_daily and has_100gb_unlimited and rules["daily_100gb_unlimited_template"] and "daily_100gb_unlimited" in templates:
        return "daily_100gb_unlimited"
    if is_daily and has_100gb and rules["daily_100gb_template"] and "daily_100gb" in templates:
        return "daily_100gb"
    return "daily_generic" if is_daily else "total"


def _rich_cell_to_formatted_data(cell: Any) -> list[tuple[str, list[str]]]:
    """Keep Excel run-level bold/italic/underline; whole-cell fonts are ignored as they are usually accidental."""
    value = cell.value
    if value is None:
        return []
    from openpyxl.cell.rich_text import CellRichText

    if not isinstance(value, CellRichText):
        text = str(value).replace("_x000D_", "")
        return [(text, [])] if text else []
    data: list[tuple[str, list[str]]] = []
    for block in value:
        font = getattr(block, "font", None)
        text = str(getattr(block, "text", block)).replace("_x000D_", "")
        if not text:
            continue
        tags = [tag for tag, attr in (("bold", "b"), ("italic", "i"), ("underline", "u")) if font is not None and getattr(font, attr, None) not in (None, False, "none")]
        if data and data[-1][1] == tags:
            data[-1] = (data[-1][0] + text, tags)
        else:
            data.append((text, tags))
    return data


def _headers_for_sheet(sheet: Any) -> dict[str, int]:
    return {normalize_cell(cell.value): index for index, cell in enumerate(sheet[1], start=1) if normalize_cell(cell.value)}


def create_translation_task_template_xlsx(path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    languages = ALL_TEMPLATE_LANGUAGES
    headers = [
        "productId", "{apn}",
        *[f"{{product_name_{lang}}}" for lang in languages],
        *[f"{{mm_detail_{lang}}}" for lang in languages],
        *[f"{{coverage_{lang}}}" for lang in languages],
        *[f"{{carrier_{lang}}}" for lang in languages],
        "{connection_gen}",
        "{usedup_throttle}", "{reset_time_calc}", "{inst_adv}",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "数据模板"
    ws.append(headers)
    ws.append([
        "100000003",
        "plus",
        *["" for _ in range(len(headers) - 6)],
        "5G",
        "后降速",
        "+8零点",
        "可提前安装",
    ])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column, header in enumerate(headers, start=1):
        if re.search(r"_(ar)\}$", str(header or "")):
            ws.cell(1, column).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
            for row in range(2, ws.max_row + 1):
                ws.cell(row, column).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
    for column in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 24 if column > 2 else 18
    ws.freeze_panes = "A2"

    source = wb.create_sheet("数据源")
    source.append(["用完处理", "重置计时", "安装说明"])
    for index, value in enumerate(["后降速", "后停用", "吃到饱", "总量停用每日降速"], start=2):
        source.cell(index, 1, value)
    for index, value in enumerate(["+7零点", "+8零点", "+9零点", "激活后每24h"], start=2):
        source.cell(index, 2, value)
    for index, value in enumerate(["可提前安装", "不可提前安装"], start=2):
        source.cell(index, 3, value)
    validation_columns = {
        "{usedup_throttle}": "'数据源'!$A$2:$A$5",
        "{reset_time_calc}": "'数据源'!$B$2:$B$5",
        "{inst_adv}": "'数据源'!$C$2:$C$3",
    }
    for header, formula in validation_columns.items():
        column = headers.index(header) + 1
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        letter = openpyxl.utils.get_column_letter(column)
        dv.add(f"{letter}2:{letter}1000")
    wb.save(path)


DETAIL_FORMAT_GUIDE_SHEET_NAME = "格式说明"
DETAIL_FORMAT_GUIDE_ROWS = [
    ("标题", "行首写 ## 或 #，井号后要有空格", "## 使用前须知", "大号标题 <h2>"),
    ("加粗", "在 Excel 中选中部分文字设为加粗；或写 **文字**", "请在 **出发前** 安装", "加粗文字 <strong>"),
    ("项目列表", "行首写 - 或 *，后面要有空格；连续多行合并为同一个列表", "- 支持热点共享\n- 可提前安装", "带圆点的列表 <ul><li>"),
    ("链接", "[显示文字](网址)；没写 https:// 会自动补上，也支持 mailto: 和 tel:", "[联系客服](www.trip.com/help)", "可点击的链接 <a>"),
    ("换行 / 空行", "单元格内换行（Windows：Alt+Enter；Mac：Option+Return）分段；空行保留为空段落", "第一段\n\n第二段", "段落 <p>"),
]
DETAIL_FORMAT_GUIDE_NOTES = [
    "仅“产品说明”“注意事项”“订后出行提醒”三列支持以上格式；“使用方法”“携程处理注意事项”为纯文本。",
    "整个单元格设为加粗不会生效，请只选中需要强调的文字。",
    "斜体、下划线、颜色等携程编辑器不支持的格式会被忽略，只保留文字。",
    "•、■ 等符号按普通文字输出，不会变成项目列表；需要列表请用“- ”。",
    "阿拉伯语的标题、段落和列表项会自动右对齐。",
    "{变量名} 占位符可与以上格式混用，例如：- **覆盖地区**：{coverage_en}",
]
DETAIL_RICH_HEADER_SUFFIX = "（支持格式填写）"
DETAIL_HEADER_NOTE_RE = re.compile(r"\s*[（(][^（）()]*[）)]\s*$")
DETAIL_FORMAT_HEADER_HINT = "支持格式：\n## 标题\n**加粗**（或 Excel 选中部分文字加粗）\n- 项目列表\n[显示文字](网址)\n详见“格式说明”工作表"


def create_detail_template_xlsx(path: str, languages: list[str] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font

    languages = languages or ALL_TEMPLATE_LANGUAGES
    wb = Workbook()
    ws = wb.active
    ws.title = DETAIL_SHEET_NAME
    column_names = ["语言", "使用方法", "产品说明", "注意事项", "订后出行提醒", "携程处理注意事项"]
    rich_names = {name for name, key in DETAIL_COLUMNS if key in DETAIL_RICH_KEYS}
    headers = [f"{name}{DETAIL_RICH_HEADER_SUFFIX}" if name in rich_names else name for name in column_names]
    ws.append(headers)
    for lang in languages:
        ws.append([lang, *["" for _ in headers[1:]]])
    widths = [12, 42, 58, 58, 58, 58]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 96
        lang = normalize_cell(ws.cell(row, 1).value).lower()
        for column in range(1, ws.max_column + 1):
            if is_rtl_lang_group(lang):
                ws.cell(row, column).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
            else:
                ws.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
    for column, name in enumerate(column_names, start=1):
        if name in rich_names:
            comment = Comment(DETAIL_FORMAT_HEADER_HINT, "Trip.com规格填充工具")
            comment.width, comment.height = 260, 150
            ws.cell(1, column).comment = comment
    _append_detail_format_guide_sheet(wb)
    wb.save(path)


def _append_detail_format_guide_sheet(wb: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    guide = wb.create_sheet(DETAIL_FORMAT_GUIDE_SHEET_NAME)
    guide.append(["富文本格式写法（适用于详情模板的 产品说明 / 注意事项 / 订后出行提醒 列）"])
    guide["A1"].font = Font(size=14, bold=True)
    guide.append([])
    guide.append(["效果", "模板写法", "示例", "在携程编辑器中显示为"])
    for cell in guide[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in DETAIL_FORMAT_GUIDE_ROWS:
        guide.append(list(row))
    guide.append([])
    guide.append(["注意事项"])
    guide.cell(guide.max_row, 1).font = Font(bold=True)
    for index, note in enumerate(DETAIL_FORMAT_GUIDE_NOTES, start=1):
        guide.append([f"{index}. {note}"])
        guide.merge_cells(start_row=guide.max_row, start_column=1, end_row=guide.max_row, end_column=4)
    for row in guide.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for letter, width in zip("ABCD", (14, 52, 36, 26)):
        guide.column_dimensions[letter].width = width


def create_package_template_xlsx(path: str, template_data: dict[str, Any] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    template_data = template_data or build_default_package_template_schema()
    wb = Workbook()
    combo = wb.active
    combo.title = PACKAGE_COMBINATION_SHEET
    combo.append(["套餐类型", "模板文本"])
    combo_data = template_data.get("combination_template", {})
    if isinstance(combo_data, str):
        combo_data = {"daily": combo_data, "total": combo_data}
    for key in ("daily", "total"):
        combo.append([key, combo_data.get(key, "")])

    capacity = wb.create_sheet(PACKAGE_CAPACITY_SHEET)
    capacity.append(["用完处理", "容量场景", "语言", "模板文本"])
    for script_name, script_templates in (template_data.get("scripts") or {}).items():
        if not isinstance(script_templates, dict):
            continue
        for capacity_key, lang_templates in script_templates.items():
            if not isinstance(lang_templates, dict):
                continue
            for lang in ALL_TEMPLATE_LANGUAGES:
                capacity.append([script_name, capacity_key, lang, lang_templates.get(lang, "")])

    install = wb.create_sheet(PACKAGE_INSTALL_SHEET)
    install.append(["安装类型", "语言", "模板文本"])
    for install_type, lang_templates in (template_data.get("install_instructions") or {}).items():
        if isinstance(lang_templates, dict):
            for lang in ALL_TEMPLATE_LANGUAGES:
                install.append([install_type, lang, lang_templates.get(lang, "")])

    day_calc = wb.create_sheet(PACKAGE_DAY_CALC_SHEET)
    day_calc.append(["语言", "模板文本"])
    for lang in ALL_TEMPLATE_LANGUAGES:
        day_calc.append([lang, (template_data.get("day_calc") or {}).get(lang, "")])

    reset = wb.create_sheet(PACKAGE_RESET_CALC_SHEET)
    reset.append(["重置类型", "语言", "模板文本"])
    for reset_type, lang_templates in (template_data.get("reset_time_calc") or {}).items():
        if isinstance(lang_templates, dict):
            for lang in ALL_TEMPLATE_LANGUAGES:
                reset.append([reset_type, lang, lang_templates.get(lang, "")])

    reset_options = wb.create_sheet(PACKAGE_RESET_OPTIONS_SHEET)
    reset_options.append(["语言", "选项1", "选项2", "选项3", "选项4"])
    for lang in ALL_TEMPLATE_LANGUAGES:
        options = list((template_data.get("reset_time_options") or {}).get(lang, []))
        reset_options.append([lang, *options[:4], *["" for _ in range(max(0, 4 - len(options)))]] )

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            sheet.column_dimensions[letter].width = 76 if letter == sheet.cell(1, sheet.max_column).column_letter else 18
        for row in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 64
            headers = _headers_for_sheet(sheet)
            lang_col = headers.get("语言")
            lang = normalize_cell(sheet.cell(row, lang_col).value).lower() if lang_col else ""
            for column in range(1, sheet.max_column + 1):
                if is_rtl_lang_group(lang):
                    sheet.cell(row, column).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True, readingOrder=2)
                else:
                    sheet.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


def load_detail_template_xlsx(path: str) -> dict[str, dict[str, Any]]:
    wb = openpyxl.load_workbook(path, data_only=False, rich_text=True)
    if DETAIL_SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"详情模板Excel缺少'{DETAIL_SHEET_NAME}'工作表")
    ws = wb[DETAIL_SHEET_NAME]
    # Headers may carry a trailing note such as "产品说明（支持格式填写）".
    headers = {DETAIL_HEADER_NOTE_RE.sub("", name): column for name, column in _headers_for_sheet(ws).items()}
    missing = [name for name in DETAIL_REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"工作表'{DETAIL_SHEET_NAME}'缺少列: {', '.join(missing)}")
    data: dict[str, dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        lang = normalize_cell(ws.cell(row, headers["语言"]).value).lower()
        if not lang:
            continue
        item: dict[str, Any] = {}
        for column_name, key in DETAIL_COLUMNS:
            if column_name not in headers:
                continue
            cell = ws.cell(row, headers[column_name])
            item[key] = _rich_cell_to_formatted_data(cell) if key in DETAIL_RICH_KEYS else normalize_cell(cell.value)
        item.setdefault("usage", "")
        item.setdefault("description", [])
        item.setdefault("notice", [])
        item.setdefault("post_booking_reminder", [])
        item.setdefault("operation_attention", "")
        data[lang] = item
    if not data:
        raise ValueError("详情模板Excel未找到有效语言行")
    return data


def build_default_package_template_schema() -> dict[str, Any]:
    return {
        "combination_template": {"daily": "", "total": ""},
        "reset_time_options": {lang: [] for lang in ALL_TEMPLATE_LANGUAGES},
        "day_calc": {lang: "" for lang in ALL_TEMPLATE_LANGUAGES},
        "install_instructions": {"pre_install_allowed": {lang: "" for lang in ALL_TEMPLATE_LANGUAGES}, "pre_install_not_allowed": {lang: "" for lang in ALL_TEMPLATE_LANGUAGES}},
        "reset_time_calc": {key: {lang: "" for lang in ALL_TEMPLATE_LANGUAGES} for key in ["+7零点", "+8零点", "+9零点", "激活后每24h"]},
        "scripts": {
            script: {name: {lang: "" for lang in ALL_TEMPLATE_LANGUAGES} for name in ["daily_100gb", "daily_100gb_unlimited", "daily_20gb", "daily_generic", "total"]}
            for script in ["后降速", "后停用", "吃到饱_提前安装", "总量停用每日降速"]
        },
    }


def load_package_template_xlsx(path: str) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True)
    data = build_default_package_template_schema()
    if PACKAGE_COMBINATION_SHEET in wb.sheetnames:
        ws = wb[PACKAGE_COMBINATION_SHEET]
        headers = _headers_for_sheet(ws)
        for row in range(2, ws.max_row + 1):
            name = normalize_cell(ws.cell(row, headers.get("套餐类型", 1)).value)
            text = str(ws.cell(row, headers.get("模板文本", 2)).value or "").replace("_x000D_", "")
            if name in {"daily", "每日", "每日套餐"}:
                data["combination_template"]["daily"] = text
            elif name in {"total", "总量", "总量套餐"}:
                data["combination_template"]["total"] = text
    if PACKAGE_CAPACITY_SHEET in wb.sheetnames:
        ws = wb[PACKAGE_CAPACITY_SHEET]
        headers = _headers_for_sheet(ws)
        for row in range(2, ws.max_row + 1):
            script = normalize_cell(ws.cell(row, headers.get("用完处理", 1)).value)
            capacity = normalize_cell(ws.cell(row, headers.get("容量场景", 2)).value)
            lang = normalize_cell(ws.cell(row, headers.get("语言", 3)).value).lower()
            text = str(ws.cell(row, headers.get("模板文本", 4)).value or "").replace("_x000D_", "")
            if script == "吃到饱":
                script = "吃到饱_提前安装"
            if script and capacity and lang:
                data["scripts"].setdefault(script, {}).setdefault(capacity, {})[lang] = text
    for sheet_name, target_key, first_col in [
        (PACKAGE_INSTALL_SHEET, "install_instructions", "安装类型"),
        (PACKAGE_RESET_CALC_SHEET, "reset_time_calc", "重置类型"),
    ]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = _headers_for_sheet(ws)
        for row in range(2, ws.max_row + 1):
            kind = normalize_cell(ws.cell(row, headers.get(first_col, 1)).value)
            lang = normalize_cell(ws.cell(row, headers.get("语言", 2)).value).lower()
            text = str(ws.cell(row, headers.get("模板文本", 3)).value or "").replace("_x000D_", "")
            if kind in {"可提前安装", "提前安装"}:
                kind = "pre_install_allowed"
            elif kind in {"不可提前安装", "不提前安装"}:
                kind = "pre_install_not_allowed"
            if kind and lang:
                data[target_key].setdefault(kind, {})[lang] = text
    if PACKAGE_DAY_CALC_SHEET in wb.sheetnames:
        ws = wb[PACKAGE_DAY_CALC_SHEET]
        headers = _headers_for_sheet(ws)
        for row in range(2, ws.max_row + 1):
            lang = normalize_cell(ws.cell(row, headers.get("语言", 1)).value).lower()
            text = str(ws.cell(row, headers.get("模板文本", 2)).value or "").replace("_x000D_", "")
            if lang:
                data["day_calc"][lang] = text
    if PACKAGE_RESET_OPTIONS_SHEET in wb.sheetnames:
        ws = wb[PACKAGE_RESET_OPTIONS_SHEET]
        headers = _headers_for_sheet(ws)
        option_cols = sorted((idx, name) for name, idx in headers.items() if name.startswith("选项"))
        for row in range(2, ws.max_row + 1):
            lang = normalize_cell(ws.cell(row, headers.get("语言", 1)).value).lower()
            if not lang:
                continue
            data["reset_time_options"][lang] = [normalize_cell(ws.cell(row, idx).value) for idx, _name in option_cols if normalize_cell(ws.cell(row, idx).value)]
    return data


# Template markup for the Ctrip translate page's Quill toolbar (header 2 / bold / bullet list / link).
# Excel has no native headings, lists or inline links, so these are written as Markdown-style markers.
QUILL_BLOCK_MARKER_RE = re.compile(r"^\s*(#{1,6}|[-*])\s+(?=\S)")
QUILL_BOLD_MARKER_RE = re.compile(r"\*\*(?=\S)(.+?)(?<=\S)\*\*")
QUILL_LINK_MARKER_RE = re.compile(r"\[([^\[\]]+)\]\(\s*([^()\s]+)\s*\)")


def _drop_leading_chars(runs: list[tuple[str, list[str]]], count: int) -> list[tuple[str, list[str]]]:
    result: list[tuple[str, list[str]]] = []
    for text, tags in runs:
        if count >= len(text):
            count -= len(text)
            continue
        result.append((text[count:], tags))
        count = 0
    return result


def _quill_link_href(raw_url: str) -> str:
    url = re.sub(r"<[^>]+>", "", raw_url).replace("&amp;", "&").strip()
    if not re.match(r"^(https?:|mailto:|tel:)", url, re.IGNORECASE):
        url = "https://" + url
    return url.replace("&", "&amp;").replace('"', "&quot;")


def _quill_inline_html(runs: list[tuple[str, list[str]]]) -> str:
    html = ""
    for text, tags in runs:
        if not text:
            continue
        escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        # The Quill editor only whitelists bold among inline text marks (plus links).
        if "bold" in tags and escaped.strip():
            escaped = f"<strong>{escaped}</strong>"
        html += escaped
    html = QUILL_BOLD_MARKER_RE.sub(r"<strong>\1</strong>", html)
    return QUILL_LINK_MARKER_RE.sub(
        lambda match: f'<a href="{_quill_link_href(match.group(2))}" rel="noopener noreferrer" target="_blank">{match.group(1)}</a>', html
    )


class DetailTemplateExcelStore:
    def __init__(self) -> None:
        self.template_data: dict[str, dict[str, Any]] = {}
        self.current_template_path: str | None = None
        self.current_template_mtime_ns: int | None = None

    def load_template_from_path(self, path: str) -> dict[str, dict[str, Any]]:
        self.template_data = load_detail_template_xlsx(path)
        self.current_template_path = path
        self.current_template_mtime_ns = os.stat(path).st_mtime_ns
        return self.template_data

    def get_template_content(self, lang_group: str) -> dict[str, Any] | None:
        return self.template_data.get(lang_group)

    def replace_placeholders(self, template_text: str, lang_code: str, data_row: dict[str, str]) -> str:
        if not template_text:
            return ""
        suffix = get_lang_group(lang_code)

        def replace_match(match: re.Match[str]) -> str:
            placeholder = match.group(1)
            direct_key = "{" + placeholder + "}"
            lower_key = "{" + placeholder.lower() + "}"
            if lower_key in data_row:
                return str(data_row.get(lower_key, match.group(0)))
            if direct_key in data_row:
                return str(data_row.get(direct_key, match.group(0)))
            parts = re.match(r"^([^_]+)_(.+)$", placeholder)
            if parts and parts.group(2).lower() != suffix:
                suffix_key = "{" + parts.group(1).lower() + "_" + suffix + "}"
                if suffix_key in data_row:
                    return str(data_row.get(suffix_key, match.group(0)))
            return match.group(0)

        return re.sub(r"\{([^}]+)\}", replace_match, str(template_text))

    def expand_placeholders_in_formatted_data(self, data: list[tuple[str, list[str]]], lang_code: str, data_row: dict[str, str]) -> list[tuple[str, list[str]]]:
        if not data:
            return []
        return [(self.replace_placeholders(text, lang_code, data_row), tags) for text, tags in data]

    def convert_formatted_to_html(self, data: Any, lang_code: str = "") -> str:
        if not isinstance(data, list):
            return self._format_text_for_quill(str(data or ""), lang_code)
        # Split runs per line first so inline tags never span block boundaries.
        lines: list[list[tuple[str, list[str]]]] = [[]]
        for text, tags in data:
            for index, part in enumerate(str(text).split("\n")):
                if index:
                    lines.append([])
                if part:
                    lines[-1].append((part, list(tags or [])))
        attrs = html_paragraph_attrs(lang_code)
        blocks: list[str] = []
        list_items: list[str] = []

        def flush_list() -> None:
            if list_items:
                blocks.append("<ul>" + "".join(list_items) + "</ul>")
                list_items.clear()

        for runs in lines:
            if runs:
                runs[-1] = (runs[-1][0].rstrip(), runs[-1][1])
            plain = "".join(text for text, _tags in runs)
            if not plain.strip():
                flush_list()
                blocks.append(f"<p{attrs}><br></p>")
                continue
            block_tag = "p"
            marker = QUILL_BLOCK_MARKER_RE.match(plain)
            if marker:
                block_tag = "h2" if marker.group(1).startswith("#") else "li"
                runs = _drop_leading_chars(runs, marker.end())
            html = _quill_inline_html(runs)
            if block_tag == "li":
                list_items.append(f"<li{attrs}>{html}</li>")
                continue
            flush_list()
            blocks.append(f"<{block_tag}{attrs}>{html}</{block_tag}>")
        flush_list()
        return "".join(blocks)

    def _format_text_for_quill(self, text: str, lang_code: str = "") -> str:
        attrs = html_paragraph_attrs(lang_code)
        return "".join(f"<p{attrs}>{line.strip().replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;') if line.strip() else '<br>'}</p>" for line in text.split("\n")) if text else ""


def sales_property_lang_candidates(locale: str) -> list[str]:
    raw = str(locale or "").strip().replace("_", "-").lower()
    candidates: list[str] = []

    def add(value: str | None) -> None:
        if value and value not in candidates:
            candidates.append(value)

    add(get_lang_group(raw))
    add(raw)
    add(raw.split("-")[0] if raw else "")
    # 中国大陆简体（zh / zh-CN）不在翻译件管理范围内，不做别名映射，直接跳过
    aliases = {"zh-hans": "zh-sgmy", "zh-sg": "zh-sgmy", "zh-my": "zh-sgmy", "zh-hant": "zh-hk"}
    for item in list(candidates):
        add(aliases.get(item))
    return candidates


def classify_daily_100gb_sales_property(text: str) -> str:
    normalized = re.sub(r"\s+", "", str(text or "")).upper()
    if re.match(r"^每日-?100GB-?不限速$", normalized):
        return "daily_100gb_unlimited"
    if re.match(r"^每日-?100GB(?:-?限速[0-9.]+[KMG]?BPS)?$", normalized):
        return "daily_100gb_standard"
    return ""


def extract_speed_limit(text: str) -> tuple[str, str]:
    """从“限速10Mbps”这类原文中取出 (数值, 单位)，单位归一为 Kbps/Mbps/Gbps/bps；取不到时回退默认值。"""
    match = re.search(r"限速\s*([0-9]+(?:\.[0-9]+)?)\s*([KMG]?)BPS", str(text or ""), flags=re.IGNORECASE)
    if not match:
        return DEFAULT_SPEED_LIMIT_MBPS, "Mbps"
    prefix = match.group(2).upper()
    return match.group(1), f"{prefix}bps" if prefix else "bps"


def _format_speed_limit(spec: dict, text: str) -> str:
    value, unit = extract_speed_limit(text)
    if unit == "Mbps":
        return spec["speed"].format(v=value)
    return f"{value}{unit}"


def _sales_property_format(locale: str) -> dict | None:
    for candidate in sales_property_lang_candidates(locale):
        spec = SALES_PROPERTY_FORMATS.get(candidate)
        if spec:
            return spec
    return None


def _format_data_amount(amount: str, spec: dict) -> str:
    number = str(amount).replace(".", spec["decimal"])
    return f"{number} {spec['unit']}" if spec["unit_space"] else f"{number}{spec['unit']}"


def translate_sales_property_value(origin: str, locale: str, mode: str = "strict") -> str:
    text = str(origin or "").strip()
    category = classify_daily_100gb_sales_property(text)
    if category:
        spec = _sales_property_format(locale)
        if not spec:
            return ""
        n = _format_data_amount("100", spec)
        if category == "daily_100gb_unlimited":
            return spec["unlimited"].format(n=n)
        speed = _format_speed_limit(spec, text)
        return spec["speed_limited"].format(n=n, speed=speed)
    strict_match = re.match(r"^(?:[45]G网络-)?(每日|总量)-(\d+(?:\.\d+)?)GB\s*$", text)
    prefix_match = re.match(r"^(?:[45]G网络-)?(每日|总量)-(\d+(?:\.\d+)?)GB", text) if mode == "fuzzy_prefix" else None
    match = strict_match or prefix_match
    if not match:
        return ""
    suffix = text[match.end():] if prefix_match and not strict_match else ""
    spec = _sales_property_format(locale)
    if not spec:
        return ""
    period, amount = match.group(1), match.group(2)
    template = spec["daily"] if period == "每日" else spec["total"]
    return template.format(n=_format_data_amount(amount, spec)) + suffix


def package_info_from_origin_with_descriptor(origin: str) -> tuple[str, str, str, str]:
    text = str(origin or "")
    match = re.search(r"QR code-(\d+)天-.*?([总總]量|每日)-([^<\n\r]+)", text)
    if not match:
        return "", "", "", ""
    days, data_type, data_cap_raw = match.groups()
    capacity_descriptor = data_cap_raw.strip()
    cap_match = re.search(r"(\d+(?:\.\d+)?\s*[Gg][Bb])", data_cap_raw)
    return days, data_type, cap_match.group(1) if cap_match else capacity_descriptor, capacity_descriptor


def package_info_from_origin(origin: str) -> tuple[str, str, str]:
    days, data_type, data_capacity, _capacity_descriptor = package_info_from_origin_with_descriptor(origin)
    return days, data_type, data_capacity


def package_template_text(
    package_template: dict[str, Any],
    locale: str,
    data_capacity: str,
    data_type: str,
    data_row: dict[str, str],
    config: TranslationExecutionConfig,
    days: str,
    capacity_descriptor: str = "",
) -> str:
    usedup_mode = data_row.get("{usedup_throttle}") or "后降速"
    if usedup_mode == "吃到饱":
        script_key = "吃到饱_提前安装"
    elif usedup_mode == "总量停用每日降速":
        script_key = "总量停用每日降速"
    elif usedup_mode == "后停用":
        script_key = "后停用"
    else:
        script_key = "后降速"
    force_total = should_force_total_stop(days, data_type, config.package_special_rules)
    forced = get_long_day_forced_script_key(days, config.long_day_package_template_mode)
    effective_script = forced or ("后停用" if force_total and "每日" not in data_type else script_key)
    lang_group = get_lang_group(locale)
    is_daily = "每日" in data_type
    combination_data = package_template.get("combination_template", {})
    combination = combination_data.get("daily" if is_daily else "total") if isinstance(combination_data, dict) else str(combination_data or "")
    script_templates = package_template.get("scripts", {}).get(effective_script, {})

    def fallback(mapping: dict[str, str] | None) -> str:
        if not isinstance(mapping, dict):
            return ""
        return str(mapping.get(lang_group) or mapping.get("en") or "")

    selector_capacity = capacity_descriptor or data_capacity
    template_name = "total" if force_total and not is_daily else select_usedup_template_key(effective_script, is_daily, selector_capacity, script_templates, config.package_special_rules)
    if template_name is None:
        return ""
    usedup = fallback(script_templates.get(template_name, {}))
    if usedup_mode and not forced and not force_total:
        excel_script = {"吃到饱": "吃到饱_提前安装"}.get(usedup_mode, usedup_mode)
        excel_templates = package_template.get("scripts", {}).get(excel_script, {})
        excel_template_name = select_usedup_template_key(excel_script, is_daily, selector_capacity, excel_templates, config.package_special_rules)
        if excel_template_name:
            usedup = fallback(excel_templates.get(excel_template_name, {}))
    reset = ""
    if data_row.get("{reset_time_calc}"):
        reset = fallback(package_template.get("reset_time_calc", {}).get(data_row["{reset_time_calc}"], {}))
    install_key = "pre_install_allowed"
    if data_row.get("{inst_adv}"):
        install_key = {"可提前安装": "pre_install_allowed", "提前安装": "pre_install_allowed", "不可提前安装": "pre_install_not_allowed", "不提前安装": "pre_install_not_allowed"}.get(data_row["{inst_adv}"], install_key)
    inst = fallback(package_template.get("install_instructions", {}).get(install_key, {}))
    day_calc = fallback(package_template.get("day_calc", {}))
    if not combination or not usedup or not inst or not day_calc:
        return ""
    usedup = usedup.format(data_capacity=data_capacity)
    result = combination.replace("{usedup_throttle}", usedup).replace("{inst_adv}", inst).replace("{reset_time_calc}", reset).replace("{Day_calc}", day_calc)
    attrs = html_paragraph_attrs(locale)
    return f"<p{attrs}>" + result.replace("\n", f"</p><p{attrs}>") + "</p>"


def detail_template_module_reuses_english(store: DetailTemplateExcelStore, locale: str, module_key: str) -> bool:
    lang_group = get_lang_group(locale)
    if lang_group == "en":
        return False
    localized = (store.get_template_content(lang_group) or {}).get(module_key, [])
    english = (store.get_template_content("en") or {}).get(module_key, [])

    def normalized_text(value: Any) -> str:
        if isinstance(value, list):
            value = "".join(str(text) for text, _tags in value)
        return re.sub(r"\s+", " ", str(value or "")).strip()

    localized_text = normalized_text(localized)
    english_text = normalized_text(english)
    return bool(localized_text and english_text and localized_text == english_text)


def generated_translation_static_targets(store: DetailTemplateExcelStore, locale: str, row_data: dict[str, str], selected_modules: list[str]) -> dict[str, str]:
    lang_group = get_lang_group(locale)
    templates = store.get_template_content(lang_group) or {}
    fixed = FIXED_CONTENT.get(fixed_content_key(locale), {})
    targets: dict[str, str] = {}
    default_product_name = "{product_name_zh-hk}" if lang_group == "zh-hk" else "{product_name_zh-tw}" if lang_group == "zh-tw" else f"{{product_name_{lang_group}}}"
    default_mm_detail = "{mm_detail_zh-hk}" if lang_group == "zh-hk" else "{mm_detail_zh-tw}" if lang_group == "zh-tw" else f"{{mm_detail_{lang_group}}}"
    plain_templates = {
        "product_name": templates.get("product_name", default_product_name),
        "mm_detail": templates.get("mm_detail", default_mm_detail),
        "fee_include": fixed.get("fee_include") or templates.get("fee_include", ""),
        "fee_exclude": fixed.get("fee_exclude") or templates.get("fee_exclude", ""),
        "usage": templates.get("usage", ""),
        "operation_attention": templates.get("operation_attention", ""),
    }
    field_code_by_module = {
        "product_name": "productName",
        "mm_detail": "productDesc",
        "fee_include": "feeInclude",
        "fee_exclude": "feeExclude",
        "usage": "howToUse",
        "operation_attention": "operationAttention",
    }
    for module_key, template in plain_templates.items():
        if module_key in selected_modules:
            value = template if module_key in {"fee_include", "fee_exclude"} and fixed else store.replace_placeholders(str(template), locale, row_data)
            if value.strip():
                targets[field_code_by_module[module_key]] = value
    for module_key, field_code in [("description", "productIntroduction"), ("notice", "attention"), ("post_booking_reminder", "postBookingTravelReminder")]:
        if module_key not in selected_modules:
            continue
        if module_key == "post_booking_reminder" and detail_template_module_reuses_english(store, locale, module_key):
            continue
        expanded = store.expand_placeholders_in_formatted_data(templates.get(module_key, []), locale, row_data)
        value = store.convert_formatted_to_html(expanded, locale)
        if re.sub(r"<[^>]+>", "", value).strip():
            targets[field_code] = value
    return targets
