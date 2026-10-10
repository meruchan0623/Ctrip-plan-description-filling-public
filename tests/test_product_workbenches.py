import os
import tempfile
import threading
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook

import app

QT_APP = app.QApplication.instance() or app.QApplication([])


class ProductWorkbenchSpecTests(unittest.TestCase):
    def test_unified_product_edit_workbench_is_explicit(self):
        self.assertEqual(app.product_workbench_spec(app.ProductWorkbench.PRODUCT_INFO).title, "产品信息接口修改")
        self.assertEqual(
            app.product_workbench_spec(app.ProductWorkbench.PRODUCT_INFO).allowed_field_names,
            ("costPrice", "sale", "productName", "vendorPackageId", "connectionCode", "introduction"),
        )
        self.assertEqual(app.product_workbench_spec(app.ProductWorkbench.PRICE_INVENTORY).title, "价格修改")
        self.assertEqual(app.product_workbench_spec(app.ProductWorkbench.PACKAGE_CONNECT).title, "套餐/API 对接修改")

        self.assertEqual(app.product_workbench_spec(app.ProductWorkbench.PRICE_INVENTORY).allowed_field_names, ("costPrice", "sale"))
        self.assertEqual(
            app.product_workbench_spec(app.ProductWorkbench.PACKAGE_CONNECT).allowed_field_names,
            ("vendorPackageId", "connectionCode", "introduction"),
        )

    def test_unified_product_edit_snapshot_needs_follow_target_fields(self):
        preload = app.product_snapshot_needs_for_workbench(app.ProductWorkbench.PRODUCT_INFO, [])
        name = app.product_snapshot_needs_for_workbench(
            app.ProductWorkbench.PRODUCT_INFO,
            [app.EditTemplateRow(2, "100000005", "", "", "", target_product_name="新产品名")],
        )
        package = app.product_snapshot_needs_for_workbench(
            app.ProductWorkbench.PRODUCT_INFO,
            [app.EditTemplateRow(2, "100000005", "1001", "2001", "套餐 A", target_vendor_package_id="V-2")],
        )
        price = app.product_snapshot_needs_for_workbench(
            app.ProductWorkbench.PRODUCT_INFO,
            [app.EditTemplateRow(2, "100000005", "1001", "2001", "套餐 A", target_cost_price="11")],
        )
        sale = app.product_snapshot_needs_for_workbench(
            app.ProductWorkbench.PRODUCT_INFO,
            [app.EditTemplateRow(2, "100000005", "1001", "2001", "套餐 A", target_sale_status="关")],
        )

        self.assertEqual(preload, app.ProductSnapshotNeeds(product_info=True, price_stock=True, currency_context=True))
        self.assertEqual(name, app.ProductSnapshotNeeds(product_info=True, price_stock=False, currency_context=False))
        self.assertEqual(package, app.ProductSnapshotNeeds(product_info=False, price_stock=False, currency_context=False))
        self.assertEqual(price, app.ProductSnapshotNeeds(product_info=False, price_stock=True, currency_context=True))
        self.assertEqual(sale, app.ProductSnapshotNeeds(product_info=False, price_stock=True, currency_context=True))


def make_snapshot(product_id="100000005"):
    draft = app.ProductDraft(
        product_id=product_id,
        target_id=product_id,
        package_info={
            "packageResourceInfo": [
                {
                    "packageId": "1001",
                    "packageName": "套餐 A",
                    "vendorPackageId": "V-1",
                    "introduction": [],
                    "resourceInfoList": [
                        {
                            "resourceId": "2001",
                            "active": True,
                            "connectionSystemKey": "C-1",
                            "connectionPropertyList": [{"key": "4", "value": "C-1"}],
                        }
                    ],
                }
            ]
        },
        package_resources=[
            {
                "packageId": "1001",
                "packageName": "套餐 A",
                "vendorPackageId": "V-1",
                "introduction": [],
                "resourceInfoList": [
                    {
                        "resourceId": "2001",
                        "active": True,
                        "connectionSystemKey": "C-1",
                        "connectionPropertyList": [{"key": "4", "value": "C-1"}],
                    }
                ],
            }
        ],
        product_version=7,
        draft_raw={"productBasicInfo": {"name": "旧产品名"}},
    )
    sku = app.SkuSnapshot(
        product_id=product_id,
        package_id="1001",
        resource_id="2001",
        package_name="套餐 A",
        vendor_package_id="V-1",
        connection_code="C-1",
        introduction="",
        cost_price="10",
        price_and_stock={"resourcePriceList": [{"resourceId": "2001", "priceList": [{"day": "00010101", "costPrice": 10}]}]},
    )
    return app.ProductSnapshot(
        product_id=product_id,
        created_at="2026-06-30T10:00:00",
        draft=draft,
        sku_snapshots={app.sku_key(product_id, "1001", "2001"): sku},
    )


class ProductWorkbenchTemplateTests(unittest.TestCase):
    def test_product_info_template_exports_all_product_edit_fields(self):
        snapshot = make_snapshot()
        next(iter(snapshot.sku_snapshots.values())).sale_status = "开"
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "product_info.xlsx"
            app.build_edit_template_workbook(
                str(path),
                {"100000005": snapshot},
                include_price=False,
                include_connection=False,
                include_intro=False,
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            )
            wb = load_workbook(path)
            headers = [cell.value for cell in wb[app.EDIT_TEMPLATE_SHEET][1]]
            self.assertIn("target_product_name", headers)
            self.assertIn("target_cost_price", headers)
            self.assertIn("old_sale_status", headers)
            self.assertIn("target_sale_status", headers)
            self.assertIn("target_vendor_package_id", headers)
            self.assertIn("target_connection_code", headers)
            self.assertIn("target_introduction", headers)
            self.assertNotIn("usedup_throttle", headers)
            self.assertNotIn("reset_time_calc", headers)
            self.assertNotIn("inst_adv", headers)
            self.assertNotIn(app.INTRO_FILL_SHEET, wb.sheetnames)
            ws = wb[app.EDIT_TEMPLATE_SHEET]
            self.assertEqual(ws.cell(2, headers.index("old_sale_status") + 1).value, "开")
            validations = list(ws.data_validations.dataValidation)
            self.assertTrue(any(validation.formula1 == '"开,关"' for validation in validations))

    def test_product_info_template_reads_sale_target_and_builds_diff(self):
        snapshot = make_snapshot()
        sku = next(iter(snapshot.sku_snapshots.values()))
        sku.sale_status = "开"
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "product_info.xlsx"
            app.build_edit_template_workbook(
                str(path),
                {snapshot.product_id: snapshot},
                include_price=False,
                include_connection=False,
                include_intro=False,
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            )
            wb = load_workbook(path)
            ws = wb[app.EDIT_TEMPLATE_SHEET]
            headers = [cell.value for cell in ws[1]]
            ws.cell(2, headers.index("target_sale_status") + 1, "关")
            wb.save(path)

            rows = app.read_edit_template(str(path), app.ProductWorkbench.PRODUCT_INFO)
            plans = app.build_diff_plans_from_rows(
                {snapshot.product_id: snapshot}, rows, app.ProductWorkbench.PRODUCT_INFO
            )

        self.assertEqual(len(plans), 1)
        self.assertEqual(plans[0].field_name, "sale")
        self.assertEqual(plans[0].old_value, "开")
        self.assertEqual(plans[0].new_value, "关")

    def test_price_template_exports_current_sale_status_and_open_close_dropdown(self):
        snapshot = make_snapshot()
        sku = next(iter(snapshot.sku_snapshots.values()))
        sku.sale_status = "关"
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "price.xlsx"
            app.build_edit_template_workbook(
                str(path),
                {snapshot.product_id: snapshot},
                include_price=True,
                include_connection=False,
                include_intro=False,
                workbench=app.ProductWorkbench.PRICE_INVENTORY,
            )
            wb = load_workbook(path)
            ws = wb[app.EDIT_TEMPLATE_SHEET]
            headers = [cell.value for cell in ws[1]]

            self.assertIn("old_sale_status", headers)
            self.assertIn("target_sale_status", headers)
            self.assertEqual(ws.cell(2, headers.index("old_sale_status") + 1).value, "关")
            validations = list(ws.data_validations.dataValidation)
            self.assertTrue(any(validation.formula1 == '"开,关"' for validation in validations))

    def test_price_template_exports_cost_price_permission_message(self):
        snapshot = make_snapshot()
        sku = next(iter(snapshot.sku_snapshots.values()))
        sku.cost_price = app.COST_PRICE_NOT_EDITABLE_MESSAGE
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "price.xlsx"
            app.build_edit_template_workbook(
                str(path),
                {snapshot.product_id: snapshot},
                include_price=True,
                include_connection=False,
                include_intro=False,
                workbench=app.ProductWorkbench.PRICE_INVENTORY,
            )
            wb = load_workbook(path)
            ws = wb[app.EDIT_TEMPLATE_SHEET]
            headers = [cell.value for cell in ws[1]]

            self.assertEqual(
                ws.cell(2, headers.index("old_cost_price") + 1).value,
                "请联系业务经理开放录入底价",
            )

    def test_sale_status_template_builds_price_stock_diff(self):
        snapshot = make_snapshot()
        sku = next(iter(snapshot.sku_snapshots.values()))
        sku.sale_status = "开"
        row = app.EditTemplateRow(
            2,
            snapshot.product_id,
            sku.package_id,
            sku.resource_id,
            sku.package_name,
            target_sale_status="关",
            workbench=app.ProductWorkbench.PRICE_INVENTORY,
        )

        plans = app.build_diff_plans_from_rows(
            {snapshot.product_id: snapshot},
            [row],
            app.ProductWorkbench.PRICE_INVENTORY,
        )

        self.assertEqual(len(plans), 1)
        self.assertEqual(plans[0].field_name, "sale")
        self.assertEqual(plans[0].old_value, "开")
        self.assertEqual(plans[0].new_value, "关")
        self.assertEqual(plans[0].interface_name, "saveOrEffectProductDraft(PriceAndStock)")

    def test_price_stock_helpers_read_permission_and_sale_status(self):
        price_and_stock = {
            "basicInfoList": [{"resourceId": "2001", "isAllowVendorEditCostPrice": False}],
            "resourcePriceList": [
                {"resourceId": "2001", "priceList": [{"day": "00010101", "costPrice": 10, "sale": False}]}
            ],
        }

        self.assertFalse(app.vendor_can_edit_cost_price(price_and_stock, "2001"))
        self.assertEqual(app.current_sale_status(price_and_stock, "2001"), "关")

        package_price_stock = {"packageId": "1001", "price": {"resourcePriceList": price_and_stock["resourcePriceList"]}}
        app.set_package_sale_status_draft(package_price_stock, "2001", "开")
        self.assertTrue(package_price_stock["price"]["resourcePriceDraftList"][0]["priceList"][0]["sale"])

    def test_sale_status_reads_draft_first_and_treats_empty_price_as_closed(self):
        open_price = {
            "resourcePriceList": [{"resourceId": "2001", "priceList": []}],
            "resourcePriceDraftList": [{"resourceId": "2001", "priceList": [{"sale": True}]}],
        }
        closed_price = {"resourcePriceList": [{"resourceId": "2001", "priceList": []}]}

        self.assertEqual(app.current_sale_status(open_price, "2001"), "开")
        self.assertEqual(app.current_sale_status(closed_price, "2001"), "关")

    def test_product_info_template_mixed_fields_build_split_diff_plans(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "product_info.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = app.EDIT_TEMPLATE_SHEET
            ws.append([
                "productId",
                "packageId",
                "resourceId",
                "packageName",
                "target_product_name",
                "target_cost_price",
                "target_vendor_package_id",
                "target_connection_code",
                "target_introduction",
            ])
            ws.append(["100000005", "1001", "2001", "套餐 A", "新产品名", "11", "V-2", "C-2", "新套餐说明"])
            wb.save(path)

            rows = app.read_edit_template(str(path), app.ProductWorkbench.PRODUCT_INFO)
            plans = app.build_diff_plans_from_rows(
                {"100000005": make_snapshot()},
                rows,
                app.ProductWorkbench.PRODUCT_INFO,
            )

        self.assertEqual([plan.field_name for plan in plans], ["productName", "costPrice", "vendorPackageId", "connectionCode", "introduction"])
        self.assertEqual(
            [plan.interface_name for plan in plans],
            [
                "saveOrEffectProductDraft(ProductInfo)",
                "saveOrEffectProductDraft(PriceAndStock)",
                "saveOrEffectProductDraft(type=2)",
                "saveOrEffectProductDraft(type=2)",
                "saveOrEffectProductDraft(type=2)",
            ],
        )
        self.assertTrue(all(plan.workbench == app.ProductWorkbench.PRODUCT_INFO for plan in plans))

    def test_legacy_price_template_reads_from_unified_product_edit_entry(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "price.xlsx"
            app.build_edit_template_workbook(
                str(path),
                {"100000005": make_snapshot()},
                include_price=True,
                include_connection=False,
                include_intro=False,
                workbench=app.ProductWorkbench.PRICE_INVENTORY,
            )
            wb = load_workbook(path)
            ws = wb[app.EDIT_TEMPLATE_SHEET]
            headers = [cell.value for cell in ws[1]]
            ws.cell(2, headers.index("target_cost_price") + 1, "11")
            wb.save(path)
            rows = app.read_edit_template(str(path), app.ProductWorkbench.PRODUCT_INFO)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].product_id, "100000005")
        self.assertEqual(rows[0].package_id, "1001")
        self.assertEqual(rows[0].resource_id, "2001")
        self.assertEqual(rows[0].target_cost_price, "11")

    def test_unified_product_edit_accepts_mixed_template_row(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "mixed.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = app.EDIT_TEMPLATE_SHEET
            ws.append(["productId", "packageId", "resourceId", "packageName", "target_cost_price", "target_vendor_package_id"])
            ws.append(["100000005", "1001", "2001", "套餐 A", "11", "V-2"])
            wb.save(path)

            rows = app.read_edit_template(str(path), app.ProductWorkbench.PRODUCT_INFO)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].target_cost_price, "11")
        self.assertEqual(rows[0].target_vendor_package_id, "V-2")

    def test_price_template_ignores_blank_target_rows_before_sku_lookup(self):
        row = app.EditTemplateRow(
            row_number=2,
            product_id="100000005",
            package_id="missing-package",
            resource_id="missing-resource",
            package_name="模板里只用于定位的套餐",
            target_cost_price="",
            workbench=app.ProductWorkbench.PRICE_INVENTORY,
        )

        plans = app.build_diff_plans_from_rows(
            {"100000005": make_snapshot()},
            [row],
            app.ProductWorkbench.PRICE_INVENTORY,
        )

        self.assertEqual([], plans)


class ProductInfoPayloadTests(unittest.TestCase):
    def test_build_product_info_name_payload_preserves_required_sections(self):
        product_info = {
            "productId": 100000005,
            "productVersion": 8,
            "productBasicInfo": {"name": "旧产品名", "productId": 100000005},
            "productDescInfo": {"addInfoList": []},
            "productRecommendActivityInfo": {"activityList": []},
            "productAdditionSaleInfo": {"items": []},
            "productPriceStock": {},
            "productExtendInfo": {},
        }

        payload = app.build_product_info_name_payload("100000005", product_info, "新产品名")

        self.assertEqual(payload["productId"], 100000005)
        self.assertEqual(payload["type"], 2)
        self.assertEqual(payload["draftTypes"], ["ProductInfo"])
        self.assertEqual(payload["draftTypesRequired"], ["ProductInfo"])
        self.assertEqual(payload["productBasicInfo"]["name"], "新产品名")
        self.assertEqual(payload["productDescInfo"], {"addInfoList": []})
        self.assertEqual(payload["productRecommendActivityInfo"], {"activityList": []})
        self.assertEqual(payload["productAdditionSaleInfo"], {"items": []})

    def test_build_product_info_name_payload_rejects_missing_basic_info(self):
        with self.assertRaisesRegex(RuntimeError, "ProductInfo 中 productBasicInfo 不是对象"):
            app.build_product_info_name_payload("100000005", {"productVersion": 1, "productBasicInfo": None}, "新产品名")


class ProductInfoApplyFakeClient:
    def __init__(self):
        self.saved_payloads = []
        self.product_info = {
            "productId": 100000005,
            "productVersion": 8,
            "productBasicInfo": {"name": "旧产品名", "productId": 100000005},
            "productDescInfo": {},
            "productRecommendActivityInfo": {},
            "productAdditionSaleInfo": {},
        }

    def load_product_info_draft(self, product_id, _stop_event):
        return self.product_info

    def save_product_info_draft(self, product_id, payload, _stop_event):
        self.saved_payloads.append((product_id, payload))
        self.product_info = payload
        return {"responseBaseData": {"success": True}}


class ProductInfoPackageApplyFakeClient:
    def __init__(self):
        self.saved_package_infos = []
        self.draft = make_snapshot().draft
        self.draft.package_info["packageResourceInfo"][0]["introduction"] = ""
        self.draft.package_resources[0]["introduction"] = ""

    def load_product(self, product_id, _stop_event):
        return app.ProductDraft(
            product_id=product_id,
            target_id=product_id,
            package_info=app.copy.deepcopy(self.draft.package_info),
            package_resources=app.copy.deepcopy(self.draft.package_resources),
            product_version=self.draft.product_version,
            draft_raw=app.copy.deepcopy(self.draft.draft_raw),
        )

    def save_package_info(self, draft, package_info, _stop_event, currency_context=None):
        self.saved_package_infos.append(package_info)
        self.draft = app.ProductDraft(
            product_id=draft.product_id,
            target_id=draft.target_id,
            package_info=app.copy.deepcopy(package_info),
            package_resources=app.copy.deepcopy(package_info.get("packageResourceInfo") or []),
            product_version=draft.product_version,
            draft_raw=app.copy.deepcopy(draft.draft_raw),
        )
        return {"responseBaseData": {"success": True}}


class PriceApplyConfirmRetryFakeClient:
    def __init__(self):
        self.main_save_payloads = []
        self.price_stock_draft = {
            "productVersion": 7,
            "packagePriceAndStockList": [
                {
                    "packageId": "1001",
                    "price": {
                        "basicInfoList": [{"resourceId": "2001"}],
                        "resourcePriceList": [
                            {"resourceId": "2001", "priceList": [{"costPrice": 10, "salePrice": 12}]}
                        ],
                    },
                    "stock": {},
                }
            ],
        }

    def load_product_price_stock_draft(self, product_id, _stop_event, refresh=False):
        return app.copy.deepcopy(self.price_stock_draft)

    def load_product(self, product_id, _stop_event):
        return make_snapshot(product_id).draft

    def save_product_price_stock_draft(self, product_id, price_stock_draft, _stop_event, package_ids=None, currency_context=None):
        self.price_stock_draft = app.copy.deepcopy(price_stock_draft)
        return {"responseBaseData": {"success": True}, "getProductDraft": {"productVersion": 8}}

    def build_price_stock_save_payload(self, product_id, price_stock_draft, package_ids=None, currency_context=None):
        return app.CtripRequestsClient.build_price_stock_save_payload(self, product_id, price_stock_draft, package_ids, currency_context)

    def _request_json(self, _url, _product_id, payload, _stop_event, timeout=60):
        self.main_save_payloads.append(app.copy.deepcopy(payload))
        if len(self.main_save_payloads) == 1:
            return {
                "responseBaseData": {
                    "success": False,
                    "errorCode": "key.v.api.resource.currency.not.same",
                    "message": "份数卖价币种与其他资源维护不一致,请确认是否需要修改",
                }
            }
        return {"responseBaseData": {"success": True}}

    def get_resource_price_and_stock_direct(self, product_id, package_id, resource_id, _stop_event):
        return {
            "response": {
                "priceAndStock": {
                    "resourcePriceList": [
                        {"resourceId": resource_id, "priceList": [{"costPrice": 11, "salePrice": 13.2}]}
                    ]
                }
            }
        }


class ProductInfoApplyTests(unittest.TestCase):
    def test_apply_product_info_name_plan_saves_and_verifies(self):
        worker = app.ProductWorkflowWorker(
            out_queue=app.queue.Queue(),
            stop_event=threading.Event(),
            preloaded_products={"100000005": make_snapshot()},
            pending_diff_plans=[],
            last_saved_products=set(),
            client_factory=lambda: ProductInfoApplyFakeClient(),
        )
        fake_client = ProductInfoApplyFakeClient()
        plan = app.DiffPlan(
            row=app.EditTemplateRow(
                row_number=2,
                product_id="100000005",
                package_id="",
                resource_id="",
                package_name="",
                target_product_name="新产品名",
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            ),
            product_id="100000005",
            package_id="",
            resource_id="",
            package_name="",
            field_name="productName",
            old_value="旧产品名",
            new_value="新产品名",
            interface_name="saveOrEffectProductDraft(ProductInfo)",
            workbench=app.ProductWorkbench.PRODUCT_INFO,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            saved = worker._apply_product_info_plans("100000005", [plan], Path(tmpdir), fake_client, audit_writer=None)

        self.assertTrue(saved)
        self.assertEqual(fake_client.saved_payloads[0][1]["productBasicInfo"]["name"], "新产品名")
        self.assertEqual(plan.status, "saved")

    def test_product_info_package_fields_save_through_package_info(self):
        fake_client = ProductInfoPackageApplyFakeClient()
        plans = [
            app.DiffPlan(
                row=app.EditTemplateRow(
                    row_number=2,
                    product_id="100000005",
                    package_id="1001",
                    resource_id="2001",
                    package_name="套餐 A",
                    target_vendor_package_id="V-2",
                    workbench=app.ProductWorkbench.PRODUCT_INFO,
                ),
                product_id="100000005",
                package_id="1001",
                resource_id="2001",
                package_name="套餐 A",
                field_name="vendorPackageId",
                old_value="V-1",
                new_value="V-2",
                interface_name="saveOrEffectProductDraft(type=2)",
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            ),
            app.DiffPlan(
                row=app.EditTemplateRow(
                    row_number=2,
                    product_id="100000005",
                    package_id="1001",
                    resource_id="2001",
                    package_name="套餐 A",
                    target_connection_code="C-2",
                    workbench=app.ProductWorkbench.PRODUCT_INFO,
                ),
                product_id="100000005",
                package_id="1001",
                resource_id="2001",
                package_name="套餐 A",
                field_name="connectionCode",
                old_value="C-1",
                new_value="C-2",
                interface_name="saveOrEffectProductDraft(type=2)",
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            ),
            app.DiffPlan(
                row=app.EditTemplateRow(
                    row_number=2,
                    product_id="100000005",
                    package_id="1001",
                    resource_id="2001",
                    package_name="套餐 A",
                    target_introduction="新套餐说明",
                    workbench=app.ProductWorkbench.PRODUCT_INFO,
                ),
                product_id="100000005",
                package_id="1001",
                resource_id="2001",
                package_name="套餐 A",
                field_name="introduction",
                old_value="",
                new_value="新套餐说明",
                interface_name="saveOrEffectProductDraft(type=2)",
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            ),
        ]
        worker = app.ProductWorkflowWorker(
            out_queue=app.queue.Queue(),
            stop_event=threading.Event(),
            preloaded_products={"100000005": make_snapshot()},
            pending_diff_plans=plans,
            last_saved_products=set(),
            client_factory=lambda: fake_client,
            workbench=app.ProductWorkbench.PRODUCT_INFO,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            saved = worker._apply_product_plans("100000005", plans, Path(tmpdir), fake_client, audit_writer=None)

        self.assertTrue(saved)
        saved_resource = fake_client.saved_package_infos[0]["packageResourceInfo"][0]
        self.assertEqual(saved_resource["vendorPackageId"], "V-2")
        self.assertEqual(app.connection_code_from_package_resource(saved_resource), "C-2")
        self.assertEqual(app.ctrip_intro_to_text(saved_resource["introduction"]), "新套餐说明")
        self.assertTrue(all(plan.status == "saved" for plan in plans))

    def test_product_info_introduction_newlines_save_without_html_breaks(self):
        fake_client = ProductInfoPackageApplyFakeClient()
        plan = app.DiffPlan(
            row=app.EditTemplateRow(
                row_number=2,
                product_id="100000005",
                package_id="1001",
                resource_id="2001",
                package_name="套餐 A",
                target_introduction="第一行\n第二行",
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            ),
            product_id="100000005",
            package_id="1001",
            resource_id="2001",
            package_name="套餐 A",
            field_name="introduction",
            old_value="",
            new_value="第一行\n第二行",
            interface_name="saveOrEffectProductDraft(type=2)",
            workbench=app.ProductWorkbench.PRODUCT_INFO,
        )
        worker = app.ProductWorkflowWorker(
            out_queue=app.queue.Queue(),
            stop_event=threading.Event(),
            preloaded_products={"100000005": make_snapshot()},
            pending_diff_plans=[plan],
            last_saved_products=set(),
            client_factory=lambda: fake_client,
            workbench=app.ProductWorkbench.PRODUCT_INFO,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            saved = worker._apply_product_plans("100000005", [plan], Path(tmpdir), fake_client, audit_writer=None)

        saved_intro = fake_client.saved_package_infos[0]["packageResourceInfo"][0]["introduction"]
        self.assertEqual(saved_intro, "第一行\n第二行")
        self.assertNotIn("<br", saved_intro.lower())
        self.assertNotIn("</br", saved_intro.lower())
        self.assertTrue(saved)
        self.assertEqual(plan.status, "saved")


class ProductPriceApplyTests(unittest.TestCase):
    def test_price_effect_main_save_confirms_currency_warning_and_retries(self):
        fake_client = PriceApplyConfirmRetryFakeClient()
        row = app.EditTemplateRow(
            row_number=2,
            product_id="100000005",
            package_id="1001",
            resource_id="2001",
            package_name="套餐 A",
            target_cost_price="11",
            workbench=app.ProductWorkbench.PRICE_INVENTORY,
        )
        plan = app.DiffPlan(
            row=row,
            product_id="100000005",
            package_id="1001",
            resource_id="2001",
            package_name="套餐 A",
            field_name="costPrice",
            old_value="10",
            new_value="11",
            interface_name="saveOrEffectProductDraft(PriceAndStock)",
            workbench=app.ProductWorkbench.PRICE_INVENTORY,
        )
        worker = app.ProductWorkflowWorker(
            out_queue=app.queue.Queue(),
            stop_event=threading.Event(),
            preloaded_products={"100000005": make_snapshot()},
            pending_diff_plans=[plan],
            last_saved_products=set(),
            client_factory=lambda: fake_client,
            workbench=app.ProductWorkbench.PRICE_INVENTORY,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            saved = worker._apply_product_plans("100000005", [plan], Path(tmpdir), fake_client, audit_writer=None)

        self.assertTrue(saved)
        self.assertEqual(plan.status, "saved")
        self.assertEqual(len(fake_client.main_save_payloads), 2)
        self.assertEqual(fake_client.main_save_payloads[0]["effectMarks"], [3, 4])
        self.assertEqual(fake_client.main_save_payloads[1]["effectMarks"], [3])


class ProductWorkbenchWorkerRoutingTests(unittest.TestCase):
    def test_preview_template_uses_workbench_snapshot_needs(self):
        calls = []

        class FakeClient:
            def load_product(self, product_id, _stop_event):
                calls.append(("load_product", product_id))
                return make_snapshot(product_id).draft

            def load_product_price_stock(self, product_id, _stop_event):
                calls.append(("load_product_price_stock", product_id))
                return {}

            def load_product_info_draft(self, product_id, _stop_event):
                calls.append(("load_product_info_draft", product_id))
                return {"productId": int(product_id), "productBasicInfo": {"name": "旧产品名"}}

            def load_product_currency_context(self, product_id, _stop_event):
                calls.append(("load_product_currency_context", product_id))
                return app.CurrencyRewriteContext(product_id=product_id)

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "product_info.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = app.EDIT_TEMPLATE_SHEET
            ws.append(["productId", "old_product_name", "target_product_name"])
            ws.append(["100000005", "旧产品名", "新产品名"])
            wb.save(path)

            worker = app.ProductWorkflowWorker(
                out_queue=app.queue.Queue(),
                stop_event=threading.Event(),
                preloaded_products={},
                pending_diff_plans=[],
                last_saved_products=set(),
                client_factory=lambda: FakeClient(),
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            )
            worker.run_preview_template(str(path))

        self.assertIn(("load_product", "100000005"), calls)
        self.assertIn(("load_product_info_draft", "100000005"), calls)
        self.assertNotIn(("load_product_price_stock", "100000005"), calls)


class ProductWorkbenchUiTests(unittest.TestCase):
    def test_main_window_has_single_product_edit_tab(self):
        window = app.App()
        try:
            tabs = window.findChild(app.QTabWidget)
            labels = [tabs.tabText(index) for index in range(tabs.count())]
            self.assertEqual(labels[0], "翻译件管理")
            self.assertIn("产品信息接口修改", labels)
            self.assertNotIn("价格修改", labels)
            self.assertNotIn("产品信息修改", labels)
            self.assertNotIn("套餐/API 对接修改", labels)
            self.assertNotIn("审核与报告", labels)
            self.assertIn("产品图片配置", labels)
            self.assertEqual(window.active_product_workbench, app.ProductWorkbench.PRODUCT_INFO)
        finally:
            window.close()
            QT_APP.processEvents()

    def test_unified_product_workbench_state_is_active(self):
        window = app.App()
        try:
            state = window.current_product_workbench_state(app.ProductWorkbench.PRODUCT_INFO)
            state.results.append(app.ResultRow(product_id="1", package_name="", status="preloaded", workbench=app.ProductWorkbench.PRODUCT_INFO))
            window._activate_product_workbench(app.ProductWorkbench.PRODUCT_INFO)
            self.assertEqual(window.results, state.results)
            self.assertIn(app.ProductWorkbench.PRODUCT_INFO, window.product_workbench_widgets)
        finally:
            window.close()
            QT_APP.processEvents()

    def test_product_info_detail_text_includes_all_before_after_fields_and_diff(self):
        row = app.ResultRow(
            product_id="P100",
            package_id="PK200",
            resource_id="R300",
            package_name="日本 3日",
            field_name="产品信息",
            old_value="旧产品名",
            new_value="新产品名",
            old_vendor_package_id="old-vendor",
            new_vendor_package_id="new-vendor",
            old_connection_code="old-code",
            new_connection_code="new-code",
            old_introduction="旧说明第一行\n旧说明第二行",
            new_introduction="新说明第一行\n新说明第二行",
            interface_name="ProductInfo + PackageInfo",
            status="preview",
            message="等待传出",
            row_number=7,
            submit_audit=True,
            workbench=app.ProductWorkbench.PRODUCT_INFO,
        )
        window = app.App()
        try:
            detail = window.product_result_detail_text(row)
        finally:
            window.close()
            QT_APP.processEvents()

        for expected in [
            "productId: P100",
            "packageId: PK200",
            "resourceId: R300",
            "packageName: 日本 3日",
            "field: 产品信息",
            "oldValue: 旧产品名",
            "newValue: 新产品名",
            "oldVendorPackageId: old-vendor",
            "newVendorPackageId: new-vendor",
            "oldConnectionCode: old-code",
            "newConnectionCode: new-code",
            "submitAudit: 是",
            "status: preview",
            "row: 7",
            "message:",
            "等待传出",
            "oldIntroduction:",
            "旧说明第一行",
            "newIntroduction:",
            "新说明第一行",
            "diff:",
            "--- 修改前",
            "+++ 修改后",
            "-旧产品名",
            "+新产品名",
            "-旧说明第一行",
            "+新说明第一行",
        ]:
            self.assertIn(expected, detail)

    def test_product_info_detail_uses_unified_product_table(self):
        window = app.App()
        try:
            product_widgets = window.product_workbench_widgets[app.ProductWorkbench.PRODUCT_INFO]
            product_model = product_widgets["product_result_model"]
            product_detail = product_widgets["detail_text"]
            product_table = product_widgets["table"]

            product_model.set_results([
                app.ResultRow(
                    product_id="PRODUCT_INFO_TAB",
                    package_id="PK200",
                    package_name="产品资料",
                    field_name="产品信息",
                    old_value="旧产品名",
                    new_value="新产品名",
                    old_vendor_package_id="old-vendor",
                    new_vendor_package_id="new-vendor",
                    old_connection_code="old-code",
                    new_connection_code="new-code",
                    old_introduction="旧说明",
                    new_introduction="新说明",
                    status="preview",
                    workbench=app.ProductWorkbench.PRODUCT_INFO,
                )
            ])
            product_table.selectRow(0)

            window.render_product_result_detail(product_table, product_model, product_detail)

            self.assertIn("productId: PRODUCT_INFO_TAB", product_detail.toPlainText())
            self.assertIn("newVendorPackageId: new-vendor", product_detail.toPlainText())
            self.assertIn("+新说明", product_detail.toPlainText())
        finally:
            window.close()
            QT_APP.processEvents()

    def test_product_info_selection_signal_updates_own_detail_panel(self):
        window = app.App()
        try:
            widgets = window.product_workbench_widgets[app.ProductWorkbench.PRODUCT_INFO]
            model = widgets["product_result_model"]
            table = widgets["table"]
            detail_text = widgets["detail_text"]
            model.set_results([
                app.ResultRow(
                    product_id="P100",
                    package_name="产品资料",
                    field_name="产品信息",
                    old_value="旧产品名",
                    new_value="新产品名",
                    old_introduction="旧说明",
                    new_introduction="新说明",
                    status="preview",
                    workbench=app.ProductWorkbench.PRODUCT_INFO,
                )
            ])

            table.selectRow(0)
            QT_APP.processEvents()
            window.product_detail_timer.stop()
            window.render_product_result_detail(table, model, detail_text)

            self.assertIn("productId: P100", detail_text.toPlainText())
            self.assertIn("+新说明", detail_text.toPlainText())
        finally:
            window.close()
            QT_APP.processEvents()

    def test_open_product_result_preview_uses_selected_product_info_row(self):
        window = app.App()
        created = []

        class FakePreviewDialog:
            def __init__(self, snapshots, results, parent):
                created.append((snapshots, results, parent))

            def exec(self):
                return 0

        original_dialog = app.PreviewDialog
        app.PreviewDialog = FakePreviewDialog
        try:
            widgets = window.product_workbench_widgets[app.ProductWorkbench.PRODUCT_INFO]
            model = widgets["product_result_model"]
            table = widgets["table"]
            selected_row = app.ResultRow(
                product_id="P100",
                package_name="产品资料",
                field_name="产品信息",
                old_value="旧产品名",
                new_value="新产品名",
                status="preview",
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            )
            other_row = app.ResultRow(
                product_id="P200",
                package_name="其他资料",
                status="preview",
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            )
            model.set_results([selected_row, other_row])
            table.selectRow(0)

            window.open_product_result_preview(
                table,
                model,
                window.current_product_workbench_state(app.ProductWorkbench.PRODUCT_INFO).preloaded_products,
                model.rows,
            )

            self.assertEqual(1, len(created))
            self.assertEqual([selected_row], created[0][1])
            self.assertIs(window, created[0][2])
        finally:
            app.PreviewDialog = original_dialog
            window.close()
            QT_APP.processEvents()

    def test_unified_product_table_shows_detail_and_previews_selected_row(self):
        window = app.App()
        created = []

        class FakePreviewDialog:
            def __init__(self, snapshots, results, parent):
                created.append((snapshots, results, parent))

            def exec(self):
                return 0

        original_dialog = app.PreviewDialog
        app.PreviewDialog = FakePreviewDialog
        try:
            widgets = window.product_workbench_widgets[app.ProductWorkbench.PRODUCT_INFO]
            table = widgets["table"]
            model = widgets["product_result_model"]
            detail_text = widgets["detail_text"]
            row = app.ResultRow(
                product_id="P100",
                package_id="PK200",
                resource_id="R300",
                package_name="产品编辑套餐",
                field_name="产品信息",
                old_value="旧内容",
                new_value="新内容",
                old_cost_price="10",
                new_cost_price="12",
                old_vendor_package_id="old-vendor",
                new_vendor_package_id="new-vendor",
                old_connection_code="old-code",
                new_connection_code="new-code",
                old_introduction="旧说明",
                new_introduction="新说明",
                interface_name="saveOrEffectProductDraft",
                status="preview",
                message="等待传出",
                row_number=8,
                submit_audit=True,
                workbench=app.ProductWorkbench.PRODUCT_INFO,
            )
            model.set_results([row])
            table.selectRow(0)

            window.render_product_result_detail(table, model, detail_text)
            text = detail_text.toPlainText()

            self.assertIn("productId: P100", text)
            self.assertIn("oldValue: 旧内容", text)
            self.assertIn("newValue: 新内容", text)
            self.assertIn("oldCostPrice: 10", text)
            self.assertIn("newCostPrice: 12", text)
            self.assertIn("oldVendorPackageId: old-vendor", text)
            self.assertIn("newVendorPackageId: new-vendor", text)
            self.assertIn("diff:", text)
            self.assertIn("-旧内容", text)
            self.assertIn("+新内容", text)

            window.open_product_result_preview(table, model, {}, model.rows)

            self.assertEqual(1, len(created))
            self.assertEqual(["P100"], [item.product_id for item in created[0][1]])
        finally:
            app.PreviewDialog = original_dialog
            window.close()
            QT_APP.processEvents()


class ProductWorkbenchReportTests(unittest.TestCase):
    def test_report_headers_include_workbench(self):
        headers = app.report_headers(app.ProductWorkbench.PRICE_INVENTORY)
        self.assertEqual(headers[0], "工作台")
        self.assertIn("产品ID", headers)
        self.assertIn("调用接口", headers)

    def test_result_row_values_start_with_workbench_title(self):
        row = app.ResultRow(
            product_id="100000005",
            package_name="套餐 A",
            status="saved",
            workbench=app.ProductWorkbench.PRICE_INVENTORY,
        )
        values = app.result_row_values(row)
        self.assertEqual(values[0], "价格修改")
