import copy
import json
import os
import queue
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest import mock
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PyQt6.QtGui import QPalette
import requests

import app
import translation_rules

QT_APP = app.QApplication.instance() or app.QApplication([])

PRODUCT_ID = "100000005"


class ProductSnapshotFakeClient:
    def __init__(self):
        self.calls = []

    def load_product(self, product_id, _stop_event):
        self.calls.append(("load_product", product_id))
        return app.ProductDraft(
            product_id=product_id,
            target_id=product_id,
            package_info={"packageResourceInfo": []},
            package_resources=[
                {
                    "packageId": "1001",
                    "packageName": "套餐 A",
                    "vendorPackageId": "V-1",
                    "introduction": [],
                    "resourceInfoList": [{"resourceId": "2001", "active": True}],
                }
            ],
            product_version=1,
            draft_raw={},
        )

    def load_product_info_draft(self, product_id, _stop_event):
        self.calls.append(("load_product_info_draft", product_id))
        return {"productBasicInfo": {"productName": "name"}}

    def load_product_currency_context(self, product_id, _stop_event):
        self.calls.append(("load_product_currency_context", product_id))
        return app.CurrencyRewriteContext(product_id=product_id)

    def load_product_price_stock(self, product_id, _stop_event):
        self.calls.append(("load_product_price_stock", product_id))
        return {}


class TranslationApiResponseTests(unittest.TestCase):
    def test_available_locales_business_error_raises_clear_message(self):
        client = app.CtripRequestsClient.__new__(app.CtripRequestsClient)
        client.log = lambda _message: None

        def fake_request(_url, _product_id, _body, _stop_event, timeout=60):
            return {"responseBaseData": {"success": False, "errorCode": "LOGIN_ERROR", "message": "无权限"}}

        client._request_json = fake_request

        with self.assertRaisesRegex(RuntimeError, "读取翻译件可用语言失败.*无权限.*LOGIN_ERROR"):
            client.load_translation_available_locales(PRODUCT_ID, threading.Event())

    def test_translation_list_business_error_raises_clear_message(self):
        client = app.CtripRequestsClient.__new__(app.CtripRequestsClient)
        client.log = lambda _message: None

        def fake_request(_url, _product_id, _body, _stop_event, timeout=60):
            return {"responseBaseData": {"success": False, "errorCode": "LOGIN_ERROR", "message": "无权限"}}

        client._request_json = fake_request

        with self.assertRaisesRegex(RuntimeError, "读取翻译件内容失败.*无权限.*LOGIN_ERROR"):
            client.load_translation_list(PRODUCT_ID, ["ar-XX"], threading.Event())


class TranslationWorkerLimitTests(unittest.TestCase):
    def test_translation_runtime_settings_clamp_parallel_to_five(self):
        settings = app.RuntimeRequestSettings(request_interval_seconds=0.1, parallel_requests=99)
        snapshot = settings.snapshot()
        self.assertEqual(snapshot.request_interval_seconds, 0.2)
        self.assertEqual(snapshot.parallel_requests, 5)


class AppStartupTabTests(unittest.TestCase):
    def test_app_import_survives_missing_pyqt_without_importing_ui_theme(self):
        script = r'''
import builtins
import sys

real_import = builtins.__import__
def blocked_import(name, *args, **kwargs):
    if name == "PyQt6" or name.startswith("PyQt6."):
        raise ImportError("PyQt6 intentionally blocked")
    return real_import(name, *args, **kwargs)

builtins.__import__ = blocked_import
sys.modules.pop("app", None)
sys.modules.pop("ui_theme", None)
import app
assert app.QT_IMPORT_ERROR is not None
assert app.ui_theme is None
assert "ui_theme" not in sys.modules
'''
        result = subprocess.run(
            [sys.executable, "-c", script],
            cwd=Path(app.__file__).parent,
            capture_output=True,
            text=True,
            timeout=20,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_translation_manual_link_opens_bundled_manual_with_system_default(self):
        window = app.App()
        try:
            link = window.translation_rule_manual_link
            self.assertIn("操作手册", link.text())
            self.assertIs(link.parent(), window.translation_rule_detail_manager_btn.parent())
            manual = app.translation_ops_manual_path()
            self.assertTrue(manual.is_file(), manual)
            with mock.patch.object(app, "open_with_system_default") as opener:
                link.linkActivated.emit("translation-manual")
            opener.assert_called_once_with(str(manual))
        finally:
            window.close()
            QT_APP.processEvents()

    def test_translation_config_exposes_operation_attention_as_default_module(self):
        window = app.App()
        window.translation_rule_runtime.config = app.default_translation_execution_config()
        captured = {}

        def inspect_and_reject_dialog():
            dialog = next(
                widget
                for widget in QT_APP.topLevelWidgets()
                if isinstance(widget, app.QDialog)
                and widget.objectName() == "TranslationRuleConfigDialog"
            )
            checkboxes = [
                item
                for item in dialog.findChildren(app.QCheckBox)
                if item.text() == "携程处理注意事项"
            ]
            captured["checkboxes"] = checkboxes
            if checkboxes:
                captured["checked"] = checkboxes[0].isChecked()
                captured["tooltip"] = checkboxes[0].toolTip()
            dialog.reject()

        try:
            app.QTimer.singleShot(0, inspect_and_reject_dialog)
            window.open_translation_rule_config_dialog()
            self.assertEqual(len(captured["checkboxes"]), 1)
            self.assertTrue(captured["checked"])
            self.assertIn("operationAttention", captured["tooltip"])
        finally:
            window.close()
            QT_APP.processEvents()

    def test_coordinator_apply_refreshes_once_and_external_palette_refreshes_once(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        original = qt_app.palette()
        native = QPalette(original)
        coordinator = app.ui_theme.SystemThemeCoordinator(
            qt_app,
            fluent_enabled=False,
            palette_source=lambda: QPalette(native),
        )
        window = app.App()
        window._theme_coordinator = coordinator
        calls = []
        original_apply = window._apply_global_theme

        def counted_apply():
            calls.append(True)
            original_apply()

        coordinator.themeApplied.connect(counted_apply)
        try:
            coordinator.apply_theme()
            qt_app.processEvents()
            self.assertEqual(len(calls), 1)

            calls.clear()
            external = QPalette(native)
            external.setColor(external.ColorRole.Window, app.QColor("#dcdcdc"))
            native = QPalette(external)
            qt_app.setPalette(external)
            qt_app.processEvents()
            qt_app.processEvents()
            self.assertEqual(len(calls), 1)
        finally:
            coordinator.stop()
            window.close()
            qt_app.setPalette(original)


    def test_translation_config_dialog_theme_uses_palette_selection_roles(self):
        dialog = app.QDialog()
        dialog.setObjectName("TranslationRuleConfigDialog")
        palette = dialog.palette()
        palette.setColor(palette.ColorRole.Window, app.QColor("#202124"))
        palette.setColor(palette.ColorRole.Base, app.QColor("#292a2d"))
        palette.setColor(palette.ColorRole.Text, app.QColor("#f1f3f4"))
        palette.setColor(palette.ColorRole.Highlight, app.QColor("#4c8bf5"))
        palette.setColor(palette.ColorRole.HighlightedText, app.QColor("#ffffff"))
        dialog.setPalette(palette)
        app.apply_translation_config_dialog_theme(dialog)
        stylesheet = dialog.styleSheet().lower()
        self.assertIn("selection-background-color: #4c8bf5", stylesheet)
        self.assertIn("selection-color: #ffffff", stylesheet)

    def test_translation_config_dialog_refresh_preserves_controls_and_releases_closed_dialog(self):
        qt_app = app.QApplication.instance() or app.QApplication([])

        for result in (app.QDialog.DialogCode.Rejected, app.QDialog.DialogCode.Accepted):
            with self.subTest(result=result):
                window = app.App()
                captured_dialogs = []

                def fake_exec(dialog):
                    captured_dialogs.append(dialog)
                    dialog.show()
                    qt_app.processEvents()
                    language_check = next(
                        checkbox
                        for checkbox in dialog.findChildren(app.QCheckBox)
                        if checkbox.text().startswith("JA ")
                    )
                    module_check = next(
                        checkbox
                        for checkbox in dialog.findChildren(app.QCheckBox)
                        if checkbox.text() == "产品名称"
                    )
                    online_check = next(
                        checkbox
                        for checkbox in dialog.findChildren(app.QCheckBox)
                        if checkbox.text().startswith("保存并上线")
                    )
                    sales_radio = next(
                        radio
                        for radio in dialog.findChildren(app.QRadioButton)
                        if radio.text() == "请求机器翻译"
                    )
                    long_day_combo = dialog.findChild(app.QComboBox)
                    self.assertIsNotNone(long_day_combo)

                    for checkbox in dialog.findChildren(app.QCheckBox):
                        checkbox.setChecked(False)
                    language_check.setChecked(True)
                    module_check.setChecked(True)
                    sales_radio.setChecked(True)
                    long_day_combo.setCurrentText("套用后降速模板")
                    online_check.setChecked(False)
                    state_before = (
                        language_check.isChecked(),
                        module_check.isChecked(),
                        sales_radio.isChecked(),
                        long_day_combo.currentText(),
                        online_check.isChecked(),
                    )
                    palette = dialog.palette()
                    palette.setColor(palette.ColorRole.Window, app.QColor("#202124"))
                    palette.setColor(palette.ColorRole.Base, app.QColor("#292a2d"))
                    palette.setColor(palette.ColorRole.Text, app.QColor("#f1f3f4"))
                    palette.setColor(palette.ColorRole.Highlight, app.QColor("#4c8bf5"))
                    palette.setColor(palette.ColorRole.HighlightedText, app.QColor("#ffffff"))
                    dialog.setPalette(palette)
                    window._refresh_open_dialog_themes()
                    self.assertEqual(
                        state_before,
                        (
                            language_check.isChecked(),
                            module_check.isChecked(),
                            sales_radio.isChecked(),
                            long_day_combo.currentText(),
                            online_check.isChecked(),
                        ),
                    )
                    self.assertIn("#4c8bf5", dialog.styleSheet().lower())
                    self.assertFalse(dialog.testAttribute(app.Qt.WidgetAttribute.WA_DeleteOnClose))
                    if result == app.QDialog.DialogCode.Accepted:
                        dialog.accept()
                    else:
                        dialog.reject()
                    return result

                try:
                    with (
                        mock.patch.object(app.QDialog, "exec", fake_exec),
                        mock.patch.object(window, "persist_app_settings"),
                    ):
                        window.open_translation_rule_config_dialog()
                    qt_app.processEvents()
                    qt_app.processEvents()
                    self.assertFalse(hasattr(window, "_translation_rule_long_day_combo"))
                    self.assertFalse(hasattr(window, "_translation_rule_online_check"))
                    self.assertEqual(1, len(captured_dialogs))
                    self.assertNotIn(captured_dialogs[0], qt_app.topLevelWidgets())
                    if result == app.QDialog.DialogCode.Accepted:
                        self.assertEqual(["ja"], window.translation_rule_runtime.config.selected_languages)
                        self.assertEqual(["product_name"], window.translation_rule_runtime.config.selected_modules)
                        self.assertEqual("ai_translation", window.translation_rule_runtime.config.sales_attr_mode)
                        self.assertEqual("throttle", window.translation_rule_runtime.config.long_day_package_template_mode)
                        self.assertFalse(window.translation_rule_runtime.config.online)
                finally:
                    window.close()
                    qt_app.processEvents()
                    qt_app.processEvents()

    def test_translation_config_accept_reads_controls_before_deferred_deletion(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        destroyed = []

        def accept_dialog():
            dialog = next(
                widget
                for widget in qt_app.topLevelWidgets()
                if isinstance(widget, app.QDialog)
                and widget.objectName() == "TranslationRuleConfigDialog"
            )
            dialog.destroyed.connect(lambda: destroyed.append(True))
            language_check = next(
                checkbox
                for checkbox in dialog.findChildren(app.QCheckBox)
                if checkbox.text().startswith("JA ")
            )
            module_check = next(
                checkbox
                for checkbox in dialog.findChildren(app.QCheckBox)
                if checkbox.text() == "产品名称"
            )
            for checkbox in dialog.findChildren(app.QCheckBox):
                checkbox.setChecked(False)
            language_check.setChecked(True)
            module_check.setChecked(True)
            dialog.accept()

        try:
            app.QTimer.singleShot(0, accept_dialog)
            with mock.patch.object(window, "persist_app_settings"):
                window.open_translation_rule_config_dialog()
            self.assertEqual(["ja"], window.translation_rule_runtime.config.selected_languages)
            self.assertEqual(["product_name"], window.translation_rule_runtime.config.selected_modules)
            self.assertFalse(hasattr(window, "_translation_rule_long_day_combo"))
            self.assertFalse(hasattr(window, "_translation_rule_online_check"))
            qt_app.processEvents()
            qt_app.processEvents()
            self.assertEqual([True], destroyed)
            self.assertFalse(
                any(
                    isinstance(widget, app.QDialog)
                    and widget.objectName() == "TranslationRuleConfigDialog"
                    for widget in qt_app.topLevelWidgets()
                )
            )
        finally:
            window.close()
            qt_app.processEvents()
            qt_app.processEvents()

    def test_global_theme_refresh_updates_all_tables_without_losing_state(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        original_palette = QPalette(qt_app.palette())
        try:
            inactive_tab = window._build_product_workbench_operational_tab(app.ProductWorkbench.PRICE_INVENTORY)
            inactive_widgets = window.product_workbench_widgets[app.ProductWorkbench.PRICE_INVENTORY]
            window.template_path.setText("product.xlsx")
            window.translation_runtime_status_label.setText("运行中｜本次 9/1")
            window.translation_log_text.setPlainText("翻译日志")
            window.product_image_workspace.product_ids_edit.setPlainText("100000001")
            inactive_widgets["log_text"].setPlainText("未激活价格日志")
            inactive_widgets["detail_text"].setPlainText("未激活价格详情")
            active_workbench = window.active_product_workbench
            active_aliases = (window.table, window.log_text, window.detail_text, window.status_label)
            workbench_models = {
                key: widgets["table"].model()
                for key, widgets in window.product_workbench_widgets.items()
            }
            workbench_status_styles = {
                key: widgets["status_label"].styleSheet()
                for key, widgets in window.product_workbench_widgets.items()
            }
            translation_model = window.translation_table.model()
            palette = QPalette(original_palette)
            palette.setColor(palette.ColorRole.Window, app.QColor("#202124"))
            palette.setColor(palette.ColorRole.Base, app.QColor("#292a2d"))
            palette.setColor(palette.ColorRole.Text, app.QColor("#f1f3f4"))
            palette.setColor(palette.ColorRole.WindowText, app.QColor("#f1f3f4"))
            palette.setColor(palette.ColorRole.Highlight, app.QColor("#4c8bf5"))
            palette.setColor(palette.ColorRole.HighlightedText, app.QColor("#ffffff"))
            palette.setColor(palette.ColorRole.Link, app.QColor("#4c8bf5"))
            qt_app.setPalette(palette)
            qt_app.processEvents()
            window._apply_global_theme()

            self.assertEqual(window.template_path.text(), "product.xlsx")
            self.assertEqual(window.translation_runtime_status_label.text(), "运行中｜本次 9/1")
            self.assertEqual(window.translation_log_text.toPlainText(), "翻译日志")
            self.assertEqual(window.product_image_workspace.product_ids_edit.toPlainText(), "100000001")
            self.assertEqual(inactive_widgets["log_text"].toPlainText(), "未激活价格日志")
            self.assertEqual(inactive_widgets["detail_text"].toPlainText(), "未激活价格详情")
            self.assertEqual(window.active_product_workbench, active_workbench)
            self.assertEqual(active_aliases, (window.table, window.log_text, window.detail_text, window.status_label))
            self.assertEqual(
                workbench_models,
                {
                    key: widgets["table"].model()
                    for key, widgets in window.product_workbench_widgets.items()
                },
            )
            self.assertIs(translation_model, window.translation_table.model())
            for key, widgets in window.product_workbench_widgets.items():
                self.assertIn("#292a2d", widgets["table"].styleSheet().lower())
                self.assertEqual(widgets["status_label"].palette().windowText().color().name(), "#4c8bf5")
                self.assertEqual(widgets["status_label"].styleSheet(), workbench_status_styles[key])
                self.assertNotIn("color:", widgets["status_label"].styleSheet().lower())
                self.assertIn("font-size: 13px", widgets["status_label"].styleSheet())
                self.assertIn("font-weight: 600", widgets["status_label"].styleSheet())
            self.assertIn("#292a2d", window.translation_table.styleSheet().lower())
        finally:
            qt_app.setPalette(original_palette)
            qt_app.processEvents()
            window.close()

    def test_translation_tab_is_first_and_default_area(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            tabs = window.findChild(app.QTabWidget)
            self.assertIsNotNone(tabs)
            self.assertEqual(tabs.tabText(0), "翻译件管理")
            self.assertEqual(tabs.currentIndex(), 0)
            self.assertEqual(window.active_worker_area, "translation")
            titles = [tabs.tabText(index) for index in range(tabs.count())]
            self.assertIn("产品信息接口修改", titles)
            self.assertNotIn("价格修改", titles)
            self.assertNotIn("产品信息修改", titles)
            self.assertNotIn("套餐/API 对接修改", titles)
            self.assertNotIn("审核与报告", titles)

            tab = tabs.widget(titles.index("产品信息接口修改"))
            button_texts = {button.text() for button in tab.findChildren(app.QPushButton)}
            self.assertTrue(any("输入产品ID" in text for text in button_texts), button_texts)
            self.assertTrue(any("导出" in text and "模板" in text for text in button_texts), button_texts)
            self.assertTrue(any("加载" in text and "预览" in text for text in button_texts), button_texts)
            self.assertTrue(any("传出" in text for text in button_texts), button_texts)
            self.assertIn("提交审核", button_texts)
            self.assertTrue(any("报告" in text for text in button_texts), button_texts)
            self.assertIn("刷新登录态", button_texts)
            self.assertIn("停止", button_texts)
        finally:
            window.close()

    def test_product_edit_tab_exposes_staged_workflow_layout(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            tabs = window.findChild(app.QTabWidget)
            titles = [tabs.tabText(index) for index in range(tabs.count())]
            tab = tabs.widget(titles.index("产品信息接口修改"))

            self.assertIsNotNone(tab.findChild(app.QFrame, "productWorkflowHeader"))
            self.assertIsNotNone(tab.findChild(app.QFrame, "productTemplateCard"))
            self.assertIsNotNone(tab.findChild(app.QFrame, "productExecutionCard"))
            self.assertIsNotNone(tab.findChild(app.QFrame, "productResultsCard"))
            self.assertEqual(window.template_path.placeholderText(), "选择已回填的 .xlsx 修改模板")
            self.assertEqual(window.preload_btn.objectName(), "productPrimaryAction")
            self.assertEqual(window.apply_btn.objectName(), "productDangerAction")
            self.assertEqual(window.stop_btn.objectName(), "productStopAction")
            self.assertEqual(window.status_label.objectName(), "productStatusLabel")
            self.assertEqual(tab.styleSheet(), "")
            for frame_name in (
                "productWorkflowHeader",
                "productTemplateCard",
                "productExecutionCard",
                "productResultsCard",
            ):
                frame = tab.findChild(app.QFrame, frame_name)
                self.assertEqual(frame.frameShape(), app.QFrame.Shape.StyledPanel)
            self.assertIsInstance(window.apply_btn, app.PrimaryPushButton)
            self.assertNotIn("color", window.status_label.styleSheet())
        finally:
            window.close()

    def test_auth_refresh_is_the_top_action_with_status_in_both_main_workbenches(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            window.show()
            qt_app.processEvents()

            product_header = window.findChild(app.QFrame, "productWorkflowHeader")
            translation_header = window.findChild(app.QFrame, "translationWorkflowHeader")
            self.assertTrue(product_header.isAncestorOf(window.auth_btn))
            self.assertTrue(product_header.isAncestorOf(window.product_auth_status_label))
            self.assertTrue(translation_header.isAncestorOf(window.translation_rule_auth_btn))
            self.assertTrue(translation_header.isAncestorOf(window.translation_auth_status_label))
            product_layout = product_header.parentWidget().layout()
            self.assertLess(product_layout.indexOf(product_header), product_layout.indexOf(window.findChild(app.QFrame, "productTemplateCard")))
            translation_layout = window.translation_workbench_tab.layout()
            self.assertLess(
                translation_layout.indexOf(translation_header),
                translation_layout.indexOf(window.findChild(app.QSplitter, "translationMainSplitter")),
            )
        finally:
            window.close()

    def test_transmit_changes_uses_the_native_primary_button_component(self):
        window = app.App()
        try:
            self.assertIsInstance(window.apply_btn, app.PrimaryPushButton)
            self.assertEqual(window.apply_btn.metaObject().className(), "PrimaryPushButton")
        finally:
            window.close()

    def test_auth_check_result_updates_status_beside_both_refresh_buttons(self):
        window = app.App()
        result = app.AuthCheckResult(
            status="valid",
            cookie_digest=app.cookie_digest("current-cookie"),
        )
        try:
            with mock.patch.object(app, "read_plain_cookie", return_value="current-cookie"):
                window._handle_startup_auth_result(result)

            self.assertEqual(window.product_auth_status_label.text(), "登录态：有效")
            self.assertEqual(window.translation_auth_status_label.text(), "登录态：有效")
        finally:
            window.close()

    def test_table_style_uses_the_active_system_palette(self):
        table = app.QTableView()
        palette = table.palette()
        palette.setColor(palette.ColorRole.Base, app.QColor("#112233"))
        palette.setColor(palette.ColorRole.AlternateBase, app.QColor("#223344"))
        palette.setColor(palette.ColorRole.Text, app.QColor("#ddeeff"))
        palette.setColor(palette.ColorRole.Highlight, app.QColor("#336699"))
        table.setPalette(palette)

        app.apply_table_readable_style(table)

        stylesheet = table.styleSheet().lower()
        self.assertIn("background-color: #112233", stylesheet)
        self.assertIn("alternate-background-color: #223344", stylesheet)
        self.assertIn("color: #ddeeff", stylesheet)
        self.assertIn("selection-background-color: #336699", stylesheet)

    def test_translation_workbench_palette_is_derived_from_active_palette(self):
        widget = app.QWidget()
        palette = widget.palette()
        palette.setColor(palette.ColorRole.Window, app.QColor("#101820"))
        palette.setColor(palette.ColorRole.Base, app.QColor("#18242e"))
        palette.setColor(palette.ColorRole.Button, app.QColor("#22313d"))
        palette.setColor(palette.ColorRole.Text, app.QColor("#edf6fb"))
        palette.setColor(palette.ColorRole.WindowText, app.QColor("#edf6fb"))
        palette.setColor(palette.ColorRole.PlaceholderText, app.QColor("#9aabb6"))
        palette.setColor(palette.ColorRole.Mid, app.QColor("#596a76"))
        palette.setColor(palette.ColorRole.Highlight, app.QColor("#35a7ff"))
        palette.setColor(palette.ColorRole.HighlightedText, app.QColor("#07131c"))
        widget.setPalette(palette)

        colors = app.translation_workbench_palette(widget)

        self.assertEqual(colors["window"], "#101820")
        self.assertEqual(colors["text"], "#edf6fb")
        self.assertEqual(colors["muted"], "#9aabb6")
        self.assertEqual(colors["border"], "#596a76")
        self.assertEqual(colors["accent"], "#35a7ff")
        self.assertEqual(colors["accent_text"], "#07131c")
        self.assertNotIn(colors["card"], {"#ffffff", "#000000"})

    def test_mix_qcolors_clamps_weight(self):
        first = app.QColor("#000000")
        second = app.QColor("#ffffff")

        self.assertEqual(app.mix_qcolors(first, second, -1).name(), "#000000")
        self.assertEqual(app.mix_qcolors(first, second, 2).name(), "#ffffff")
        self.assertEqual(app.mix_qcolors(first, second, 0.5).name(), "#808080")

    def test_translation_tab_exposes_compact_sidebar_layout(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            tab = window.translation_workbench_tab
            main_splitter = tab.findChild(app.QSplitter, "translationMainSplitter")
            sidebar = tab.findChild(app.QWidget, "translationSidebar")
            results_pane = tab.findChild(app.QFrame, "translationResultsPane")
            results_splitter = tab.findChild(app.QSplitter, "translationResultsSplitter")
            log_detail_splitter = tab.findChild(app.QSplitter, "translationLogDetailSplitter")

            self.assertIsNotNone(main_splitter)
            self.assertEqual(main_splitter.orientation(), app.Qt.Orientation.Horizontal)
            self.assertEqual(main_splitter.count(), 2)
            self.assertIs(main_splitter.widget(0), sidebar)
            self.assertIs(main_splitter.widget(1), results_pane)
            self.assertEqual(sidebar.findChildren(app.QSplitter), [])

            self.assertIsNotNone(results_splitter)
            self.assertEqual(results_splitter.orientation(), app.Qt.Orientation.Vertical)
            self.assertEqual(results_splitter.count(), 2)
            self.assertIsNotNone(log_detail_splitter)
            self.assertEqual(log_detail_splitter.orientation(), app.Qt.Orientation.Horizontal)
            self.assertEqual(log_detail_splitter.count(), 2)

            for label in (
                window.translation_runtime_summary_label,
                window.translation_runtime_status_label,
                window.translation_rule_task_status_label,
                window.translation_rule_detail_status_label,
                window.translation_rule_package_status_label,
            ):
                self.assertTrue(label.wordWrap())
                self.assertGreater(label.maximumWidth(), 1000)

            button_texts = {button.text() for button in tab.findChildren(app.QPushButton)}
            expected_buttons = {
                "加载任务模板", "下载任务模板", "编辑任务模板", "重新加载任务",
                "加载详情", "管理详情", "下载详情",
                "加载规格", "管理规格", "下载规格",
                "执行配置", "生成Diff", "确认传出", "导出翻译报告",
                "刷新登录态", "停止", "前一表格", "后一表格", "重置程序",
            }
            self.assertTrue(expected_buttons.issubset(button_texts))

            window.resize(1100, 700)
            window.show()
            qt_app.processEvents()
            self.assertGreater(results_pane.width(), sidebar.width())
            self.assertGreater(window.translation_table.height(), log_detail_splitter.height())
        finally:
            window.close()

    def test_translation_compact_layout_preserves_full_status_content(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            summary = "语言 2｜模块 1｜销售属性 strict｜保存并上线"
            runtime = "就绪｜请求 累计 123/4｜今日 56/2｜本次 7/1"
            task = "已加载：示例任务模板.xlsx"
            detail = "已加载：示例详情模板.xlsx"
            package = "已加载：示例规格模板.xlsx"

            window.translation_runtime_summary_label.setText(summary)
            window.translation_runtime_status_label.setText(runtime)
            window.translation_rule_task_status_label.setText(task)
            window.translation_rule_detail_status_label.setText(detail)
            window.translation_rule_package_status_label.setText(package)
            window._apply_translation_workbench_theme()
            qt_app.processEvents()

            self.assertEqual(window.translation_runtime_summary_label.text(), summary)
            self.assertEqual(window.translation_runtime_status_label.text(), runtime)
            self.assertEqual(window.translation_rule_task_status_label.text(), task)
            self.assertEqual(window.translation_rule_detail_status_label.text(), detail)
            self.assertEqual(window.translation_rule_package_status_label.text(), package)
            self.assertEqual(window.translation_rule_task_status_label.toolTip(), task)
            self.assertEqual(window.translation_rule_detail_status_label.toolTip(), detail)
            self.assertEqual(window.translation_rule_package_status_label.toolTip(), package)
        finally:
            window.close()

    def test_translation_splitters_resize_without_losing_runtime_state(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            window.show()
            qt_app.processEvents()
            main = window.translation_workbench_tab.findChild(app.QSplitter, "translationMainSplitter")
            vertical = window.translation_workbench_tab.findChild(app.QSplitter, "translationResultsSplitter")
            window.translation_runtime_status_label.setText("运行中：本次 9/1")
            window.translation_log_text.setPlainText("保留日志")
            main.setSizes([460, 900])
            vertical.setSizes([420, 240])
            qt_app.processEvents()
            self.assertEqual(window.translation_runtime_status_label.text(), "运行中：本次 9/1")
            self.assertEqual(window.translation_log_text.toPlainText(), "保留日志")
            self.assertGreater(main.sizes()[0], 340)
            self.assertGreater(vertical.sizes()[0], vertical.sizes()[1])

            main.setSizes([0, 1000])
            qt_app.processEvents()
            self.assertGreaterEqual(main.sizes()[0], 340)

            main.setSizes([1000, 0])
            qt_app.processEvents()
            self.assertGreaterEqual(main.sizes()[1], 480)
        finally:
            window.close()

    def test_translation_sidebar_controls_stay_usable_inside_supported_viewports(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            tab = window.translation_workbench_tab
            sidebar = tab.findChild(app.QWidget, "translationSidebar")
            sections = [
                tab.findChild(app.QFrame, name)
                for name in (
                    "translationTaskSection",
                    "translationDetailSection",
                    "translationPackageSection",
                    "translationPreviewSection",
                    "translationDeliverySection",
                )
            ]
            status_labels = (
                window.translation_rule_task_status_label,
                window.translation_rule_detail_status_label,
                window.translation_rule_package_status_label,
            )
            window.translation_rule_task_status_label.setText(
                "已加载：示例任务模板.xlsx"
            )
            window.translation_rule_detail_status_label.setText(
                "已加载：示例详情模板.xlsx"
            )
            window.translation_rule_package_status_label.setText(
                "已加载：示例规格模板.xlsx"
            )

            for width, height in ((980, 620), (1100, 700), (1700, 900)):
                window.resize(width, height)
                window.show()
                qt_app.processEvents()

                for section in sections:
                    self.assertTrue(sidebar.rect().contains(section.geometry()), (width, height, section.objectName()))
                    for button in section.findChildren(app.QPushButton):
                        self.assertTrue(button.isVisible(), (width, height, button.text()))
                        self.assertGreaterEqual(button.height(), 24, (width, height, button.text()))
                        self.assertTrue(section.rect().contains(button.geometry()), (width, height, button.text()))
                for label in status_labels:
                    section = label.parentWidget()
                    self.assertTrue(label.isVisible(), (width, height, label.text()))
                    self.assertTrue(section.rect().contains(label.geometry()), (width, height, label.text()))
                    self.assertGreaterEqual(
                        label.height(),
                        label.heightForWidth(label.width()),
                        (width, height, label.text()),
                    )
        finally:
            window.close()

    def test_translation_template_status_colors_survive_light_and_dark_theme_refresh(self):
        window = app.App()
        try:
            loaded_label = window.translation_rule_task_status_label
            missing_label = window.translation_rule_detail_status_label
            window.set_translation_rule_status_label(loaded_label, "/tmp/loaded-task.xlsx", "已加载", "未加载任务模板")
            window.set_translation_rule_status_label(missing_label, "", "已加载", "未加载详情模板")
            loaded_text = loaded_label.text()
            missing_text = missing_label.text()

            for window_color, base_color, text_color, muted_color, accent_color in (
                ("#f2f3f5", "#ffffff", "#1c2024", "#7a838c", "#087f78"),
                ("#202225", "#2b2e32", "#f2f4f7", "#98a2ad", "#35a7ff"),
            ):
                palette = window.translation_workbench_tab.palette()
                palette.setColor(palette.ColorRole.Window, app.QColor(window_color))
                palette.setColor(palette.ColorRole.Base, app.QColor(base_color))
                palette.setColor(palette.ColorRole.Button, app.QColor(base_color))
                palette.setColor(palette.ColorRole.Text, app.QColor(text_color))
                palette.setColor(palette.ColorRole.WindowText, app.QColor(text_color))
                palette.setColor(palette.ColorRole.PlaceholderText, app.QColor(muted_color))
                palette.setColor(palette.ColorRole.Mid, app.QColor("#69737d"))
                palette.setColor(palette.ColorRole.Highlight, app.QColor(accent_color))
                palette.setColor(palette.ColorRole.HighlightedText, app.QColor("#07131c"))
                palette.setColor(palette.ColorRole.Link, app.QColor(accent_color))
                window.translation_workbench_tab.setPalette(palette)

                window._apply_translation_workbench_theme()

                self.assertIn(accent_color, loaded_label.styleSheet().lower())
                self.assertIn(muted_color, missing_label.styleSheet().lower())
                self.assertNotEqual(loaded_label.styleSheet(), missing_label.styleSheet())
                self.assertEqual(loaded_label.text(), loaded_text)
                self.assertEqual(missing_label.text(), missing_text)
                self.assertEqual(loaded_label.toolTip(), loaded_text)
                self.assertEqual(missing_label.toolTip(), missing_text)
        finally:
            window.close()

    def test_translation_theme_uses_custom_light_and_dark_palettes(self):
        window = app.App()
        try:
            tab = window.translation_workbench_tab
            for window_color, base_color, text_color, accent_color in (
                ("#f2f3f5", "#ffffff", "#1c2024", "#087f78"),
                ("#202225", "#2b2e32", "#f2f4f7", "#35a7ff"),
            ):
                palette = tab.palette()
                palette.setColor(palette.ColorRole.Window, app.QColor(window_color))
                palette.setColor(palette.ColorRole.Base, app.QColor(base_color))
                palette.setColor(palette.ColorRole.Button, app.QColor(base_color))
                palette.setColor(palette.ColorRole.Text, app.QColor(text_color))
                palette.setColor(palette.ColorRole.WindowText, app.QColor(text_color))
                palette.setColor(palette.ColorRole.PlaceholderText, app.QColor("#8a949e"))
                palette.setColor(palette.ColorRole.Mid, app.QColor("#69737d"))
                palette.setColor(palette.ColorRole.Highlight, app.QColor(accent_color))
                palette.setColor(palette.ColorRole.HighlightedText, app.QColor("#07131c"))
                tab.setPalette(palette)

                window._apply_translation_workbench_theme()

                stylesheet = tab.styleSheet().lower()
                self.assertIn(accent_color, stylesheet)
                self.assertIn(text_color, stylesheet)
                if base_color != "#ffffff":
                    self.assertNotIn("background-color: #ffffff", stylesheet)
                    self.assertIn(
                        f"background-color: {base_color}",
                        window.translation_table.styleSheet().lower(),
                    )
                    self.assertIn(
                        f"color: {text_color}",
                        window.translation_table.styleSheet().lower(),
                    )
                    self.assertIn("#07131c", window.translation_rule_generate_btn.styleSheet().lower())
                    self.assertIn(text_color, window.translation_rule_apply_btn.styleSheet().lower())
                    self.assertIn(text_color, window.translation_request_interval_edit.styleSheet().lower())
                    self.assertIn(text_color, window.translation_log_text.styleSheet().lower())
                    self.assertIn(text_color, window.translation_detail_text.styleSheet().lower())
        finally:
            window.close()

    def test_translation_palette_refresh_preserves_runtime_content(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        try:
            window.translation_log_text.setPlainText("保留日志")
            window.translation_detail_text.setPlainText("保留 Diff")
            before_model = window.translation_table.model()

            window.changeEvent(app.QEvent(app.QEvent.Type.PaletteChange))
            qt_app.processEvents()

            self.assertEqual(window.translation_log_text.toPlainText(), "保留日志")
            self.assertEqual(window.translation_detail_text.toPlainText(), "保留 Diff")
            self.assertIs(window.translation_table.model(), before_model)
        finally:
            window.close()

    def test_translation_detail_text_includes_old_new_diff_and_parameters(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        row = app.TranslationResultRow(
            product_id="P100",
            locale="ja-JP",
            module_code="DETAIL",
            field_code="description",
            field_name="详情",
            self_key="detail.description",
            status="preview",
            old_value="旧翻译第一行\n旧翻译第二行",
            new_value="新翻译第一行\n新翻译第二行",
            online=False,
            interface_name="productTranslationPush",
            message="等待传出",
            row_number=9,
        )
        window = app.App()
        try:
            detail = window.translation_result_detail_text(row)
        finally:
            window.close()
            qt_app.processEvents()

        for expected in [
            "productId: P100",
            "locale: ja-JP",
            "moduleCode: DETAIL",
            "fieldCode: description",
            "fieldName: 详情",
            "selfKey: detail.description",
            "online: 否",
            "interface: productTranslationPush",
            "status: preview",
            "row: 9",
            "message:",
            "等待传出",
            "diff:",
            "--- 当前翻译",
            "+++ 目标翻译",
            "-旧翻译第一行",
            "+新翻译第一行",
            "oldValue:",
            "newValue:",
        ]:
            self.assertIn(expected, detail)

    def test_translation_double_click_preview_uses_selected_row(self):
        qt_app = app.QApplication.instance() or app.QApplication([])
        window = app.App()
        created = []

        class FakeTranslationPreviewDialog:
            def __init__(self, snapshots, results, parent):
                created.append((snapshots, results, parent))

            def exec(self):
                return 0

        original_dialog = app.TranslationPreviewDialog
        app.TranslationPreviewDialog = FakeTranslationPreviewDialog
        try:
            selected_row = app.TranslationResultRow(
                product_id="P100",
                locale="ja-JP",
                module_code="DETAIL",
                field_code="description",
                field_name="详情",
                self_key="detail.description",
                status="preview",
                old_value="旧翻译",
                new_value="新翻译",
                row_number=9,
            )
            other_row = app.TranslationResultRow(
                product_id="P200",
                locale="ko-KR",
                module_code="DETAIL",
                field_code="description",
                field_name="详情",
                self_key="detail.description",
                status="preview",
            )
            window.translation_result_model.set_results([selected_row, other_row])
            window.translation_table.selectRow(0)

            window.render_translation_result_detail(window.translation_table, window.translation_result_model, window.translation_detail_text)
            self.assertIn("productId: P100", window.translation_detail_text.toPlainText())
            self.assertIn("+新翻译", window.translation_detail_text.toPlainText())

            window.open_translation_result_preview(
                window.translation_table,
                window.translation_result_model,
                window.translation_snapshots,
                window.translation_result_model.rows,
            )

            self.assertEqual(1, len(created))
            self.assertEqual([selected_row], created[0][1])
            self.assertIs(window, created[0][2])
        finally:
            app.TranslationPreviewDialog = original_dialog
            window.close()
            qt_app.processEvents()



class FileDialogOptionTests(unittest.TestCase):
    def test_mac_excel_file_dialogs_use_qt_dialog_to_keep_xlsx_selectable(self):
        if app.QT_IMPORT_ERROR is not None:
            self.skipTest(f"PyQt6 unavailable: {app.QT_IMPORT_ERROR}")

        with mock.patch.object(app.platform, "system", return_value="Darwin"):
            options = app.file_dialog_options_for_excel()

        self.assertTrue(options & app.QFileDialog.Option.DontUseNativeDialog)

    def test_non_mac_excel_file_dialogs_keep_default_native_dialog(self):
        if app.QT_IMPORT_ERROR is not None:
            self.skipTest(f"PyQt6 unavailable: {app.QT_IMPORT_ERROR}")

        with mock.patch.object(app.platform, "system", return_value="Windows"):
            options = app.file_dialog_options_for_excel()

        self.assertFalse(options & app.QFileDialog.Option.DontUseNativeDialog)


class ReportChineseHeaderTests(unittest.TestCase):
    def test_product_and_translation_report_headers_are_chinese(self):
        self.assertEqual(
            app.report_headers(),
            [
                "工作台",
                "产品ID",
                "套餐ID",
                "资源ID",
                "套餐名称",
                "修改字段",
                "修改前内容",
                "修改后内容",
                "调用接口",
                "用完处理模板",
                "重置时间规则",
                "安装说明模板",
                "原底价",
                "新底价",
                "原供应商套餐ID",
                "新供应商套餐ID",
                "原对接代码",
                "新对接代码",
                "是否提交审核",
                "处理状态",
                "原套餐说明",
                "新套餐说明",
                "结果说明",
                "模板行号",
            ],
        )
        self.assertEqual(
            app.translation_report_headers(),
            [
                "产品ID",
                "语言地区",
                "模块编码",
                "字段编码",
                "字段名称",
                "字段键",
                "修改前译文",
                "修改后译文",
                "是否上线",
                "调用接口",
                "处理状态",
                "结果说明",
                "模板行号",
            ],
        )

    def test_main_gui_result_table_headers_are_chinese(self):
        self.assertEqual(
            app.PRODUCT_RESULT_HEADERS,
            ["产品ID", "套餐ID", "资源ID", "套餐名称", "修改字段", "修改前内容", "修改后内容", "调用接口", "处理状态", "结果说明", "模板行号"],
        )
        self.assertEqual(
            app.TRANSLATION_RESULT_HEADERS,
            ["产品ID", "语言地区", "模块", "翻译字段", "修改前译文", "修改后译文", "是否上线", "调用接口", "处理状态", "结果说明", "模板行号"],
        )

    def test_preview_snapshot_headers_are_chinese(self):
        product_snapshot = app.ProductSnapshot(
            product_id=PRODUCT_ID,
            created_at="2026-06-30T10:00:00",
            draft=app.ProductDraft(product_id=PRODUCT_ID, target_id=PRODUCT_ID, package_info={}, package_resources=[]),
            sku_snapshots={
                "sku-1": app.SkuSnapshot(
                    product_id=PRODUCT_ID,
                    package_id="pkg-1",
                    resource_id="res-1",
                    package_name="套餐A",
                    cost_price="10",
                    vendor_package_id="vendor-1",
                    connection_code="code-1",
                    introduction="套餐说明",
                    price_error="",
                )
            },
        )
        translation_snapshot = app.TranslationSnapshot(
            product_id=PRODUCT_ID,
            created_at="2026-06-30T10:00:00",
            locales=["en-US"],
            raw_by_locale={},
            fields={
                "field-1": app.TranslationFieldSnapshot(
                    product_id=PRODUCT_ID,
                    locale="en-US",
                    module_code="basic",
                    module_name="基础信息",
                    field_code="productName",
                    field_name="产品名称",
                    value_key="value",
                    self_key="self-1",
                    path="basic.productName",
                    origin="原文",
                    old_translation="Old name",
                )
            },
        )

        product_headers, product_rows = app.snapshot_preview_rows({PRODUCT_ID: product_snapshot})
        translation_headers, translation_rows = app.translation_snapshot_preview_rows({PRODUCT_ID: translation_snapshot})

        self.assertEqual(product_headers, ["产品ID", "套餐ID", "资源ID", "套餐名称", "当前底价", "当前供应商套餐ID", "当前对接代码", "当前套餐说明", "价格读取状态", "快照时间"])
        self.assertEqual(translation_headers, ["产品ID", "语言地区", "模块编码", "模块名称", "字段编码", "字段名称", "字段键", "原文内容", "当前译文", "快照时间"])
        self.assertEqual(len(product_headers), len(product_rows[0]))
        self.assertEqual(len(translation_headers), len(translation_rows[0]))

    def test_xlsx_reports_write_chinese_headers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            product_path = Path(tmpdir) / "product_report.xlsx"
            translation_path = Path(tmpdir) / "translation_report.xlsx"
            product_results = [
                app.ResultRow(
                    product_id=PRODUCT_ID,
                    package_id="pkg-1",
                    resource_id="res-1",
                    package_name="套餐A",
                    field_name="套餐说明",
                    old_value="旧值",
                    new_value="新值",
                    interface_name="saveOrEffectProductDraft(type=2)",
                    status="saved",
                    message="完成",
                    row_number=2,
                )
            ]
            translation_results = [
                app.TranslationResultRow(
                    product_id=PRODUCT_ID,
                    locale="en-US",
                    module_code="basic",
                    field_code="productName",
                    field_name="产品名称",
                    self_key="self-1",
                    old_value="Old",
                    new_value="New",
                    status="saved",
                    message="完成",
                    row_number=2,
                )
            ]

            app.write_report(str(product_path), product_results)
            app.write_translation_report(str(translation_path), translation_results)

            product_ws = load_workbook(product_path, read_only=True)["结果"]
            translation_ws = load_workbook(translation_path, read_only=True)["翻译结果"]
            product_headers = [product_ws.cell(1, col).value for col in range(1, product_ws.max_column + 1)]
            translation_headers = [translation_ws.cell(1, col).value for col in range(1, translation_ws.max_column + 1)]

            self.assertEqual(product_headers, app.report_headers())
            self.assertEqual(translation_headers, app.translation_report_headers())


class AuthCookieRefreshTests(unittest.TestCase):
    def test_default_plain_auth_paths_use_system_config_dir(self):
        self.assertIn("CtripPlanDescriptionFilling", str(app.PLAIN_AUTH_DIR))
        self.assertEqual(app.PLAIN_COOKIE_FILE.parent.name, "auth")
        self.assertFalse(hasattr(app, "DEFAULT_TRIP_SESSION_FILE"))
        self.assertFalse(hasattr(app, "LEGACY_PLAIN_COOKIE_FILE"))

    def test_windows_plain_auth_path_uses_appdata(self):
        with mock.patch.object(app.platform, "system", return_value="Windows"), \
             mock.patch.dict(app.os.environ, {"APPDATA": r"C:\Users\operator\AppData\Roaming"}):
            auth_dir = app.default_app_auth_dir()

        normalized = str(auth_dir).replace("\\", "/")
        self.assertEqual(normalized, "C:/Users/operator/AppData/Roaming/CtripPlanDescriptionFilling/auth")

    def test_read_plain_cookie_reads_only_requested_system_file(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            system_path = Path(tmpdir) / "system" / "ctrip_cookie_header.txt"
            system_path.parent.mkdir()
            system_path.write_text("system_cookie=1", encoding="utf-8")

            self.assertEqual(app.read_plain_cookie(system_path), "system_cookie=1")

    def test_read_plain_cookie_does_not_accept_legacy_fallback(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            system_path = Path(tmpdir) / "system" / "ctrip_cookie_header.txt"
            legacy_path = Path(tmpdir) / "legacy" / "ctrip_cookie_header.txt"
            legacy_path.parent.mkdir()
            legacy_path.write_text("legacy_cookie=1", encoding="utf-8")

            with self.assertRaises(TypeError):
                app.read_plain_cookie(system_path, legacy_path)

    def test_startup_auth_check_reports_missing_without_validation(self):
        checker = getattr(app, "check_saved_plain_auth", None)
        self.assertIsNotNone(checker)
        with tempfile.TemporaryDirectory() as tmpdir:
            validator = mock.Mock()
            result = checker(Path(tmpdir) / "missing.txt", validator=validator)

        self.assertEqual(result.status, "missing")
        validator.assert_not_called()

    def test_startup_auth_check_treats_empty_file_as_missing(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            cookie_path.write_text("  \n", encoding="utf-8")
            validator = mock.Mock()
            result = app.check_saved_plain_auth(cookie_path, validator=validator)

        self.assertEqual(result.status, "missing")
        validator.assert_not_called()

    def test_startup_auth_check_accepts_saved_login_cookie_without_product_probe(self):
        checker = getattr(app, "check_saved_plain_auth", None)
        self.assertIsNotNone(checker)
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            cookie_path.write_text("_bfa=tracking; vbkticket=valid", encoding="utf-8")
            validator = mock.Mock()
            result = checker(cookie_path, validator=validator)

        self.assertEqual(result.status, "valid")
        self.assertTrue(result.cookie_digest)
        validator.assert_not_called()

    def test_startup_auth_check_rejects_tracking_cookies_without_login_cookie(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            cookie_path.write_text("_bfa=tracking; GUID=anonymous", encoding="utf-8")

            result = app.check_saved_plain_auth(cookie_path)

        self.assertEqual(result.status, "invalid")

    def test_startup_auth_check_classifies_login_error_as_invalid(self):
        checker = getattr(app, "check_saved_plain_auth", None)
        self.assertIsNotNone(checker)
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            cookie_path.write_text("vbkticket=expired", encoding="utf-8")
            validator = mock.Mock(side_effect=RuntimeError("登录态校验失败：无权限 | LOGIN_ERROR"))
            result = checker(cookie_path, product_id="7654321", validator=validator)

        self.assertEqual(result.status, "invalid")

    def test_startup_auth_check_keeps_permission_and_network_errors_distinct(self):
        checker = getattr(app, "check_saved_plain_auth", None)
        self.assertIsNotNone(checker)
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            cookie_path.write_text("vbkticket=unknown", encoding="utf-8")

            permission_result = checker(
                cookie_path,
                product_id="7654321",
                validator=mock.Mock(side_effect=RuntimeError("产品无权限")),
            )
            network_result = checker(
                cookie_path,
                product_id="7654321",
                validator=mock.Mock(side_effect=requests.Timeout("连接超时")),
            )

        self.assertEqual(permission_result.status, "error")
        self.assertEqual(network_result.status, "error")

    def test_auth_refresh_without_loaded_product_skips_product_probe(self):
        window = app.App()
        try:
            with mock.patch.object(window, "_auth_validation_product_id", return_value=""), \
                 mock.patch.object(app, "refresh_plain_auth", return_value=(Path("cookie.txt"), 1, "current_browser_cdp")) as refresh:
                window._refresh_auth_for_area("product")
        finally:
            window.close()

        refresh.assert_called_once_with(
            validation_product_id="",
            area="product",
        )

    def test_translation_refresh_without_loaded_product_skips_product_probe(self):
        window = app.App()
        try:
            with mock.patch.object(window, "_auth_validation_product_id", return_value=""), \
                 mock.patch.object(app, "refresh_plain_auth", return_value=(Path("cookie.txt"), 1, "current_browser_cdp")) as refresh:
                window._refresh_auth_for_area("translation")
        finally:
            window.close()

        refresh.assert_called_once_with(
            validation_product_id="",
            area="product",
        )

    def test_missing_auth_does_not_automatically_refresh_browser(self):
        window = app.App()
        window.active_worker_area = "product"
        try:
            with mock.patch.object(app, "read_plain_cookie", return_value=""), \
                 mock.patch.object(app, "refresh_plain_auth") as refresh, \
                 mock.patch.object(app.QMessageBox, "warning"):
                ready = window._ensure_auth_ready("7654321", "product")
        finally:
            window.close()

        self.assertFalse(ready)
        refresh.assert_not_called()

    def test_startup_auth_result_ignores_stale_cookie_and_prompts_current_failure(self):
        window = app.App()
        stale = app.AuthCheckResult(status="invalid", message="LOGIN_ERROR", cookie_digest="old")
        current = app.AuthCheckResult(
            status="invalid",
            message="LOGIN_ERROR",
            cookie_digest=app.cookie_digest("current-cookie"),
        )
        try:
            with mock.patch.object(app, "read_plain_cookie", return_value="current-cookie"), \
                 mock.patch.object(app.QMessageBox, "warning") as warning:
                window._handle_startup_auth_result(stale)
                warning.assert_not_called()
                window._handle_startup_auth_result(current)
        finally:
            window.close()

        warning.assert_called_once()

    def test_successful_refresh_suppresses_matching_startup_failure(self):
        window = app.App()
        current_digest = app.cookie_digest("current-cookie")
        window.last_validated_auth_digest = current_digest
        result = app.AuthCheckResult(status="error", message="连接超时", cookie_digest=current_digest)
        try:
            with mock.patch.object(app, "read_plain_cookie", return_value="current-cookie"), \
                 mock.patch.object(app.QMessageBox, "warning") as warning:
                window._handle_startup_auth_result(result)
        finally:
            window.close()

        warning.assert_not_called()

    def test_startup_auth_check_posts_background_result(self):
        starter = getattr(app.App, "start_startup_auth_check", None)
        self.assertIsNotNone(starter)
        window = app.App()

        class ImmediateThread:
            def __init__(self, target, daemon):
                self.target = target

            def is_alive(self):
                return False

            def start(self):
                self.target()

        expected = app.AuthCheckResult(status="missing")
        try:
            with mock.patch.object(app, "check_saved_plain_auth", return_value=expected), \
                 mock.patch.object(app.threading, "Thread", ImmediateThread):
                window.start_startup_auth_check()
            self.assertEqual(window.out_queue.get_nowait(), ("startup_auth_check", expected))
        finally:
            window.close()

    def test_macos_browser_launches_use_separate_ports_and_profiles(self):
        with mock.patch.object(app.platform, "system", return_value="Darwin"), \
             mock.patch.object(app, "browser_is_installed", return_value=True, create=True):
            launches = app.browser_cdp_launches()

        self.assertEqual([launch.browser_name for launch in launches], ["Microsoft Edge", "Google Chrome"])
        self.assertEqual([getattr(launch, "port", None) for launch in launches], [9222, 9223])
        profile_args = [next(arg for arg in launch.command if arg.startswith("--user-data-dir=")) for launch in launches]
        self.assertNotEqual(profile_args[0], profile_args[1])

    def test_windows_browser_launches_include_edge_and_chrome(self):
        with mock.patch.object(app.platform, "system", return_value="Windows"), \
             mock.patch.object(app, "browser_is_installed", return_value=True, create=True):
            launches = app.browser_cdp_launches()

        self.assertEqual([launch.browser_name for launch in launches], ["Microsoft Edge", "Google Chrome"])
        self.assertEqual([getattr(launch, "port", None) for launch in launches], [9222, 9223])

    def test_windows_browser_launch_uses_discovered_executable_path(self):
        finder = getattr(app, "browser_executable_path", None)
        self.assertIsNotNone(finder)
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            edge_path = root / "Microsoft" / "Edge" / "Application" / "msedge.exe"
            edge_path.parent.mkdir(parents=True)
            edge_path.touch()
            with mock.patch.object(app.platform, "system", return_value="Windows"), \
                 mock.patch.object(app.shutil, "which", return_value=None), \
                 mock.patch.dict(app.os.environ, {"PROGRAMFILES": str(root), "PROGRAMFILES(X86)": "", "LOCALAPPDATA": ""}):
                self.assertEqual(finder("edge"), str(edge_path))
                launches = app.browser_cdp_launches()

            self.assertEqual([launch.browser_key for launch in launches], ["edge"])
            self.assertEqual(launches[0].command[0], str(edge_path))

    def test_existing_cdp_discovery_identifies_only_edge_and_chrome(self):
        versions = {
            9222: {"Browser": "Edg/126.0"},
            9223: {"Browser": "Chrome/126.0"},
            9224: {"Browser": "Unrelated/1.0"},
        }
        with mock.patch.object(app, "browser_cdp_version", side_effect=lambda port: versions.get(port)):
            sessions = app.discover_existing_browser_cdp_sessions(range(9222, 9225))

        self.assertEqual([(session.browser_key, session.port) for session in sessions], [("edge", 9222), ("chrome", 9223)])

    def test_browser_launches_skip_uninstalled_edge_and_keep_chrome(self):
        with mock.patch.object(app.platform, "system", return_value="Darwin"), \
             mock.patch.object(app, "browser_is_installed", side_effect=lambda browser_key: browser_key == "chrome", create=True):
            launches = app.browser_cdp_launches()

        self.assertEqual([launch.browser_name for launch in launches], ["Google Chrome"])
        profile_arg = next(arg for arg in launches[0].command if arg.startswith("--user-data-dir="))
        self.assertEqual(Path(profile_arg.split("=", 1)[1]), app.default_browser_cdp_user_data_dir())

    def test_browser_launches_skip_uninstalled_chrome_and_keep_edge(self):
        with mock.patch.object(app.platform, "system", return_value="Windows"), \
             mock.patch.object(app, "browser_is_installed", side_effect=lambda browser_key: browser_key == "edge", create=True):
            launches = app.browser_cdp_launches()

        self.assertEqual([launch.browser_name for launch in launches], ["Microsoft Edge"])

    def test_windows_chrome_only_keeps_its_independent_profile(self):
        with mock.patch.object(app.platform, "system", return_value="Windows"), \
             mock.patch.object(app, "browser_is_installed", side_effect=lambda browser_key: browser_key == "chrome", create=True):
            launches = app.browser_cdp_launches()

        self.assertEqual([launch.browser_name for launch in launches], ["Google Chrome"])
        self.assertEqual(Path(launches[0].profile_dir).name, "browser-cdp-profile-chrome")

    def test_browser_launches_move_to_fallback_ports_when_preferred_ports_are_occupied(self):
        with mock.patch.object(app.platform, "system", return_value="Darwin"), \
             mock.patch.object(app, "browser_is_installed", return_value=True, create=True):
            launches = app.browser_cdp_launches(occupied_ports={9222, 9223})

        self.assertEqual([launch.port for launch in launches], [9224, 9225])

    def test_parallel_browser_cookie_read_uses_first_success(self):
        reader = getattr(app, "read_browser_cookie_candidates", None)
        result_type = getattr(app, "BrowserCookieResult", None)
        self.assertIsNotNone(reader)
        self.assertIsNotNone(result_type)
        candidates = [
            app.BrowserCdpSession(False, "Microsoft Edge", browser_key="edge", port=9222),
            app.BrowserCdpSession(False, "Google Chrome", browser_key="chrome", port=9223),
        ]
        barrier = threading.Barrier(2)

        def read_candidate(candidate, stop_event, wait_seconds, poll_seconds):
            barrier.wait(timeout=1)
            if candidate.browser_key == "chrome":
                return result_type("vbkticket=chrome", 1, "chrome", "Google Chrome", 9223, "")
            self.assertTrue(stop_event.wait(1))
            raise RuntimeError("已取消")

        result = reader(candidates, reader=read_candidate, wait_seconds=1, poll_seconds=0)

        self.assertEqual(result.browser_key, "chrome")
        self.assertEqual(result.cookie_header, "vbkticket=chrome")

    def test_parallel_browser_cookie_read_continues_when_one_browser_fails(self):
        candidates = [
            app.BrowserCdpSession(False, "Microsoft Edge", browser_key="edge", port=9222),
            app.BrowserCdpSession(False, "Google Chrome", browser_key="chrome", port=9223),
        ]

        def read_candidate(candidate, stop_event, wait_seconds, poll_seconds):
            if candidate.browser_key == "edge":
                raise RuntimeError("启动失败")
            return app.BrowserCookieResult("vbkticket=chrome", 1, "chrome", "Chrome", 9223, "")

        result = app.read_browser_cookie_candidates(candidates, reader=read_candidate, wait_seconds=0, poll_seconds=0)

        self.assertEqual(result.browser_key, "chrome")

    def test_parallel_browser_cookie_read_accepts_edge_first(self):
        candidates = [
            app.BrowserCdpSession(False, "Microsoft Edge", browser_key="edge", port=9222),
            app.BrowserCdpSession(False, "Google Chrome", browser_key="chrome", port=9223),
        ]
        barrier = threading.Barrier(2)

        def read_candidate(candidate, stop_event, wait_seconds, poll_seconds):
            barrier.wait(timeout=1)
            if candidate.browser_key == "edge":
                return app.BrowserCookieResult("vbkticket=edge", 1, "edge", "Edge", 9222, "")
            self.assertTrue(stop_event.wait(1))
            raise RuntimeError("已取消")

        result = app.read_browser_cookie_candidates(candidates, reader=read_candidate, wait_seconds=1, poll_seconds=0)

        self.assertEqual(result.browser_key, "edge")

    def test_parallel_browser_cookie_read_reports_both_failures(self):
        reader = getattr(app, "read_browser_cookie_candidates", None)
        self.assertIsNotNone(reader)
        candidates = [
            app.BrowserCdpSession(False, "Microsoft Edge", browser_key="edge", port=9222),
            app.BrowserCdpSession(False, "Google Chrome", browser_key="chrome", port=9223),
        ]

        def read_candidate(candidate, stop_event, wait_seconds, poll_seconds):
            raise RuntimeError("未登录")

        with self.assertRaisesRegex(RuntimeError, "Microsoft Edge.*未登录.*Google Chrome.*未登录"):
            reader(candidates, reader=read_candidate, wait_seconds=0, poll_seconds=0)

    def test_parallel_browser_cookie_read_closes_tool_started_nonwinner(self):
        edge = app.BrowserCdpSession(True, "Microsoft Edge", browser_key="edge", port=9222)
        chrome = app.BrowserCdpSession(True, "Google Chrome", browser_key="chrome", port=9223)
        barrier = threading.Barrier(2)

        def read_candidate(candidate, stop_event, wait_seconds, poll_seconds):
            barrier.wait(timeout=1)
            if candidate.browser_key == "edge":
                return app.BrowserCookieResult("vbkticket=edge", 1, "edge", "Edge", 9222, "", session=candidate)
            self.assertTrue(stop_event.wait(1))
            return app.BrowserCookieResult("vbkticket=chrome", 1, "chrome", "Chrome", 9223, "", session=candidate)

        closed = []
        with mock.patch.object(app, "close_browser_cdp_session", side_effect=lambda session, cdp=None: closed.append(session.browser_name)):
            result = app.read_browser_cookie_candidates([edge, chrome], reader=read_candidate, wait_seconds=1, poll_seconds=0)

        self.assertEqual(result.browser_key, "edge")
        self.assertEqual(closed, ["Google Chrome"])

    def test_extract_browser_cookie_header_closes_auto_tab_after_reading(self):
        class FakeCdp:
            def __init__(self):
                self.closed = []

            def open_url(self, url):
                self.opened_url = url
                return "target-1"

            def get_cookies(self, target_id, urls):
                self.cookie_target = target_id
                return [{"name": "vbkticket", "value": "token", "domain": ".ctrip.com", "path": "/"}]

            def close(self, target_id):
                self.closed.append(target_id)

        fake = FakeCdp()
        cookie_header, count = app.extract_browser_cookie_header(fake, wait_seconds=0, poll_seconds=0)

        self.assertEqual(cookie_header, "vbkticket=token")
        self.assertEqual(count, 1)
        self.assertEqual(fake.closed, ["target-1"])

    def test_extract_browser_cookie_header_waits_for_login_cookie(self):
        class FakeCdp:
            def __init__(self):
                self.cookie_reads = 0

            def open_url(self, url):
                return "target-1"

            def get_cookies(self, target_id, urls):
                self.cookie_reads += 1
                if self.cookie_reads == 1:
                    return [{"name": "_bfa", "value": "tracking", "domain": ".ctrip.com", "path": "/"}]
                return [
                    {"name": "_bfa", "value": "tracking", "domain": ".ctrip.com", "path": "/"},
                    {"name": "vbkticket", "value": "token", "domain": ".ctrip.com", "path": "/"},
                ]

            def close(self, target_id):
                pass

        fake = FakeCdp()
        with mock.patch.object(app.time, "monotonic", side_effect=[0.0, 0.1, 0.2]), \
             mock.patch.object(app.time, "sleep"):
            cookie_header, count = app.extract_browser_cookie_header(fake, wait_seconds=1, poll_seconds=0.01)

        self.assertEqual(cookie_header, "_bfa=tracking; vbkticket=token")
        self.assertEqual(count, 2)
        self.assertEqual(fake.cookie_reads, 2)

    def test_extract_browser_cookie_header_keeps_winning_tool_browser_open(self):
        class FakeCdp:
            def __init__(self):
                self.closed_tabs = []
                self.browser_closed = False

            def open_url(self, url):
                self.opened_url = url
                return "target-1"

            def get_cookies(self, target_id, urls):
                return [{"name": "vbkticket", "value": "token", "domain": ".ctrip.com", "path": "/"}]

            def close(self, target_id):
                self.closed_tabs.append(target_id)

            def close_browser(self):
                self.browser_closed = True

        fake = FakeCdp()
        session = app.BrowserCdpSession(
            started_by_tool=True,
            browser_name="Microsoft Edge",
            process=None,
            browser_key="edge",
            port=9222,
            profile_dir="/tmp/edge-profile",
        )

        result = app.poll_browser_cookie(fake, session, threading.Event(), 0, 0)

        self.assertEqual(result.cookie_header, "vbkticket=token")
        self.assertEqual(result.cookie_count, 1)
        self.assertEqual(result.browser_name, "Edge")
        self.assertEqual(fake.closed_tabs, ["target-1"])
        self.assertFalse(fake.browser_closed)

    def test_extract_browser_cookie_header_closes_tool_started_browser_on_cookie_failure(self):
        class FakeCdp:
            def __init__(self):
                self.closed_tabs = []
                self.browser_closed = False

            def open_url(self, url):
                return "target-1"

            def get_cookies(self, target_id, urls):
                return []

            def close(self, target_id):
                self.closed_tabs.append(target_id)

            def close_browser(self):
                self.browser_closed = True

        fake = FakeCdp()
        session = app.BrowserCdpSession(
            started_by_tool=True,
            browser_name="Microsoft Edge",
            process=None,
            browser_key="edge",
            port=9222,
            profile_dir="/tmp/edge-profile",
        )

        with self.assertRaisesRegex(RuntimeError, "未检测到携程后台登录 Cookie"):
            app.poll_browser_cookie(fake, session, threading.Event(), 0, 0)

        self.assertEqual(fake.closed_tabs, ["target-1"])
        self.assertTrue(fake.browser_closed)

    def test_extract_browser_cookie_header_does_not_close_existing_browser_session(self):
        class FakeCdp:
            def __init__(self):
                self.closed_tabs = []
                self.browser_closed = False

            def open_url(self, url):
                return "target-1"

            def get_cookies(self, target_id, urls):
                return [{"name": "vbkticket", "value": "token", "domain": ".ctrip.com", "path": "/"}]

            def close(self, target_id):
                self.closed_tabs.append(target_id)

            def close_browser(self):
                self.browser_closed = True

        fake = FakeCdp()
        session = app.BrowserCdpSession(
            started_by_tool=False,
            browser_name="Google Chrome",
            process=None,
            browser_key="chrome",
            port=9223,
        )

        app.poll_browser_cookie(fake, session, threading.Event(), 0, 0)

        self.assertEqual(fake.closed_tabs, ["target-1"])
        self.assertFalse(fake.browser_closed)

    def test_cookie_header_from_cdp_cookies_is_stable_and_ctrip_only(self):
        cookies = [
            {"name": "z_cookie", "value": "3", "domain": ".ctrip.com", "path": "/"},
            {"name": "ignore", "value": "x", "domain": ".example.com", "path": "/"},
            {"name": "a_cookie", "value": "1", "domain": ".ctrip.com", "path": "/"},
            {"name": "a_cookie", "value": "2", "domain": ".vbooking.ctrip.com", "path": "/"},
            {"name": "", "value": "missing", "domain": ".ctrip.com", "path": "/"},
        ]

        self.assertEqual(
            app.cookie_header_from_cdp_cookies(cookies),
            "a_cookie=1; a_cookie=2; z_cookie=3",
        )

    def test_persist_plain_auth_does_not_overwrite_existing_cookie_when_validation_fails(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            auth_json_path = Path(tmpdir) / "ctrip_auth_plain.json"
            cookie_path.write_text("old_cookie=1", encoding="utf-8")

            with self.assertRaisesRegex(RuntimeError, "无权限"):
                app.refresh_plain_auth_with_validation(
                    cookie_header="new_cookie=2",
                    source="current_browser_cdp",
                    loaded_cookie_count=1,
                    validator=lambda _cookie: (_ for _ in ()).throw(RuntimeError("无权限")),
                    cookie_path=cookie_path,
                    auth_json_path=auth_json_path,
                )

            self.assertEqual(cookie_path.read_text(encoding="utf-8"), "old_cookie=1")
            self.assertFalse(auth_json_path.exists())

    def test_persist_plain_auth_writes_after_validation_passes(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            auth_json_path = Path(tmpdir) / "ctrip_auth_plain.json"

            path, count, source = app.refresh_plain_auth_with_validation(
                cookie_header="new_cookie=2",
                source="current_browser_cdp",
                loaded_cookie_count=1,
                validator=lambda cookie: None if cookie == "new_cookie=2" else (_ for _ in ()).throw(RuntimeError("wrong cookie")),
                cookie_path=cookie_path,
                auth_json_path=auth_json_path,
            )

            self.assertEqual(path, cookie_path)
            self.assertEqual(count, 1)
            self.assertEqual(source, "current_browser_cdp")
            self.assertEqual(cookie_path.read_text(encoding="utf-8"), "new_cookie=2")
            meta = json.loads(auth_json_path.read_text(encoding="utf-8"))
            self.assertEqual(meta["source"], "current_browser_cdp")
            self.assertEqual(meta["loadedTripDetailCookies"], 1)

    def test_refresh_plain_auth_without_product_does_not_run_product_validation(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            auth_json_path = Path(tmpdir) / "ctrip_auth_plain.json"
            browser_result = app.BrowserCookieResult(
                "vbkticket=token",
                1,
                "chrome",
                "Chrome",
                9223,
                "/tmp/chrome-profile",
            )
            with mock.patch.object(app, "extract_browser_cookie_result", return_value=browser_result), \
                 mock.patch.object(app, "validate_plain_auth_cookie") as validate:
                app.refresh_plain_auth(
                    validation_product_id="",
                    cookie_path=cookie_path,
                    auth_json_path=auth_json_path,
                )

            validate.assert_not_called()
            self.assertEqual(cookie_path.read_text(encoding="utf-8"), "vbkticket=token")
            metadata = json.loads(auth_json_path.read_text(encoding="utf-8"))
            self.assertEqual(metadata["browser"], "chrome")
            self.assertEqual(metadata["browserName"], "Chrome")
            self.assertEqual(metadata["cdpPort"], 9223)
            self.assertEqual(metadata["profileDir"], "/tmp/chrome-profile")
            metadata_text = json.dumps(metadata).lower()
            self.assertNotIn("vbkticket", metadata_text)
            self.assertNotIn("token", metadata_text)

    def test_refresh_plain_auth_with_product_runs_real_product_validation(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            cookie_path = Path(tmpdir) / "ctrip_cookie_header.txt"
            auth_json_path = Path(tmpdir) / "ctrip_auth_plain.json"
            browser_result = app.BrowserCookieResult(
                "vbkticket=token",
                1,
                "edge",
                "Edge",
                9222,
                "/tmp/edge-profile",
            )
            with mock.patch.object(app, "extract_browser_cookie_result", return_value=browser_result), \
                 mock.patch.object(app, "validate_plain_auth_cookie") as validate:
                app.refresh_plain_auth(
                    validation_product_id="7654321",
                    area="translation",
                    cookie_path=cookie_path,
                    auth_json_path=auth_json_path,
                )

            validate.assert_called_once_with("vbkticket=token", "7654321", "translation")

    def test_translation_auth_refresh_failure_hides_python_traceback(self):
        window = app.App()
        try:
            with mock.patch.object(window, "_refresh_auth_for_area", side_effect=RuntimeError("请在 Edge 登录携程后台")):
                window._run_translation_refresh_auth()

            self.assertEqual(window.out_queue.get_nowait(), ("auth_status", "登录态：刷新失败"))
            kind, message = window.out_queue.get_nowait()
            self.assertEqual(kind, "fatal")
            self.assertEqual(message, "请在 Edge 登录携程后台")
            self.assertNotIn("Traceback", message)
        finally:
            window.close()

    def test_auth_refresh_success_displays_winning_browser(self):
        window = app.App()
        try:
            with mock.patch.object(window, "_refresh_auth_for_area", return_value=(Path("cookie.txt"), 1, "Chrome")):
                window._run_translation_refresh_auth()

            self.assertEqual(window.out_queue.get_nowait(), ("auth_status", "登录态：Chrome 登录态有效"))
        finally:
            window.close()

    def test_cdp_guide_copies_both_browser_debugging_urls(self):
        window = app.App()
        try:
            with mock.patch.object(app.QMessageBox, "information"):
                window.show_cdp_guide()

            clipboard_text = app.QGuiApplication.clipboard().text()
            self.assertIn(app.EDGE_REMOTE_DEBUGGING_URL, clipboard_text)
            self.assertIn(app.CHROME_REMOTE_DEBUGGING_URL, clipboard_text)
        finally:
            window.close()

    def test_refresh_plain_auth_does_not_accept_legacy_session_file(self):
        with self.assertRaises(TypeError):
            app.refresh_plain_auth(session_file=Path("session_data.json"))


class TranslationReportDefaultNameTests(unittest.TestCase):
    def test_uses_task_template_stem_with_result_suffix(self):
        self.assertEqual(
            app.translation_report_default_filename("/Users/example/Downloads/任務模版_本地單地 10M恆速.xlsx"),
            "任務模版_本地單地 10M恆速_result.xlsx",
        )

    def test_falls_back_to_timestamp_name_without_task_template(self):
        self.assertEqual(
            app.translation_report_default_filename("", stamp_provider=lambda: "20260625_172233"),
            "ctrip_translation_report_20260625_172233.xlsx",
        )


class TranslationReportMetadataTests(unittest.TestCase):
    def test_xlsx_report_includes_execution_summary_sheet(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "translation_report.xlsx"
            results = [
                app.TranslationResultRow(
                    product_id=PRODUCT_ID,
                    locale="ar-XX",
                    module_code="basic",
                    field_code="productName",
                    field_name="产品名称",
                    self_key="k1",
                    status="saved",
                ),
                app.TranslationResultRow(
                    product_id=PRODUCT_ID,
                    locale="ar-XX",
                    module_code="basic",
                    field_code="feeInclude",
                    field_name="费用包含",
                    self_key="k2",
                    status="error",
                    message="无权限",
                ),
            ]
            metadata = app.build_translation_report_metadata(
                task_template_path="/Users/example/Downloads/task.xlsx",
                detail_template_path="/Users/example/Downloads/detail.xlsx",
                package_template_path="/Users/example/Downloads/package.xlsx",
                execution_summary="语言=ar | 模块=产品名称",
                results=results,
                request_stats={
                    "summary": "请求 累计 5/1 | 今日 3/1 | 本次 2/1",
                    "total_requests": 5,
                    "failed_requests": 1,
                    "today_requests": 3,
                    "today_failed_requests": 1,
                    "run_requests": 2,
                    "run_failed_requests": 1,
                },
            )

            app.write_translation_report(str(path), results, metadata=metadata)

            wb = load_workbook(path)
            self.assertIn("执行摘要", wb.sheetnames)
            rows = list(wb["执行摘要"].iter_rows(values_only=True))
            self.assertEqual(rows[0], ("项目", "值"))
            summary = dict(rows[1:])
            self.assertEqual(summary["任务模板"], "/Users/example/Downloads/task.xlsx")
            self.assertEqual(summary["详情模板"], "/Users/example/Downloads/detail.xlsx")
            self.assertEqual(summary["规格模板"], "/Users/example/Downloads/package.xlsx")
            self.assertEqual(summary["结果总数"], 2)
            self.assertEqual(summary["状态 saved"], 1)
            self.assertEqual(summary["状态 error"], 1)
            self.assertEqual(summary["本次请求数"], 2)
            self.assertEqual(summary["本次错误数"], 1)


class ProductSnapshotNeedsTests(unittest.TestCase):
    def test_price_snapshot_exports_permission_message_and_current_sale_status(self):
        client = ProductSnapshotFakeClient()
        client.load_product_price_stock = lambda _product_id, _stop_event: {
            "1001": {
                "packageId": "1001",
                "price": {
                    "basicInfoList": [{"resourceId": "2001", "isAllowVendorEditCostPrice": False}],
                    "resourcePriceList": [
                        {"resourceId": "2001", "priceList": []}
                    ],
                    "resourcePriceDraftList": [
                        {"resourceId": "2001", "priceList": [{"day": "00010101", "costPrice": 0, "sale": True}]}
                    ],
                },
                "stock": {"basicInfoList": [{"resourceId": "2001"}]},
            }
        }

        snapshot = app.build_product_snapshot(
            client,
            PRODUCT_ID,
            threading.Event(),
            app.ProductSnapshotNeeds(product_info=False, price_stock=True, currency_context=False),
        )

        sku = next(iter(snapshot.sku_snapshots.values()))
        self.assertFalse(sku.cost_price_editable)
        self.assertEqual(sku.cost_price, "请联系业务经理开放录入底价")
        self.assertEqual(sku.sale_status, "开")

    def test_package_only_snapshot_does_not_read_price_stock_or_product_info(self):
        client = ProductSnapshotFakeClient()

        snapshot = app.build_product_snapshot(
            client,
            PRODUCT_ID,
            threading.Event(),
            app.ProductSnapshotNeeds(product_info=False, price_stock=False, currency_context=False),
        )

        self.assertEqual(client.calls, [("load_product", PRODUCT_ID)])
        sku = next(iter(snapshot.sku_snapshots.values()))
        self.assertEqual(sku.price_error, "not_requested")
        self.assertEqual(sku.cost_price, "")

    def test_template_rows_only_request_price_stock_when_price_field_present(self):
        package_row = app.EditTemplateRow(
            row_number=2,
            product_id=PRODUCT_ID,
            package_id="1001",
            resource_id="2001",
            package_name="套餐 A",
            target_introduction="intro",
        )
        price_row = app.EditTemplateRow(
            row_number=3,
            product_id=PRODUCT_ID,
            package_id="1001",
            resource_id="2001",
            package_name="套餐 A",
            target_cost_price="1.23",
        )

        self.assertFalse(app.product_snapshot_needs_from_edit_rows([package_row]).price_stock)
        self.assertTrue(app.product_snapshot_needs_from_edit_rows([package_row, price_row]).price_stock)


def product_edit_snapshot():
    return app.ProductSnapshot(
        product_id=PRODUCT_ID,
        created_at="2026-06-30T10:00:00",
        draft=app.ProductDraft(
            product_id=PRODUCT_ID,
            target_id=PRODUCT_ID,
            package_info={"productBasicInfo": {"name": "旧产品名"}},
            package_resources=[],
            draft_raw={"productBasicInfo": {"name": "旧产品名"}},
        ),
        sku_snapshots={
            app.sku_key(PRODUCT_ID, "1001", "2001"): app.SkuSnapshot(
                product_id=PRODUCT_ID,
                package_id="1001",
                resource_id="2001",
                package_name="套餐 A",
                vendor_package_id="V-1",
                connection_code="C-1",
                introduction="旧套餐说明",
                cost_price="12.3",
                sale_status="开",
                price_and_stock={"resourceId": "2001", "price": {"costPrice": "12.3"}},
            )
        },
    )


def workbook_headers(path, sheet=app.EDIT_TEMPLATE_SHEET):
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[sheet]
        return [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    finally:
        wb.close()


class ProductEditOperationalSplitTests(unittest.TestCase):
    def test_product_edit_template_exports_single_wide_sheet(self):
        snapshots = {PRODUCT_ID: product_edit_snapshot()}
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "product_edit.xlsx"
            app.build_edit_template_workbook(str(path), snapshots, False, False, False, app.ProductWorkbench.PRODUCT_INFO)
            headers = workbook_headers(path)
            wb = load_workbook(path, read_only=True)
            try:
                sheet_names = wb.sheetnames
            finally:
                wb.close()

        self.assertEqual(
            headers,
            [
                "productId",
                "productName",
                "packageId",
                "resourceId",
                "packageName",
                "PackageNameMark",
                "old_product_name",
                "target_product_name",
                "old_cost_price",
                "target_cost_price",
                "old_sale_status",
                "target_sale_status",
                "old_vendor_package_id",
                "target_vendor_package_id",
                "old_connection_code",
                "target_connection_code",
                "old_introduction",
                "target_introduction",
            ],
        )
        self.assertNotIn(app.INTRO_FILL_SHEET, sheet_names)

    def test_legacy_package_template_reads_from_unified_product_edit_entry(self):
        snapshots = {PRODUCT_ID: product_edit_snapshot()}
        with tempfile.TemporaryDirectory() as tmpdir:
            package_path = Path(tmpdir) / "package.xlsx"
            app.build_edit_template_workbook(str(package_path), snapshots, True, True, False, app.ProductWorkbench.PACKAGE_CONNECT)
            wb = load_workbook(package_path)
            ws = wb[app.EDIT_TEMPLATE_SHEET]
            headers = [cell.value for cell in ws[1]]
            ws.cell(2, headers.index("target_vendor_package_id") + 1, "V-2")
            wb.save(package_path)

            rows = app.read_edit_template(str(package_path), app.ProductWorkbench.PRODUCT_INFO)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].product_id, PRODUCT_ID)
        self.assertEqual(rows[0].package_id, "1001")
        self.assertEqual(rows[0].resource_id, "2001")

    def test_product_edit_template_accepts_chinese_alias_headers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "alias.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = app.EDIT_TEMPLATE_SHEET
            ws.append(["备注", "产品ID", "目标底价", "资源ID", "套餐ID", "目标产品名称", "目标对接代码"])
            ws.append(["ignore", PRODUCT_ID, "9.99", "2001", "1001", "新产品名", "C-2"])
            wb.save(path)

            rows = app.read_edit_template(str(path), app.ProductWorkbench.PRODUCT_INFO)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].product_id, PRODUCT_ID)
        self.assertEqual(rows[0].target_cost_price, "9.99")
        self.assertEqual(rows[0].target_product_name, "新产品名")
        self.assertEqual(rows[0].target_connection_code, "C-2")

    def test_product_edit_conflicts_are_row_invalid(self):
        rows = [
            app.EditTemplateRow(2, PRODUCT_ID, "1001", "2001", "套餐 A", target_product_name="名称 A"),
            app.EditTemplateRow(3, PRODUCT_ID, "1001", "2001", "套餐 A", target_product_name="名称 B"),
            app.EditTemplateRow(4, PRODUCT_ID, "1001", "2001", "套餐 A", target_cost_price="9.99"),
            app.EditTemplateRow(5, PRODUCT_ID, "1001", "2001", "套餐 A", target_cost_price="8.88"),
        ]

        app.mark_edit_template_conflicts(rows)

        self.assertIn("多个目标产品名称", rows[0].validation_error)
        self.assertIn("多个目标产品名称", rows[1].validation_error)
        self.assertIn("同一 SKU", rows[2].validation_error)
        self.assertIn("同一 SKU", rows[3].validation_error)

    def test_product_info_field_specs_only_allow_product_name_first_phase(self):
        self.assertEqual(tuple(app.PRODUCT_INFO_FIELD_SPECS), ("productName",))
        self.assertEqual(app.PRODUCT_INFO_FIELD_SPECS["productName"]["label"], "产品名称")


def translation_raw(locale, value):
    return {
        "targetLocale": locale,
        "translationModules": [
            {
                "moduleCode": "basic",
                "moduleName": "基础信息",
                "fields": [
                    {
                        "fieldCode": "feeExclude",
                        "fieldName": "费用不含",
                        "selfKey": "feeExclude-key",
                        "translation": value,
                        "origin": "source",
                    }
                ],
            }
        ],
    }


def translation_snapshot_for(state):
    locales = list(state)
    return app.build_translation_snapshot_from_response(
        PRODUCT_ID,
        locales,
        {"translations": [copy.deepcopy(state[locale]) for locale in locales]},
        locale_options={locale: {"translatedLocales": [locale], "operableLocales": [locale]} for locale in locales},
    )


def operation_attention_translation_raw(locale, operation_value, reminder_value):
    return {
        "targetLocale": locale,
        "translationModules": [
            {
                "moduleCode": "basicInfo",
                "moduleName": "基本信息",
                "fields": [
                    {
                        "fieldCode": "postBookingTravelReminder",
                        "selfKey": "7-novalueKey-postBookingTravelReminder-basicInfo-basicInfo",
                        "origin": "相同中文原文",
                        "translation": reminder_value,
                    },
                    {
                        "fieldCode": "operationAttentionInfo",
                        "subItems": [
                            {
                                "fieldCode": "operationAttention",
                                "selfKey": "9-novalueKey-operationAttention-basicInfo-basicInfo",
                                "origin": "相同中文原文",
                                "translation": operation_value,
                                "ibuKey": "operationAttention",
                            }
                        ],
                    },
                ],
            }
        ],
    }


class OperationAttentionSnapshotTests(unittest.TestCase):
    def test_snapshot_exposes_operation_attention_child_as_independent_field(self):
        response = {
            "activityFieldMetaList": [
                {
                    "fieldCode": "operationAttentionInfo",
                    "fieldName": "携程处理注意事项",
                },
                {
                    "fieldCode": "operationAttention",
                    "fieldName": "携程处理注意事项",
                },
            ],
            "translations": [
                operation_attention_translation_raw(
                    "ar-XX",
                    "operation translated",
                    "reminder translated",
                )
            ],
        }

        snapshot = app.build_translation_snapshot_from_response(
            PRODUCT_ID,
            ["ar-XX"],
            response,
        )

        fields_by_code = {
            field.field_code: field
            for field in snapshot.fields.values()
        }
        self.assertNotIn("operationAttentionInfo", fields_by_code)
        operation_field = fields_by_code["operationAttention"]
        reminder_field = fields_by_code["postBookingTravelReminder"]
        self.assertEqual(operation_field.field_name, "携程处理注意事项")
        self.assertEqual(operation_field.module_code, "basicInfo")
        self.assertEqual(
            operation_field.self_key,
            "9-novalueKey-operationAttention-basicInfo-basicInfo",
        )
        self.assertEqual(operation_field.value_key, "translation")
        self.assertEqual(operation_field.old_translation, "operation translated")
        self.assertEqual(reminder_field.old_translation, "reminder translated")

    def test_rule_diff_targets_only_operation_attention(self):
        response = {
            "activityFieldMetaList": [
                {
                    "fieldCode": "operationAttention",
                    "fieldName": "携程处理注意事项",
                }
            ],
            "translations": [
                operation_attention_translation_raw(
                    "ar-XX",
                    "old operation",
                    "old reminder",
                )
            ],
        }
        snapshot = app.build_translation_snapshot_from_response(
            PRODUCT_ID,
            ["ar-XX"],
            response,
            locale_options={
                "ar-XX": {
                    "translatedLocales": ["ar-XX"],
                    "operableLocales": ["ar-XX"],
                }
            },
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            detail_path = Path(tmpdir) / "detail_templates.xlsx"
            translation_rules.create_detail_template_xlsx(
                str(detail_path),
                ["ar"],
            )
            wb = load_workbook(detail_path)
            ws = wb["详情模板"]
            headers = [cell.value for cell in ws[1]]
            operation_column = headers.index("携程处理注意事项") + 1
            ws.cell(2, operation_column, "new operation")
            wb.save(detail_path)

            task = translation_rules.TranslationRuleTask(
                row_number=2,
                product_id=PRODUCT_ID,
                source=PRODUCT_ID,
                row_data={"productId": PRODUCT_ID},
            )
            config = translation_rules.TranslationExecutionConfig(
                selected_languages=["ar"],
                selected_modules=["operation_attention"],
                online=False,
            )
            plans = app.build_translation_rule_diff_plans(
                {PRODUCT_ID: snapshot},
                [task],
                str(detail_path),
                "",
                config,
            )

        self.assertEqual(len(plans), 1)
        self.assertEqual(plans[0].field_code, "operationAttention")
        self.assertEqual(
            plans[0].self_key,
            "9-novalueKey-operationAttention-basicInfo-basicInfo",
        )
        self.assertEqual(plans[0].old_value, "old operation")
        self.assertEqual(plans[0].new_value, "new operation")
        self.assertNotEqual(plans[0].field_code, "postBookingTravelReminder")

    def test_rule_diff_reports_duplicate_english_reminder_and_keeps_korean_operation_plan(self):
        response = {
            "activityFieldMetaList": [
                {
                    "fieldCode": "postBookingTravelReminder",
                    "fieldName": "订后出行提醒",
                },
                {
                    "fieldCode": "operationAttention",
                    "fieldName": "携程处理注意事项",
                },
            ],
            "translations": [
                operation_attention_translation_raw(
                    "ko-KR",
                    "old operation",
                    "기존 한국어 안내",
                )
            ],
        }
        snapshot = app.build_translation_snapshot_from_response(
            PRODUCT_ID,
            ["ko-KR"],
            response,
            locale_options={
                "ko-KR": {
                    "translatedLocales": ["ko-KR"],
                    "operableLocales": ["ko-KR"],
                }
            },
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            detail_path = Path(tmpdir) / "detail_templates.xlsx"
            translation_rules.create_detail_template_xlsx(str(detail_path), ["en", "ko"])
            wb = load_workbook(detail_path)
            ws = wb["详情模板"]
            headers = [cell.value for cell in ws[1]]
            reminder_column = headers.index("订后出行提醒（支持格式填写）") + 1
            operation_column = headers.index("携程处理注意事项") + 1
            ws.cell(2, reminder_column, "English reminder")
            ws.cell(3, reminder_column, "English reminder")
            ws.cell(3, operation_column, "새 한국어 안내")
            wb.save(detail_path)

            task = translation_rules.TranslationRuleTask(
                row_number=2,
                product_id=PRODUCT_ID,
                source=PRODUCT_ID,
                row_data={"productId": PRODUCT_ID},
            )
            config = translation_rules.TranslationExecutionConfig(
                selected_languages=["ko"],
                selected_modules=["post_booking_reminder", "operation_attention"],
                online=False,
            )
            plans = app.build_translation_rule_diff_plans(
                {PRODUCT_ID: snapshot},
                [task],
                str(detail_path),
                "",
                config,
            )

        operation_plan = next(plan for plan in plans if plan.field_code == "operationAttention")
        reminder_plan = next(plan for plan in plans if plan.module_code == "post_booking_reminder")
        self.assertEqual(operation_plan.status, "pending")
        self.assertEqual(operation_plan.new_value, "새 한국어 안내")
        self.assertEqual(reminder_plan.status, "invalid")
        self.assertIn("与英文模板完全相同", reminder_plan.message)
        self.assertIn("保留线上现有翻译", reminder_plan.message)

    def test_apply_translation_updates_operation_child_without_touching_reminder(self):
        raw = operation_attention_translation_raw(
            "ar-XX",
            "old operation",
            "old reminder",
        )
        modules = copy.deepcopy(raw["translationModules"])

        changed = app.apply_translation_to_modules(
            modules,
            "9-novalueKey-operationAttention-basicInfo-basicInfo",
            "translation",
            "new operation",
        )

        basic_fields = modules[0]["fields"]
        reminder = basic_fields[0]
        operation = basic_fields[1]["subItems"][0]
        self.assertTrue(changed)
        self.assertEqual(operation["translation"], "new operation")
        self.assertEqual(reminder["translation"], "old reminder")
        self.assertEqual(basic_fields[1]["fieldCode"], "operationAttentionInfo")


class TranslationTemplateRemovedIdColumnTests(unittest.TestCase):
    def test_translation_template_exports_and_imports_without_removed_id_column(self):
        snapshot = translation_snapshot_for({"en-US": translation_raw("en-US", "old en")})
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "translation.xlsx"
            app.build_translation_template_workbook(str(path), {PRODUCT_ID: snapshot})

            wb = load_workbook(path)
            ws = wb[app.TRANSLATION_TEMPLATE_SHEET]
            headers = [cell.value for cell in ws[1]]
            self.assertNotIn("snapshot" + "Id", headers)
            target_col = headers.index("target_translation") + 1
            ws.cell(2, target_col).value = "new en"
            wb.save(path)
            wb.close()

            rows = app.read_translation_template(str(path))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].product_id, PRODUCT_ID)
        self.assertEqual(rows[0].target_translation, "new en")


def translation_plan(snapshot, locale, new_value):
    field = snapshot.fields[app.translation_field_key(PRODUCT_ID, locale, "feeExclude-key")]
    row = app.TranslationTemplateRow(
        row_number=2,
        product_id=PRODUCT_ID,
        locale=locale,
        module_code=field.module_code,
        field_code=field.field_code,
        field_name=field.field_name,
        value_key=field.value_key,
        self_key=field.self_key,
        origin=field.origin,
        old_translation=field.old_translation,
        target_translation=new_value,
        online=False,
    )
    return app.TranslationDiffPlan(
        row=row,
        product_id=PRODUCT_ID,
        locale=locale,
        module_code=field.module_code,
        field_code=field.field_code,
        field_name=field.field_name,
        self_key=field.self_key,
        old_value=field.old_translation,
        new_value=new_value,
        online=False,
    )


class TranslationFakeClient:
    def __init__(self, state):
        self.state = state
        self.list_calls = []
        self.save_calls = []

    def load_translation_list(self, product_id, locales, _stop_event):
        self.list_calls.append((product_id, list(locales)))
        return {"translations": [copy.deepcopy(self.state[locale]) for locale in locales]}

    def save_translation(self, product_id, locale, raw_translation, changed_plans, online, _stop_event):
        self.save_calls.append((product_id, locale, len(changed_plans), online))
        payload = {
            "productId": product_id,
            "locale": locale,
            "online": online,
            "translationModules": copy.deepcopy(raw_translation.get("translationModules") or []),
        }
        for plan in changed_plans:
            app.apply_translation_to_modules(payload["translationModules"], plan.self_key, plan.row.value_key, plan.new_value)
        self.state[locale]["translationModules"] = copy.deepcopy(payload["translationModules"])
        return payload, {"responseBaseData": {"success": True}}


class TranslationRequestCoalescingTests(unittest.TestCase):
    def test_apply_pending_keeps_unchanged_rows_for_export(self):
        state = {"en-US": translation_raw("en-US", "old en")}
        snapshot = translation_snapshot_for(state)
        unchanged_plan = translation_plan(snapshot, "en-US", "old en")
        unchanged_plan.status = "unchanged"
        unchanged_plan.message = "目标翻译与当前快照一致"
        out_queue = queue.Queue()
        worker = app.TranslationWorkflowWorker(
            out_queue,
            threading.Event(),
            {PRODUCT_ID: snapshot},
            [unchanged_plan],
            set(),
        )

        worker.run_apply_pending()

        result_rows = []
        done_rows = []
        while not out_queue.empty():
            event, payload = out_queue.get()
            if event == "result_batch":
                result_rows.extend(payload)
            elif event == "result":
                result_rows.append(payload)
            elif event == "done":
                done_rows = payload

        self.assertEqual([row.status for row in result_rows], ["unchanged"])
        self.assertEqual([row.status for row in done_rows], ["unchanged"])

    def test_same_product_multiple_locales_use_one_pre_read_and_one_verify_read(self):
        state = {
            "en-US": translation_raw("en-US", "old en"),
            "ja-JP": translation_raw("ja-JP", "old ja"),
        }
        snapshot = translation_snapshot_for(state)
        worker = app.TranslationWorkflowWorker(
            queue.Queue(),
            threading.Event(),
            {PRODUCT_ID: snapshot},
            [],
            set(),
        )
        client = TranslationFakeClient(state)
        plans = {
            "en-US": [translation_plan(snapshot, "en-US", "new en")],
            "ja-JP": [translation_plan(snapshot, "ja-JP", "new ja")],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            saved_keys = worker._apply_translation_product_groups(
                PRODUCT_ID,
                plans,
                [],
                Path(tmpdir),
                client,
            )

        self.assertEqual(saved_keys, {f"{PRODUCT_ID}:en-US", f"{PRODUCT_ID}:ja-JP"})
        self.assertEqual(
            client.list_calls,
            [
                (PRODUCT_ID, ["en-US", "ja-JP"]),
                (PRODUCT_ID, ["en-US", "ja-JP"]),
            ],
        )
        self.assertEqual(
            client.save_calls,
            [
                (PRODUCT_ID, "en-US", 1, False),
                (PRODUCT_ID, "ja-JP", 1, False),
            ],
        )
        self.assertTrue(all(result.status == "saved" for result in worker.results))

    def test_one_locale_save_failure_does_not_mark_other_locale_as_error(self):
        state = {
            "zh-TW": translation_raw("zh-TW", "old zh"),
            "ko-KR": translation_raw("ko-KR", "old ko"),
        }
        snapshot = translation_snapshot_for(state)
        worker = app.TranslationWorkflowWorker(
            queue.Queue(),
            threading.Event(),
            {PRODUCT_ID: snapshot},
            [],
            set(),
        )

        class SelectiveFailureClient(TranslationFakeClient):
            def save_translation(self, product_id, locale, raw_translation, changed_plans, online, stop_event):
                payload, response = super().save_translation(
                    product_id,
                    locale,
                    raw_translation,
                    changed_plans,
                    online,
                    stop_event,
                )
                if locale == "ko-KR":
                    response = {
                        "responseBaseData": {
                            "success": False,
                            "errorCode": "1000004",
                            "message": "当前翻译语言为ko-KR，录入信息为en-US",
                        }
                    }
                return payload, response

        client = SelectiveFailureClient(state)
        plans = {
            "zh-TW": [translation_plan(snapshot, "zh-TW", "new zh")],
            "ko-KR": [translation_plan(snapshot, "ko-KR", "new ko")],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            saved_keys = worker._apply_translation_product_groups(
                PRODUCT_ID,
                plans,
                [],
                Path(tmpdir),
                client,
            )

        self.assertEqual(saved_keys, {f"{PRODUCT_ID}:zh-TW"})
        self.assertEqual(
            {result.locale: result.status for result in worker.results},
            {"zh-TW": "saved", "ko-KR": "error"},
        )
        ko_result = next(result for result in worker.results if result.locale == "ko-KR")
        self.assertIn("1000004", ko_result.message)


class PackageTemplateDescriptorTests(unittest.TestCase):
    def test_package_info_from_origin_preserves_capacity_descriptor_for_unlimited_speed(self):
        origin = "套餐名称：QR code-5天-5G网络-每日-100GB-不限速"

        self.assertEqual(
            translation_rules.package_info_from_origin_with_descriptor(origin),
            ("5", "每日", "100GB", "100GB-不限速"),
        )
        self.assertEqual(
            translation_rules.package_info_from_origin(origin),
            ("5", "每日", "100GB"),
        )

    def test_package_template_text_uses_descriptor_for_template_selection_only(self):
        package_template = {
            "combination_template": {
                "daily": "{usedup_throttle}\n{Day_calc}\n{inst_adv}",
                "total": "{usedup_throttle}\n{Day_calc}\n{inst_adv}",
            },
            "scripts": {
                "后降速": {
                    "daily_100gb": {"ko": "STANDARD {data_capacity}"},
                    "daily_100gb_unlimited": {"ko": "UNLIMITED {data_capacity}"},
                    "daily_generic": {"ko": "GENERIC {data_capacity}"},
                    "total": {"ko": "TOTAL {data_capacity}"},
                },
            },
            "install_instructions": {"pre_install_allowed": {"ko": "INSTALL"}},
            "day_calc": {"ko": "DAY"},
            "reset_time_calc": {},
        }
        config = translation_rules.TranslationExecutionConfig(
            selected_modules=["package_detail"],
            package_special_rules=translation_rules.normalize_special_rule_config({}),
        )
        row_data = {"{usedup_throttle}": "后降速", "{inst_adv}": "可提前安装"}

        text = translation_rules.package_template_text(
            package_template,
            "ko-KR",
            "100GB",
            "每日",
            row_data,
            config,
            "5",
            capacity_descriptor="100GB-不限速",
        )

        self.assertIn("UNLIMITED 100GB", text)
        self.assertNotIn("STANDARD", text)
        self.assertNotIn("100GB-不限速", text)

    def test_package_template_text_keeps_standard_100gb_without_descriptor(self):
        package_template = {
            "combination_template": {
                "daily": "{usedup_throttle}\n{Day_calc}\n{inst_adv}",
                "total": "{usedup_throttle}\n{Day_calc}\n{inst_adv}",
            },
            "scripts": {
                "后降速": {
                    "daily_100gb": {"ko": "STANDARD {data_capacity}"},
                    "daily_100gb_unlimited": {"ko": "UNLIMITED {data_capacity}"},
                    "daily_generic": {"ko": "GENERIC {data_capacity}"},
                    "total": {"ko": "TOTAL {data_capacity}"},
                },
            },
            "install_instructions": {"pre_install_allowed": {"ko": "INSTALL"}},
            "day_calc": {"ko": "DAY"},
            "reset_time_calc": {},
        }
        config = translation_rules.TranslationExecutionConfig(
            selected_modules=["package_detail"],
            package_special_rules=translation_rules.normalize_special_rule_config({}),
        )
        row_data = {"{usedup_throttle}": "后降速", "{inst_adv}": "可提前安装"}

        text = translation_rules.package_template_text(
            package_template,
            "ko-KR",
            "100GB",
            "每日",
            row_data,
            config,
            "5",
        )

        self.assertIn("STANDARD 100GB", text)
        self.assertNotIn("UNLIMITED", text)


class DetailTemplateContractTests(unittest.TestCase):
    def test_rule_diff_uses_preload_error_instead_of_missing_snapshot_message(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            detail_path = Path(tmpdir) / "detail_templates.xlsx"
            translation_rules.create_detail_template_xlsx(str(detail_path), ["ar"])
            task = translation_rules.TranslationRuleTask(
                row_number=2,
                product_id=PRODUCT_ID,
                source=PRODUCT_ID,
                row_data={"productId": PRODUCT_ID},
            )

            plans = app.build_translation_rule_diff_plans(
                {},
                [task],
                str(detail_path),
                "",
                app.default_translation_execution_config(),
                preload_errors={PRODUCT_ID: "读取翻译件可用语言失败：无权限 | LOGIN_ERROR"},
            )

            self.assertEqual(len(plans), 1)
            self.assertEqual(plans[0].status, "invalid")
            self.assertIn("读取翻译件可用语言失败", plans[0].message)
            self.assertNotIn("内存翻译快照不存在", plans[0].message)

    def test_detail_template_export_excludes_task_template_fields(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            detail_path = Path(tmpdir) / "detail_templates.xlsx"
            task_path = Path(tmpdir) / "translation_rule_task_template.xlsx"

            translation_rules.create_detail_template_xlsx(str(detail_path), ["en"])
            translation_rules.create_translation_task_template_xlsx(str(task_path))

            detail_ws = load_workbook(detail_path, read_only=True)["详情模板"]
            detail_headers = [detail_ws.cell(1, col).value for col in range(1, detail_ws.max_column + 1)]
            task_ws = load_workbook(task_path, read_only=True)["数据模板"]
            task_headers = [task_ws.cell(1, col).value for col in range(1, task_ws.max_column + 1)]

            self.assertEqual(
                detail_headers,
                ["语言", "使用方法", "产品说明（支持格式填写）", "注意事项（支持格式填写）", "订后出行提醒（支持格式填写）", "携程处理注意事项"],
            )
            self.assertNotIn("产品名称", detail_headers)
            self.assertNotIn("图文详情", detail_headers)
            self.assertNotIn("费用包含", detail_headers)
            self.assertNotIn("费用不包含", detail_headers)
            self.assertIn("{product_name_en}", task_headers)
            self.assertIn("{mm_detail_en}", task_headers)

    def test_detail_template_exports_and_loads_operation_attention_as_plain_text(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "detail_templates.xlsx"
            translation_rules.create_detail_template_xlsx(str(path), ["en"])

            wb = load_workbook(path)
            ws = wb["详情模板"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(
                headers,
                [
                    "语言",
                    "使用方法",
                    "产品说明（支持格式填写）",
                    "注意事项（支持格式填写）",
                    "订后出行提醒（支持格式填写）",
                    "携程处理注意事项",
                ],
            )
            operation_column = headers.index("携程处理注意事项") + 1
            ws.cell(2, operation_column, "Contact {contact_en}\nKeep screenshots")
            wb.save(path)

            data = translation_rules.load_detail_template_xlsx(str(path))

            self.assertEqual(
                data["en"]["operation_attention"],
                "Contact {contact_en}\nKeep screenshots",
            )
            self.assertIsInstance(data["en"]["operation_attention"], str)
            self.assertNotIn("operation_attention", translation_rules.DETAIL_RICH_KEYS)

    def test_legacy_detail_template_without_operation_attention_still_loads(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "legacy_detail_templates.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "详情模板"
            ws.append(["语言", "使用方法", "产品说明", "注意事项", "订后出行提醒"])
            ws.append(["en", "Use", "Description", "Notice", "Reminder"])
            wb.save(path)

            data = translation_rules.load_detail_template_xlsx(str(path))

            self.assertEqual(data["en"]["operation_attention"], "")

    def test_detail_template_viewer_displays_operation_attention(self):
        viewer = app.TranslationDetailTemplateViewer(
            {
                "en": {
                    "usage": "Use",
                    "description": [("Description", [])],
                    "notice": [("Notice", [])],
                    "post_booking_reminder": [("Reminder", [])],
                    "operation_attention": "Contact support@example.com",
                }
            },
            "/tmp/detail_templates.xlsx",
        )
        try:
            text = viewer.findChild(app.QTextEdit)
            self.assertIsNotNone(text)
            self.assertIn("携程处理注意事项:", text.toPlainText())
            self.assertIn("Contact support@example.com", text.toPlainText())
        finally:
            viewer.close()
            QT_APP.processEvents()

    def test_generated_targets_map_operation_attention_without_html_wrapping(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "detail_templates.xlsx"
            translation_rules.create_detail_template_xlsx(str(path), ["en"])
            wb = load_workbook(path)
            ws = wb["详情模板"]
            headers = [cell.value for cell in ws[1]]
            operation_column = headers.index("携程处理注意事项") + 1
            ws.cell(2, operation_column, "Contact {contact_en}\nKeep screenshots")
            wb.save(path)

            store = translation_rules.DetailTemplateExcelStore()
            store.load_template_from_path(str(path))
            targets = translation_rules.generated_translation_static_targets(
                store,
                "en-US",
                {"{contact_en}": "support@example.com"},
                ["operation_attention"],
            )

            self.assertEqual(
                targets,
                {
                    "operationAttention": (
                        "Contact support@example.com\nKeep screenshots"
                    )
                },
            )
            self.assertNotIn("<p", targets["operationAttention"])

    def test_non_english_duplicate_reminder_is_skipped_but_operation_attention_remains(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "detail_templates.xlsx"
            translation_rules.create_detail_template_xlsx(str(path), ["en", "ko"])
            wb = load_workbook(path)
            ws = wb["详情模板"]
            headers = [cell.value for cell in ws[1]]
            reminder_column = headers.index("订后出行提醒（支持格式填写）") + 1
            operation_column = headers.index("携程处理注意事项") + 1
            ws.cell(2, reminder_column, "English reminder")
            ws.cell(3, reminder_column, "English reminder")
            ws.cell(2, operation_column, "Contact support")
            ws.cell(3, operation_column, "한국어 고객센터 안내")
            wb.save(path)

            store = translation_rules.DetailTemplateExcelStore()
            store.load_template_from_path(str(path))
            targets = translation_rules.generated_translation_static_targets(
                store,
                "ko-KR",
                {},
                ["post_booking_reminder", "operation_attention"],
            )

            self.assertNotIn("postBookingTravelReminder", targets)
            self.assertEqual(targets["operationAttention"], "한국어 고객센터 안내")

    def test_detail_template_loader_ignores_legacy_hidden_fields(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "legacy_detail_templates.xlsx"
            translation_rules.create_detail_template_xlsx(str(path), ["en"])

            wb = load_workbook(path)
            ws = wb["详情模板"]
            for header, value in [
                ("产品名称", "Legacy product name"),
                ("图文详情", "Legacy graphic detail"),
                ("费用包含", "Legacy fee include"),
                ("费用不包含", "Legacy fee exclude"),
            ]:
                column = ws.max_column + 1
                ws.cell(1, column, header)
                ws.cell(2, column, value)
            wb.save(path)

            data = translation_rules.load_detail_template_xlsx(str(path))

            self.assertNotIn("product_name", data["en"])
            self.assertNotIn("mm_detail", data["en"])
            self.assertNotIn("fee_include", data["en"])
            self.assertNotIn("fee_exclude", data["en"])


if __name__ == "__main__":
    unittest.main()
