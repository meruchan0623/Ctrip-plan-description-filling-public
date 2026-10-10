import unittest

import translation_rules


class PackageTemplateTextTests(unittest.TestCase):
    def test_reset_options_are_not_used_as_hidden_fallback(self):
        template = translation_rules.build_default_package_template_schema()
        template["combination_template"] = {
            "daily": "● {usedup_throttle}\n● {reset_time_calc}\n● {Day_calc}\n● {inst_adv}",
            "total": "● {usedup_throttle}\n● {Day_calc}\n● {inst_adv}",
        }
        template["scripts"]["后降速"]["daily_generic"]["en"] = "{data_capacity} per day"
        template["day_calc"]["en"] = "Valid for purchased days"
        template["install_instructions"]["pre_install_allowed"]["en"] = "Can be installed in advance"
        template["reset_time_options"]["en"] = ["Hidden reset option"]

        text = translation_rules.package_template_text(
            template,
            "en-US",
            "1GB",
            "每日",
            {"{usedup_throttle}": "后降速", "{inst_adv}": "可提前安装"},
            translation_rules.TranslationExecutionConfig(),
            "1",
        )

        self.assertNotIn("Hidden reset option", text)
        self.assertIn("<p>● 1GB per day</p>", text)


if __name__ == "__main__":
    unittest.main()


class DetailTemplateRichTextTests(unittest.TestCase):
    def test_excel_rich_runs_are_pushed_as_html_formatting(self):
        import os
        import tempfile

        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from openpyxl.styles import Font

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = translation_rules.DETAIL_SHEET_NAME
        sheet.append([name for name, _key in translation_rules.DETAIL_COLUMNS])
        sheet.insert_cols(1)
        sheet.cell(1, 1, "语言")
        sheet.cell(2, 1, "en")
        sheet.cell(2, 3).value = CellRichText([
            "Intro ", TextBlock(InlineFont(b=True), "Bold line\nnext"), " & ", TextBlock(InlineFont(i=True, u="single"), "<tip>"),
        ])
        sheet.cell(2, 4, "Whole cell\nbold").font = Font(bold=True)
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "detail.xlsx")
            workbook.save(path)
            data = translation_rules.load_detail_template_xlsx(path)

        store = translation_rules.DetailTemplateExcelStore()
        self.assertEqual(
            store.convert_formatted_to_html(data["en"]["description"], "en-US"),
            "<p>Intro <strong>Bold line</strong></p><p><strong>next</strong> &amp; &lt;tip&gt;</p>",
        )
        self.assertEqual(store.convert_formatted_to_html(data["en"]["notice"], "en-US"), "<p>Whole cell</p><p>bold</p>")
        self.assertEqual(
            store.convert_formatted_to_html([("Hi", ["bold"])], "ar-XX"),
            '<p dir="rtl" style="text-align:right;"><strong>Hi</strong></p>',
        )


class QuillMarkupTests(unittest.TestCase):
    def setUp(self):
        self.store = translation_rules.DetailTemplateExcelStore()

    def test_heading_bold_bullet_and_link_match_quill_html(self):
        data = [
            ("## Before Use\n- item ", []), ("one", ["bold"]),
            ("\n* see [site](example.com/a?b=1&c=2)\n\nplain **bold** [mail](mailto:a@b.c)", []),
        ]
        self.assertEqual(
            self.store.convert_formatted_to_html(data, "en-US"),
            '<h2>Before Use</h2><ul><li>item <strong>one</strong></li>'
            '<li>see <a href="https://example.com/a?b=1&amp;c=2" rel="noopener noreferrer" target="_blank">site</a></li></ul>'
            '<p><br></p><p>plain <strong>bold</strong> '
            '<a href="mailto:a@b.c" rel="noopener noreferrer" target="_blank">mail</a></p>',
        )

    def test_rtl_attributes_apply_to_every_block(self):
        self.assertEqual(
            self.store.convert_formatted_to_html([("# عنوان\n- أ", [])], "ar-XX"),
            '<h2 dir="rtl" style="text-align:right;">عنوان</h2><ul><li dir="rtl" style="text-align:right;">أ</li></ul>',
        )

    def test_plain_symbols_are_not_treated_as_markup(self):
        data = [("• Coverage: {x}\n#1 hotline\n-5°C * 2\n[note] (text) & <b>", [])]
        self.assertEqual(
            self.store.convert_formatted_to_html(data, "en-US"),
            "<p>• Coverage: {x}</p><p>#1 hotline</p><p>-5°C * 2</p><p>[note] (text) &amp; &lt;b&gt;</p>",
        )


class DetailTemplateFormatGuideTests(unittest.TestCase):
    def test_generated_template_documents_rich_text_markup(self):
        import os
        import tempfile

        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "detail.xlsx")
            translation_rules.create_detail_template_xlsx(path, ["en", "ar"])
            workbook = openpyxl.load_workbook(path)
            data = translation_rules.load_detail_template_xlsx(path)

        self.assertEqual(set(data), {"en", "ar"})
        self.assertEqual(workbook.sheetnames, [translation_rules.DETAIL_SHEET_NAME, translation_rules.DETAIL_FORMAT_GUIDE_SHEET_NAME])
        headers = {cell.value: cell for cell in workbook[translation_rules.DETAIL_SHEET_NAME][1]}
        self.assertEqual(
            list(headers),
            ["语言", "使用方法", "产品说明（支持格式填写）", "注意事项（支持格式填写）", "订后出行提醒（支持格式填写）", "携程处理注意事项"],
        )
        for rich_header in ("产品说明（支持格式填写）", "注意事项（支持格式填写）", "订后出行提醒（支持格式填写）"):
            self.assertIn("## 标题", headers[rich_header].comment.text)
        for plain_header in ("使用方法", "携程处理注意事项"):
            self.assertIsNone(headers[plain_header].comment)
        guide_text = "\n".join(str(cell.value) for row in workbook[translation_rules.DETAIL_FORMAT_GUIDE_SHEET_NAME].iter_rows() for cell in row if cell.value)
        for token in ("## 使用前须知", "**出发前**", "- 支持热点共享", "[联系客服](www.trip.com/help)"):
            self.assertIn(token, guide_text)

    def test_guide_examples_render_the_documented_tags(self):
        store = translation_rules.DetailTemplateExcelStore()
        for _effect, _how, example, shown in translation_rules.DETAIL_FORMAT_GUIDE_ROWS:
            html = store.convert_formatted_to_html([(example, [])], "en-US")
            tag = shown.split("<", 1)[1].split(">", 1)[0].split("><")[0]
            self.assertIn(f"<{tag}", html, example)
